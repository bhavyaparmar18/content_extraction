"""Tests for the v2.0 template extraction contract.

Covers the fidelity fixes that have no sample-document coverage: merged-cell
resolution, ``w:shd`` shading, content-hash icon deduplication, and the
shaded-versus-unshaded branching of the icon+instruction layout container.
"""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, RGBColor
from PIL import Image

from app.config.settings import Settings
from app.services.extraction.template_extractor import TemplateExtractionService
from app.services.parser.docx_parser import DocxParser
from app.services.parser.template_parser import TemplateDocxParser

BLUE = RGBColor(0x00, 0x75, 0xFF)


# ── Fixtures / OOXML helpers ───────────────────────────────────────────

@pytest.fixture
def settings(tmp_path):
    """Settings rooted in tmp_path so extraction artefacts stay isolated."""
    s = Settings(project_root=tmp_path)
    s.resolve_paths(tmp_path)
    s.ensure_directories()
    return s


def _shade_cell(cell, fill: str):
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def _shade_paragraph(paragraph, fill: str):
    paragraph._element.get_or_add_pPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def _set_left_border(cell, color: str):
    cell._element.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{color}"/>'
            f'</w:tcBorders>'
        )
    )


def _set_text_direction(cell, val: str):
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:textDirection {nsdecls("w")} w:val="{val}"/>')
    )


def _blue_paragraph(doc, text: str):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.color.rgb = BLUE
    return para


def _make_png(path: Path, color: tuple[int, int, int]) -> Path:
    Image.new("RGB", (24, 24), color).save(path)
    return path


def _put_icon(cell, png_path: Path):
    cell.paragraphs[0].add_run().add_picture(str(png_path), width=Inches(0.22))


# ── A1: merged-cell resolution ─────────────────────────────────────────

def test_merged_cells_resolve_to_single_origin(tmp_path, settings):
    """The role matrix: one vertical and one horizontal merge in the header."""
    doc = Document()
    table = doc.add_table(rows=3, cols=6)
    table.cell(0, 1).merge(table.cell(0, 5))   # "Tasks for" spans 5 columns
    table.cell(0, 0).merge(table.cell(1, 0))   # "Process Step" spans 2 rows
    table.cell(0, 0).text = "Process Step"
    table.cell(0, 1).text = "Tasks for"
    for col in range(1, 6):
        table.cell(1, col).text = f"[Role {col}]"
    table.cell(2, 0).text = "Step 1"

    path = tmp_path / "merges.docx"
    doc.save(str(path))

    raw = DocxParser(settings=settings).parse(str(path))
    tables = [
        el
        for page in raw.pages
        for el in page.elements
        if getattr(el, "grid_cols", None) is not None
    ]
    assert len(tables) == 1
    ext = tables[0]

    assert ext.grid_cols == 6

    header = ext.headers
    assert len(header) == 6
    assert header[0].content_text == "Process Step"
    assert header[0].is_merge_origin
    assert header[0].row_span == 2
    assert header[0].col_span == 1

    assert header[1].content_text == "Tasks for"
    assert header[1].is_merge_origin
    assert header[1].col_span == 5

    # The remaining header positions are continuations of "Tasks for"
    for cell in header[2:]:
        assert not cell.is_merge_origin
        assert cell.merge_origin_ref == "horizontal"
        assert cell.content_text == ""
        assert cell.col_span == 1

    # Grid row 1: col 0 continues the vertical merge, cols 1-5 are the roles
    role_row = ext.rows[0]
    assert not role_row[0].is_merge_origin
    assert role_row[0].merge_origin_ref == "vertical"
    assert [c.content_text for c in role_row[1:]] == [f"[Role {i}]" for i in range(1, 6)]

    all_cells = [c for row in [ext.headers, *ext.rows] for c in row]
    assert [c.content_text for c in all_cells].count("Process Step") == 1
    assert [c.content_text for c in all_cells].count("Tasks for") == 1


# ── A2: w:shd shading ──────────────────────────────────────────────────

def test_paragraph_and_cell_shading_extracted(tmp_path, settings):
    doc = Document()
    _shade_paragraph(doc.add_paragraph("Shaded prose"), "E2EFDA")
    _shade_paragraph(doc.add_paragraph("Auto fill prose"), "auto")
    doc.add_paragraph("Plain prose")

    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "shaded cell"
    _shade_cell(table.cell(0, 0), "FCE4D6")
    table.cell(0, 1).text = "nil cell"
    _shade_cell(table.cell(0, 1), "nil")
    table.cell(0, 2).text = "plain cell"

    path = tmp_path / "shading.docx"
    doc.save(str(path))

    raw = DocxParser(settings=settings).parse(str(path))
    elements = [el for page in raw.pages for el in page.elements]
    by_text = {el.content: el for el in elements if el.content}

    assert by_text["Shaded prose"].shading_hex == "E2EFDA"
    assert by_text["Auto fill prose"].shading_hex is None
    assert by_text["Plain prose"].shading_hex is None

    ext_table = next(el for el in elements if getattr(el, "grid_cols", None))
    cells = {c.content_text: c for c in ext_table.headers}
    assert cells["shaded cell"].shading_hex == "FCE4D6"
    assert cells["nil cell"].shading_hex is None
    assert cells["plain cell"].shading_hex is None


# ── A5: table formatting attributes ────────────────────────────────────

def test_table_formatting_attributes_extracted(tmp_path, settings):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Role"
    _set_text_direction(table.cell(0, 0), "btLr")
    table.cell(0, 0)._element.get_or_add_tcPr().append(
        parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>')
    )
    table.cell(0, 1).paragraphs[0].add_run("Bold header").bold = True
    table.rows[0]._tr.get_or_add_trPr().append(
        parse_xml(f'<w:tblHeader {nsdecls("w")}/>')
    )

    path = tmp_path / "formatting.docx"
    doc.save(str(path))

    raw = DocxParser(settings=settings).parse(str(path))
    ext_table = next(
        el
        for page in raw.pages
        for el in page.elements
        if getattr(el, "grid_cols", None)
    )

    assert ext_table.header_rows == 1
    assert len(ext_table.col_widths_pt) == 2
    assert all(width > 0 for width in ext_table.col_widths_pt)
    assert ext_table.headers[0].text_direction == "btLr"
    assert ext_table.headers[0].valign == "center"
    assert ext_table.headers[1].bold is True


# ── A4: detection_method replaces the STYLE_INSTRUCTION sentinel ───────

def test_blue_detection_reports_method_and_hex_only_colour(tmp_path, settings):
    doc = Document()
    _blue_paragraph(doc, "Follow this instruction.")
    doc.add_paragraph("Regular body text.")

    path = tmp_path / "blue.docx"
    doc.save(str(path))

    raw = TemplateDocxParser(settings=settings).parse(str(path), template_id="tpl")
    elements = [el for page in raw.pages for el in page.elements]
    blue = next(el for el in elements if el.content == "Follow this instruction.")
    plain = next(el for el in elements if el.content == "Regular body text.")

    assert blue.color_detection_method == "run_color"
    assert blue.font_color_hex == "0075FF"
    assert blue.metadata["is_instruction"] is True
    assert plain.color_detection_method is None
    assert plain.font_color_hex is None


# ── End-to-end v2.0 output ─────────────────────────────────────────────

def _build_template_docx(path: Path, icon_a: Path, icon_b: Path) -> Path:
    """A miniature template exercising global rules, callouts and icon reuse."""
    doc = Document()

    # Global rules — blue text before the first section heading
    _blue_paragraph(doc, "Change the font - Arial is the official font to be used")
    _blue_paragraph(
        doc, "Follow the instructions in blue text and delete the blue text before finalization."
    )
    _blue_paragraph(doc, "Do not change Header and Footer")

    doc.add_heading("1 PURPOSE", level=1)
    doc.add_paragraph("This section describes the purpose for [Role 1].")

    # Unshaded icon+instruction layout container — should be unwrapped
    container = doc.add_table(rows=2, cols=2)
    for row_idx, text in enumerate(
        ["Brief description of what the document is about?", "Explain the target roles."]
    ):
        _put_icon(container.cell(row_idx, 0), icon_a)
        run = container.cell(row_idx, 1).paragraphs[0].add_run(text)
        run.font.color.rgb = BLUE

    doc.add_heading("6 PROCESS", level=1)
    _blue_paragraph(
        doc, 'Use "Explanation" infographic in the document as needed to accompany the text'
    )

    # Shaded infographic box — should become a callout
    callout = doc.add_table(rows=1, cols=2)
    _put_icon(callout.cell(0, 0), icon_b)
    callout.cell(0, 1).text = "Explanation text goes here."
    for cell in callout.rows[0].cells:
        _shade_cell(cell, "E2EFDA")
        _set_left_border(cell, "385723")

    doc.save(str(path))
    return path


@pytest.mark.asyncio
async def test_v2_output_contract(tmp_path, settings):
    icon_a = _make_png(tmp_path / "icon_a.png", (200, 40, 40))
    icon_b = _make_png(tmp_path / "icon_b.png", (40, 160, 60))
    template_path = _build_template_docx(tmp_path / "mini_template.docx", icon_a, icon_b)

    service = TemplateExtractionService(settings=settings)
    output = await service.extract("mini_tpl", template_path, template_name="Mini Template")

    assert output.version == "2.0"
    assert output.totals.sections == len(output.sections)
    assert output.totals.instructions == len(output.global_rules.instructions) + sum(
        len(s.authoring_instructions) for s in output.sections
    )

    # Global rules carry directive types and enforceable machine rules
    rules = {r.machine_rule.rule: r for r in output.global_rules.instructions if r.machine_rule}
    assert rules["font_family"].machine_rule.value == "Arial"
    assert rules["strip_blue_text"].machine_rule.value is True
    assert rules["preserve_headers_footers"].machine_rule.value is True
    assert all(r.scope == "global" for r in output.global_rules.instructions)
    assert all(r.instruction_id.startswith("gr_") for r in output.global_rules.instructions)

    # Icon library is deduplicated and keyed by content hash
    assert output.icon_library, "expected icons to be registered"
    keys = [entry.icon_key for entry in output.icon_library]
    assert len(keys) == len(set(keys))
    assert all(entry.icon_key.startswith("icon_") for entry in output.icon_library)
    assert all(entry.occurrences >= 1 for entry in output.icon_library)

    # The repeated icon in the layout container is one library entry, used twice
    reused = [entry for entry in output.icon_library if entry.occurrences > 1]
    assert reused, "the repeated container icon should be deduplicated"

    # Every icon reference resolves into the library
    referenced = {
        ref.icon_key
        for section in output.sections
        for element in section.skeleton_elements
        for ref in element.icons
    } | {
        ref.icon_key
        for section in output.sections
        for instruction in section.authoring_instructions
        for ref in instruction.icons
    }
    assert referenced <= set(keys)

    # Shaded row promoted to a callout, with colours read from the template
    callouts = [
        element
        for section in output.sections
        for element in section.skeleton_elements
        if element.element_type == "callout"
    ]
    assert len(callouts) == 1
    assert callouts[0].shading_hex == "E2EFDA"
    assert callouts[0].callout_type == "explanation"

    assert len(output.callout_styles) == 1
    style = output.callout_styles[0]
    assert style.callout_type == "explanation"
    assert style.background_color_hex == "#E2EFDA"
    assert style.left_border_color_hex == "#385723"
    assert style.icon_key in set(keys)
    assert "Explanation" in (style.trigger_instruction or "")
    assert output.totals.callouts == 1

    # Unshaded layout container stays unwrapped — no table survives in PURPOSE
    purpose = next(s for s in output.sections if s.title == "1 PURPOSE")
    assert purpose.section_number == "1"
    assert [e.element_type for e in purpose.skeleton_elements].count("table") == 0
    assert len(purpose.authoring_instructions) == 2
    assert all(i.scope == "section" for i in purpose.authoring_instructions)
    assert "[Role 1]" in purpose.placeholders

    # Serialised output is portable and free of the old sentinel
    clean = output.to_clean_dict()
    assert set(clean) >= {
        "version",
        "icon_library",
        "callout_styles",
        "global_rules",
        "sections",
        "totals",
    }
    blob = str(clean)
    assert "STYLE_INSTRUCTION" not in blob
    assert str(tmp_path) not in blob
    for entry in clean["icon_library"]:
        assert not Path(entry["asset_path"]).is_absolute()
        assert ":" not in entry["asset_path"]


@pytest.mark.asyncio
async def test_legend_table_keeps_icons_inside_cells_only(tmp_path, settings):
    """A header plus shaded rows stays a table, and each icon is drawn once.

    Copying cell icons onto the table element makes the UI render them again
    above the grid, with no description.
    """
    icon = _make_png(tmp_path / "legend.png", (30, 90, 200))
    doc = Document()
    doc.add_heading("4 DEFINITIONS", level=1)

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Infographics"
    table.cell(0, 1).text = "Description"
    _put_icon(table.cell(1, 0), icon)
    table.cell(1, 1).text = "Explanation – additional information to the topic."
    for cell in table.rows[1].cells:
        _shade_cell(cell, "E2EFDA")

    path = tmp_path / "legend.docx"
    doc.save(str(path))

    output = await TemplateExtractionService(settings=settings).extract(
        "legend", path, template_name="Legend"
    )
    section = output.sections[0]
    tables = [el for el in section.skeleton_elements if el.element_type == "table"]
    assert len(tables) == 1
    assert tables[0].icons == []
    icon_cells = [cell for cell in tables[0].cells if cell.icon_key]
    assert len(icon_cells) == 1
    assert icon_cells[0].text == ""
    described = next(cell for cell in tables[0].cells if "Explanation" in cell.text)
    assert described.shading_hex == "E2EFDA"


@pytest.mark.asyncio
async def test_icon_keys_are_stable_across_extractions(tmp_path, settings):
    """Content-hash keys must not drift between runs of the same template."""
    icon_a = _make_png(tmp_path / "icon_a.png", (200, 40, 40))
    icon_b = _make_png(tmp_path / "icon_b.png", (40, 160, 60))
    template_path = _build_template_docx(tmp_path / "stable.docx", icon_a, icon_b)

    service = TemplateExtractionService(settings=settings)
    first = await service.extract("stable_1", template_path, template_name="Stable")
    second = await service.extract("stable_2", template_path, template_name="Stable")

    assert [e.icon_key for e in first.icon_library] == [
        e.icon_key for e in second.icon_library
    ]
    assert [e.occurrences for e in first.icon_library] == [
        e.occurrences for e in second.icon_library
    ]
