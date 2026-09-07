"""Unit tests for DocxStyler and CalloutBuilder."""

import pytest
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

from app.schemas.migration import MigrationElement
from app.services.migration.docx_styler import DocxStyler
from app.services.migration.callout_builder import CalloutBuilder
from app.services.migration.schemas import CalloutStyleDef


TINY_PNG_BYTES = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4"
    b"\x00\x00\x00\rIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfe\xa74;\xa8\x00\x00\x00\x00IEND\xaeB`\x82"
)


def test_docx_styler_headings_and_paragraphs(tmp_path):
    styler = DocxStyler()
    doc = Document()

    h1 = styler.insert_heading(doc, "1 PURPOSE", level=1, font_family="Arial", font_size_pt=14.0)
    assert h1.text == "1 PURPOSE"
    assert h1.style.name.startswith("Heading")

    p1 = styler.insert_paragraph(doc, "This is a body paragraph.", font_family="Arial", font_size_pt=10.0)
    assert p1.text == "This is a body paragraph."

    out_file = tmp_path / "styled.docx"
    doc.save(str(out_file))
    assert out_file.exists()


def test_insert_list_ordered_formatting(tmp_path):
    styler = DocxStyler()
    doc = Document()

    # Provide raw numbered items with existing prefix numbers
    raw_items = [
        "1. First procedure step",
        "2. Second procedure step",
        "3. Final procedure step",
    ]
    paras = styler.insert_list(doc, raw_items, font_family="Arial")

    assert len(paras) == 3
    # Check that double numbers are stripped and proper hanging indent is set
    assert paras[0].text == "1.\tFirst procedure step"
    assert paras[1].text == "2.\tSecond procedure step"
    assert paras[2].text == "3.\tFinal procedure step"

    # Hanging indent verification
    assert paras[0].paragraph_format.left_indent == Inches(0.25)
    assert paras[0].paragraph_format.first_line_indent == Inches(-0.25)

    # Spacing verification: tight intermediate spacing, wider trailing spacing
    assert paras[0].paragraph_format.space_after == Pt(2.0)
    assert paras[1].paragraph_format.space_after == Pt(2.0)
    assert paras[2].paragraph_format.space_after == Pt(6.0)

    out_file = tmp_path / "ordered_list.docx"
    doc.save(str(out_file))
    assert out_file.exists()


def test_insert_list_unordered_formatting(tmp_path):
    styler = DocxStyler()
    doc = Document()

    # Provide raw bullet/dash items
    raw_items = [
        "• Safety goggles required",
        "- Heat resistant gloves",
        "* Full lab coat",
    ]
    paras = styler.insert_list(doc, raw_items, font_family="Arial")

    assert len(paras) == 3
    # Check that double bullets are stripped and uniform "•\t" is applied
    assert paras[0].text == "•\tSafety goggles required"
    assert paras[1].text == "•\tHeat resistant gloves"
    assert paras[2].text == "•\tFull lab coat"

    # Hanging indent verification
    assert paras[0].paragraph_format.left_indent == Inches(0.25)
    assert paras[0].paragraph_format.first_line_indent == Inches(-0.25)


def test_vertical_centering_with_icon(tmp_path):
    icon_path = tmp_path / "test_icon.png"
    icon_path.write_bytes(TINY_PNG_BYTES)

    styler = DocxStyler()
    doc = Document()

    # 1. Heading with icon
    h = styler.insert_heading(doc, "Safety Rules", level=2, icon_paths=[icon_path])
    pPr_h = h._element.get_or_add_pPr()
    text_align_h = pPr_h.find(qn("w:textAlignment"))
    assert text_align_h is not None
    assert text_align_h.get(qn("w:val")) == "center"

    # 2. Paragraph with icon
    p = styler.insert_paragraph(doc, "Attention required.", icon_paths=[icon_path])
    pPr_p = p._element.get_or_add_pPr()
    text_align_p = pPr_p.find(qn("w:textAlignment"))
    assert text_align_p is not None
    assert text_align_p.get(qn("w:val")) == "center"

    # 3. List with icon on first item
    paras = styler.insert_list(doc, ["Wear gloves", "Wash hands"], icon_paths=[icon_path])
    pPr_l = paras[0]._element.get_or_add_pPr()
    text_align_l = pPr_l.find(qn("w:textAlignment"))
    assert text_align_l is not None
    assert text_align_l.get(qn("w:val")) == "center"


def test_callout_builder_shading_and_borders(tmp_path):
    icon_path = tmp_path / "callout_icon.png"
    icon_path.write_bytes(TINY_PNG_BYTES)

    builder = CalloutBuilder()
    doc = Document()

    elem = MigrationElement(
        element_type="paragraph",
        text="Executive Summary: Learn what to expect from this section.",
        icons=[{"icon_id": "ico_1", "path": str(icon_path)}],
    )
    style = CalloutStyleDef(
        callout_type="executive_summary",
        display_name="Executive Summary",
        background_color_hex="#D9E1F2",
        left_border_color_hex="#2F5597",
    )

    tbl = builder.build(doc, elem, style=style, font_family="Arial")
    assert len(tbl.rows) == 1
    assert len(tbl.columns) == 1
    assert "Executive Summary" in tbl.cell(0, 0).text

    # Cell vertical alignment verification
    from docx.enum.table import WD_ALIGN_VERTICAL
    assert tbl.cell(0, 0).vertical_alignment == WD_ALIGN_VERTICAL.CENTER

    # Paragraph text alignment verification
    pPr = tbl.cell(0, 0).paragraphs[0]._element.get_or_add_pPr()
    text_align = pPr.find(qn("w:textAlignment"))
    assert text_align is not None
    assert text_align.get(qn("w:val")) == "center"

    out_file = tmp_path / "callout.docx"
    doc.save(str(out_file))
    assert out_file.exists()
