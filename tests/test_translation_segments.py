"""Phase T4: the English final document cut into segments (claims, table cells and headers, cross-references and the
template's own text), the renderer's body hook, the paragraph encoding with formatting tags, and block paths."""

from __future__ import annotations

import io

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.schemas.v2 import ResolvedRef, SCOPE_TEMPLATE, SCOPE_XREFS, SegmentKind, SegmentSet
from app.services.migration_v2.export import render_final
from app.services.migration_v2.render.xml import W_P, W_SDT, sdt_tag
from app.services.translation.segments import NUMBER_PLACEHOLDER, block_path, build_segments, find_block, xref_source
from app.services.translation.template_text import encode_paragraph, harvest_hook, parse_tags, strip_tags
from tests.test_v2_api_ui import _accept_gap
from tests.test_v2_assembly import _export_env
from tests.test_v2_renderer import _texts, tpl  # noqa: F401  (fixture)


def _migration(tmp_path, tpl):  # noqa: F811
    """A completed migration (its required gap accepted as N/A)."""
    e = _export_env(tmp_path, tpl)
    return e, _accept_gap(e)


def _final(migration, hook=None):
    e, job = migration
    return render_final(job, e.store, e.writer, body_hook=hook)


def _segments(migration):
    collected: list = []
    final = _final(migration, harvest_hook(collected))
    return final, build_segments(final.template, final.drafts, final.source, final.assembly.document, collected,
                                 final.accepted)


# ── The renderer's body hook ──────────────────────────────────────────


def test_the_body_hook_sees_the_filled_slots_before_they_are_unwrapped_and_changes_nothing(tpl, tmp_path):  # noqa: F811
    seen = {}

    def hook(doc, generated):
        seen["calls"] = seen.get("calls", 0) + 1
        seen["controls"] = sum((sdt_tag(s) or "").startswith("CC_") for s in doc.element.body.iter(W_SDT))
        seen["generated"] = generated

    migration = _migration(tmp_path, tpl)
    with_hook = _final(migration, hook)
    without = _final(migration)
    assert seen["calls"] == 1 and seen["controls"] > 0
    claim_text = "Never close a deviation without QA approval."
    generated_text = {"".join(t.text or "" for t in p.iter(qn("w:t"))) for p in seen["generated"] if p.tag == W_P}
    assert claim_text in generated_text and "PURPOSE" not in generated_text
    assert _texts(Document(io.BytesIO(with_hook.data))) == _texts(Document(io.BytesIO(without.data)))
    assert with_hook.slot_plan is not None


# ── Segments of the fixture migration ─────────────────────────────────


def test_segments_cover_claims_tables_xrefs_and_template_text(tpl, tmp_path):  # noqa: F811
    final, segments = _segments(_migration(tmp_path, tpl))
    by_kind = {kind: [s for s in segments if s.kind == kind] for kind in SegmentKind}

    claims = {s.locator.claim_id: s for s in by_kind[SegmentKind.CLAIM]}
    rendered = [c for d in final.drafts for slot in d.slots if slot.slot_id not in final.accepted for c in slot.claims
                if not c.is_gap_marker and c.kind.value not in ("table_row", "figure")]
    assert set(claims) == {c.claim_id for c in rendered}
    assert claims["C-5-001"].source_text == "Record every deviation within 5 days ({{ref:SRC-4.2}})."
    assert claims["C-5-007"].source_text == "Assessment"  # a sub-heading
    assert "C-2-001" not in claims  # the accepted gap's slot gives nothing
    assert all(s.scope == s.locator.section_id for s in by_kind[SegmentKind.CLAIM])

    cells = {(s.locator.unit_id, s.locator.row, s.locator.col): s.source_text for s in by_kind[SegmentKind.TABLE_CELL]}
    assert cells[("SRC-3-U001", 1, 0)] == "CAPA" and cells[("SRC-3-U002", 2, 1)] == "Quality Assurance"
    headers = [(s.locator.unit_id, s.source_text) for s in by_kind[SegmentKind.TABLE_HEADER]]
    assert headers.count(("SRC-3-U001", "Term")) == 1 and not any(u == "SRC-3-U002" for u, _ in headers)  # once per table
    assert {s.scope for s in by_kind[SegmentKind.TABLE_CELL]} == {"TGT-3", "TGT-4"}

    refs = final.assembly.document.refs
    xrefs = by_kind[SegmentKind.XREF]
    assert [s.locator.ref_index for s in xrefs] == [i for i, r in enumerate(refs) if r.text.strip()]
    assert all(s.scope == SCOPE_XREFS and s.source_text == xref_source(refs[s.locator.ref_index]) for s in xrefs)

    template = [s.source_text for s in by_kind[SegmentKind.TEMPLATE_TEXT]]
    assert {"PURPOSE", "PROCESS", "OWNER", "Role", "Responsibility", "This SOP is applicable:"} <= set(template)
    assert not set(template) & {c.text for c in rendered}
    assert all(s.scope == SCOPE_TEMPLATE for s in by_kind[SegmentKind.TEMPLATE_TEXT])

    ids = [s.segment_id for s in segments]
    assert ids == [f"SEG-{n:04d}" for n in range(1, len(ids) + 1)]
    assert all(s.target_text is None for s in segments)
    assert SegmentSet(job_id="TRN-X", migration_job_id="MIG-X", language="de-DE", version=1, segments=segments).counts


def test_segments_are_the_same_on_a_second_run(tpl, tmp_path):  # noqa: F811
    migration = _migration(tmp_path, tpl)
    _, first = _segments(migration)
    _, second = _segments(migration)
    assert [s.model_dump() for s in first] == [s.model_dump() for s in second]


def test_template_text_paths_point_at_their_paragraphs(tpl, tmp_path):  # noqa: F811
    seen = []

    def hook(doc, generated):
        collected: list = []
        harvest_hook(collected)(doc, generated)
        body = doc.element.body
        seen.extend((path, strip_tags(text), find_block(body, path)) for path, text in collected)

    _final(_migration(tmp_path, tpl), hook)
    assert seen
    for path, text, p in seen:
        assert p is not None and p.tag == W_P
        assert "".join(t.text or "" for t in p.iter(qn("w:t"))) == text


def test_xref_number_becomes_a_placeholder():
    ref = ResolvedRef(claim_id="C", target="SRC-7", source_phrase="Chapter 7", text="see Chapter 9, no. 1",
                      number="9", number_at=12)
    assert xref_source(ref) == f"see Chapter {NUMBER_PLACEHOLDER}, no. 1"
    assert xref_source(ref.model_copy(update={"number": None, "number_at": None})) == "see Chapter 9, no. 1"
    assert xref_source(ref.model_copy(update={"number_at": 3})) == "see Chapter 9, no. 1"  # does not match: plain


# ── Template text: what is collected ──────────────────────────────────


def _body(xml: str):
    doc = Document()
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for el in parse_xml(f"<w:body {nsdecls('w')}>{xml}</w:body>"):
        body.find(qn("w:sectPr")).addprevious(el)
    return doc


def _p(text: str, extra: str = "") -> str:
    return f"<w:p>{extra}<w:r><w:t xml:space='preserve'>{text}</w:t></w:r></w:p>"


def test_harvest_leaves_out_the_cover_the_toc_entries_built_content_and_text_without_letters():
    doc = _body(
        _p("Cover title") + "<w:p><w:r><w:br w:type='page'/></w:r></w:p>"
        + _p("Table of Contents")
        + "<w:p><w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText>TOC \\o \"1-3\"</w:instrText></w:r>"
          "<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>1 Purpose 3</w:t></w:r></w:p>"
        + "<w:p><w:r><w:t>2 Scope 4</w:t></w:r><w:r><w:fldChar w:fldCharType='end'/></w:r></w:p>"
        + _p("Purpose") + _p("12.5 %")
        + "<w:tbl><w:tr><w:tc>" + _p("Attention") + "</w:tc><w:tc>" + _p("Built from a claim") + "</w:tc></w:tr></w:tbl>"
    )
    body = doc.element.body
    built = next(p for p in body.iter(W_P) if "Built" in "".join(t.text or "" for t in p.iter(qn("w:t"))))
    collected: list = []
    harvest_hook(collected)(doc, frozenset(built.iter()))
    assert [text for _, text in collected] == ["Table of Contents", "Purpose", "Attention"]


# ── Paragraph encoding ────────────────────────────────────────────────


def _para(xml: str):
    return parse_xml(f"<w:p {nsdecls('w', 'r')}>{xml}</w:p>")


def _r(text: str, rpr: str = "") -> str:
    return f"<w:r>{'<w:rPr>' + rpr + '</w:rPr>' if rpr else ''}<w:t xml:space='preserve'>{text}</w:t></w:r>"


def test_a_paragraph_with_one_formatting_has_no_tags():
    p = _para(_r("Record the ", "<w:lang w:val='en-GB'/>") + _r("deviation.", "<w:lang w:val='en-US'/>"))
    encoded = encode_paragraph(p)
    assert encoded.tagged == encoded.text == "Record the deviation." and not encoded.groups


def test_other_formatting_is_tagged_and_objects_are_standalone():
    p = _para(_r("Record the ") + _r("batch number", "<w:b/>") + _r(" in ") + "<w:r><w:tab/></w:r>"
              + _r("the log, ") + _r("always", "<w:b/>") + _r(".")
              + "<w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText>PAGE</w:instrText></w:r>"
                "<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>3</w:t></w:r>"
                "<w:r><w:fldChar w:fldCharType='end'/></w:r>"
              + "<w:hyperlink r:id='rId9'>" + _r("www.example.com") + "</w:hyperlink>")
    encoded = encode_paragraph(p)
    assert encoded.tagged == "Record the ⟦1⟧batch number⟦/1⟧ in ⟦2⟧the log, ⟦1⟧always⟦/1⟧.⟦3⟧⟦4⟧"
    assert encoded.text == "Record the batch number in the log, always."
    assert set(encoded.groups) == {1} and set(encoded.objects) == {2, 3, 4}
    assert len(encoded.objects[3].elements) == 5 and encoded.objects[3].whole  # the whole PAGE field
    assert parse_tags(encoded.tagged) == [("open", 1), ("close", 1), ("object", 2), ("open", 1), ("close", 1),
                                          ("object", 3), ("object", 4)]
    assert strip_tags(encoded.tagged) == encoded.text


def test_the_most_used_formatting_is_the_plain_one():
    encoded = encode_paragraph(_para(_r("Note:", "<w:b/>") + _r(" a long explanation follows here")))
    assert encoded.tagged == "⟦1⟧Note:⟦/1⟧ a long explanation follows here"


def test_a_space_in_other_formatting_opens_no_tag():
    encoded = encode_paragraph(_para(_r("Table of") + _r(" ", "<w:i/>") + _r("Contents")))
    assert encoded.tagged == "Table of Contents" and not encoded.groups


# ── Block paths ───────────────────────────────────────────────────────


def test_block_paths_round_trip_for_body_cell_control_and_text_box_paragraphs():
    doc = _body(
        _p("one") + _p("two")
        + "<w:tbl><w:tr><w:tc>" + _p("cell a") + "</w:tc><w:tc>" + _p("cell b") + _p("cell c") + "</w:tc></w:tr></w:tbl>"
        + "<w:sdt><w:sdtPr><w:tag w:val='CC_X'/></w:sdtPr><w:sdtContent>" + _p("in control") + "</w:sdtContent></w:sdt>"
        + "<w:p><w:r><w:pict><v:shape xmlns:v='urn:schemas-microsoft-com:vml'><v:textbox><w:txbxContent>"
        + _p("in text box") + "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )
    body = doc.element.body
    paths = {}
    for p in body.iter(W_P):
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        if text:
            paths[text] = block_path(p, body)
            assert find_block(body, paths[text]) is p
    assert paths["two"] == "body/p[1]"
    assert paths["cell c"] == "body/tbl[0]/tr[0]/tc[1]/p[1]"
    assert paths["in control"] == "body/sdt[0]/sdtContent[0]/p[0]"
    assert paths["in text box"].startswith("body/p[2]/r[0]/pict[0]/v:shape[0]/v:textbox[0]/txbxContent[0]/p[0]")
    assert find_block(body, "body/p[99]") is None and find_block(body, "nope") is None
