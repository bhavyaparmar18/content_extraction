"""DOCX highlight, numbering, and shaded-callout extraction."""

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image

from app.config.settings import Settings
from app.schemas.document import ElementType
from app.services.export.migration_exporter import MigrationExporter
from app.services.extraction.icons import IconExtractor
from app.services.hierarchy.ast_builder import ASTBuilder
from app.services.parser.docx_parser import DocxParser


@pytest.fixture
def settings(tmp_path):
    configured = Settings(project_root=tmp_path)
    configured.resolve_paths(tmp_path)
    configured.ensure_directories()
    return configured


def _save(doc: Document, path: Path) -> Path:
    doc.save(str(path))
    return path


def _parse(settings, path: Path):
    return DocxParser(settings=settings).parse(str(path), document_id="doc")


def _elements(raw):
    return [element for page in raw.pages for element in page.elements]


def _export(settings, path: Path) -> dict:
    raw = _parse(settings, path)
    raw = IconExtractor(settings=settings).extract(raw, document_id="doc")
    ast = ASTBuilder(settings=settings).build(raw)
    return MigrationExporter.export(document_id="doc", ast=ast).to_clean_dict()


def _shade(props_owner, fill: str) -> None:
    props_owner.append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def _mark_highlight(paragraph, color: str) -> None:
    """Highlight the paragraph mark only. That does not colour the text."""
    mark = OxmlElement("w:rPr")
    highlight = OxmlElement("w:highlight")
    highlight.set(qn("w:val"), color)
    mark.append(highlight)
    paragraph._element.get_or_add_pPr().append(mark)


def test_paragraph_mark_highlight_is_ignored(settings, tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("Learning Expert")
    _mark_highlight(paragraph, "yellow")
    raw = _parse(settings, _save(doc, tmp_path / "mark.docx"))

    element = _elements(raw)[0]
    assert element.content == "Learning Expert"
    assert element.highlight_spans == []
    assert element.highlight_color is None


def test_run_highlight_offsets(settings, tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("before ")
    highlighted = paragraph.add_run("highlighted")
    highlighted.font.highlight_color = WD_COLOR_INDEX.YELLOW
    paragraph.add_run(" after")
    # A second adjacent yellow run must merge into the same span.
    paragraph = doc.add_paragraph()
    paragraph.add_run("keep ")
    left = paragraph.add_run("this ")
    right = paragraph.add_run("phrase")
    left.font.highlight_color = WD_COLOR_INDEX.YELLOW
    right.font.highlight_color = WD_COLOR_INDEX.YELLOW

    raw = _parse(settings, _save(doc, tmp_path / "runs.docx"))
    first, second = _elements(raw)

    assert len(first.highlight_spans) == 1
    span = first.highlight_spans[0]
    assert first.content[span.start_offset:span.end_offset] == "highlighted"
    assert span.text == "highlighted"
    assert span.color_name == "yellow"
    assert span.color_hex == "FFFF00"
    assert first.highlight_color is None  # the paragraph is not entirely highlighted

    assert len(second.highlight_spans) == 1
    merged = second.highlight_spans[0]
    assert second.content[merged.start_offset:merged.end_offset] == "this phrase"


def test_white_shading_is_not_a_background(settings, tmp_path):
    doc = Document()
    white = doc.add_paragraph("white paragraph")
    _shade(white._element.get_or_add_pPr(), "FFFFFF")
    white_run = white.add_run(" still white")
    _shade(white_run._element.get_or_add_rPr(), "FFFFFF")

    shaded = doc.add_paragraph("Shaded paragraph")
    _shade(shaded._element.get_or_add_pPr(), "DBE5F1")

    table = doc.add_table(rows=1, cols=2)
    _shade(table.cell(0, 0)._element.get_or_add_tcPr(), "FFFFFF")
    table.cell(0, 0).text = "no fill"
    _shade(table.cell(0, 1)._element.get_or_add_tcPr(), "F7CAAC")
    table.cell(0, 1).text = "peach"

    raw = _parse(settings, _save(doc, tmp_path / "shading.docx"))
    paragraphs = [el for el in _elements(raw) if el.element_type == ElementType.PARAGRAPH]
    white_el, shaded_el = paragraphs[0], paragraphs[1]
    assert white_el.shading_hex is None
    assert white_el.highlight_spans == []
    assert shaded_el.shading_hex == "DBE5F1"

    tables = [el for el in _elements(raw) if el.element_type == ElementType.TABLE]
    header = tables[0].headers
    assert header[0].shading_hex is None
    assert header[1].shading_hex == "F7CAAC"

    exported = _export(settings, tmp_path / "shading.docx")
    exported_paragraphs = [
        el for sec in exported["sections"] for el in sec["elements"]
        if el["element_type"] == "paragraph"
    ]
    assert "background_color" not in exported_paragraphs[0]
    assert exported_paragraphs[1]["background_color"] == "#DBE5F1"


def _number(paragraph, num_id: str) -> None:
    """Attach direct ``w:numPr``. The style stays Normal; the default template
    already defines these numIds (1 bullet, 5 decimal)."""
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), num_id)
    num_pr.append(num)
    paragraph._element.get_or_add_pPr().append(num_pr)


def test_numbering_overrides_style_name(settings, tmp_path):
    doc = Document()
    plain_bullet = doc.add_paragraph("Plain bullet")
    plain_number = doc.add_paragraph("Plain number")
    assert plain_bullet.style.name == "Normal"
    _number(plain_bullet, "1")
    _number(plain_number, "5")

    raw = _parse(settings, _save(doc, tmp_path / "numbering.docx"))
    by_text = {el.content: el.element_type for el in _elements(raw)}
    assert by_text["Plain bullet"] == ElementType.LIST_ITEM
    assert by_text["Plain number"] == ElementType.NUMBERED_STEP


def test_content_controls_supply_table_cell_text(settings, tmp_path):
    """Preamble values live in w:sdt, which python-docx cell.text does not read."""
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Title"
    table.cell(1, 0).text = "Version"
    # Cell-level content control: w:tc/w:sdt/w:sdtContent/w:p
    value = table.cell(0, 1)._tc
    for child in list(value):
        if child.tag == qn("w:p"):
            value.remove(child)
    value.append(parse_xml(
        f'<w:sdt {nsdecls("w")}>'
        f'<w:sdtContent><w:p><w:r><w:t>Learning Management</w:t></w:r></w:p></w:sdtContent>'
        f'</w:sdt>'
    ))
    # Paragraph-level controls mixed with a separator run: "1" + "." + "0"
    version = table.cell(1, 1).paragraphs[0]._element
    for child in list(version):
        if child.tag != qn("w:pPr"):
            version.remove(child)
    version.append(parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:r><w:t>1</w:t></w:r></w:sdtContent></w:sdt>'
    ))
    version.append(parse_xml(f'<w:r {nsdecls("w")}><w:t>.</w:t></w:r>'))
    version.append(parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:r><w:t>0</w:t></w:r></w:sdtContent></w:sdt>'
    ))

    path = _save(doc, tmp_path / "sdt.docx")
    from app.services.extraction.metadata_extractor import SOPMetadataExtractor

    title, _name, _number, version, _doc_type = SOPMetadataExtractor._extract_from_docx(path)
    assert title == "Learning Management"
    assert version == "1.0"

    raw = _parse(settings, path)
    tables = [el for el in _elements(raw) if el.element_type == ElementType.TABLE]
    rows = [tables[0].headers] + tables[0].rows
    texts = [[c.content_text for c in row] for row in rows]
    assert texts[0] == ["Title", "Learning Management"]
    assert texts[1] == ["Version", "1.0"]


def _png(path: Path, size: int = 16) -> Path:
    Image.new("RGB", (size, size), "red").save(path)
    return path


def test_shaded_icon_list_callout_is_one_element(settings, tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    icon_cell, text_cell = table.cell(0, 0), table.cell(0, 1)
    _shade(icon_cell._element.get_or_add_tcPr(), "F7CAAC")
    _shade(text_cell._element.get_or_add_tcPr(), "F7CAAC")
    icon_cell.paragraphs[0].add_run().add_picture(str(_png(tmp_path / "icon.png")))

    text_cell.paragraphs[0].add_run("Required Learnings for new hires. Supervisors consider these points:")
    text_cell.add_paragraph("Assign basic compliance trainings immediately.", style="List Bullet")
    text_cell.add_paragraph("Assign all other Learnings in a staggered way.", style="List Bullet")
    closing = text_cell.add_paragraph()
    closing.add_run("Employees must be fully qualified before performing their work.").bold = True

    exported = _export(settings, _save(doc, tmp_path / "callout.docx"))
    elements = [el for sec in exported["sections"] for el in sec["elements"]]
    callouts = [el for el in elements if el["element_type"] == "callout"]
    assert len(callouts) == 1
    callout = callouts[0]
    assert callout["background_color"] == "#F7CAAC"
    assert len(callout["icons"]) == 1
    assert [block["type"] for block in callout["content"]] == ["paragraph", "list", "paragraph"]
    assert callout["content"][1]["items"] == [
        "Assign basic compliance trainings immediately.",
        "Assign all other Learnings in a staggered way.",
    ]
    assert callout["content"][2]["bold"] is True
    assert "fully qualified" in callout["content"][2]["text"]
    assert not any(el["element_type"] == "table" for el in elements)


def test_icon_inside_list_paragraph_stays_one_list(settings, tmp_path):
    doc = Document()
    first = doc.add_paragraph("Assign basic compliance trainings immediately.", style="List Bullet")
    first.add_run().add_picture(str(_png(tmp_path / "bullet-icon.png")))
    doc.add_paragraph("Assign all other Learnings in a staggered way.", style="List Bullet")

    exported = _export(settings, _save(doc, tmp_path / "split-list.docx"))
    elements = [el for sec in exported["sections"] for el in sec["elements"]]
    lists = [el for el in elements if el["element_type"] == "list"]
    assert len(lists) == 1
    assert lists[0]["items"] == [
        "Assign basic compliance trainings immediately.",
        "Assign all other Learnings in a staggered way.",
    ]
    assert len(lists[0]["icons"]) == 1
    assert not any(el["element_type"] in ("image", "paragraph") for el in elements)


def _wrap_square(paragraph, width_in: float, height_in: float, align: str = "left") -> None:
    """Turn the paragraph's inline picture into a side-wrapping float."""
    drawing = next(paragraph._element.iter(qn("w:drawing")))
    inline = next(drawing.iter(qn("wp:inline")))
    extent = inline.find(qn("wp:extent"))
    emu = 914400
    extent.set("cx", str(int(width_in * emu)))
    extent.set("cy", str(int(height_in * emu)))
    anchor = OxmlElement("wp:anchor")
    anchor.set("behindDoc", "0")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "0")
    anchor.set("distR", "0")
    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "margin")
    align_el = OxmlElement("wp:align")
    align_el.text = align
    position_h.append(align_el)
    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "paragraph")
    offset = OxmlElement("wp:posOffset")
    offset.text = "0"
    position_v.append(offset)
    wrap = OxmlElement("wp:wrapSquare")
    moved = list(inline)
    for child in moved:
        inline.remove(child)
    anchor.extend([position_h, position_v, extent, wrap])
    for child in moved:
        if child is extent:
            continue
        anchor.append(child)
    drawing.replace(inline, anchor)


def test_side_wrap_groups_the_paragraphs_beside_the_picture(settings, tmp_path):
    doc = Document()
    lead = doc.add_paragraph("Evaluation of a need of an Effectiveness Check is performed during review.")
    lead.add_run().add_picture(str(_png(tmp_path / "figure.png", size=64)))
    _wrap_square(lead, width_in=2.4, height_in=2.2)
    doc.add_paragraph("The Learning Owner must review executed Effectiveness Checks.")
    doc.add_heading("Maintenance of the learning", level=2)
    doc.add_paragraph("This paragraph starts below the figure.")

    exported = _export(settings, _save(doc, tmp_path / "float.docx"))
    elements = [el for sec in exported["sections"] for el in sec["elements"]]
    floats = [el for el in elements if el["element_type"] == "float"]
    assert len(floats) == 1
    floated = floats[0]
    assert floated["image_align"] == "left"
    assert floated["wrap"] == "square"
    assert [block["text"] for block in floated["content"]] == [
        "Evaluation of a need of an Effectiveness Check is performed during review.",
        "The Learning Owner must review executed Effectiveness Checks.",
    ]
    assert "This paragraph starts below the figure." in [el.get("text") for el in elements]
    assert not any(
        el["element_type"] == "image" and "figure" in (el.get("image_path") or "")
        for el in elements
    )


def test_wide_picture_stays_a_separate_image(settings, tmp_path):
    doc = Document()
    lead = doc.add_paragraph("A caption line above a full-width diagram.")
    lead.add_run().add_picture(str(_png(tmp_path / "wide.png", size=160)))
    _wrap_square(lead, width_in=6.2, height_in=3.0)
    doc.add_paragraph("This stays under the diagram.")

    exported = _export(settings, _save(doc, tmp_path / "wide.docx"))
    elements = [el for sec in exported["sections"] for el in sec["elements"]]
    assert not any(el["element_type"] == "float" for el in elements)
    assert any(el["element_type"] == "image" for el in elements)
    assert any(el.get("text") == "This stays under the diagram." for el in elements)
