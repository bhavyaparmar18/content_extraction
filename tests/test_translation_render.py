"""Phase T6: the translated final document. The template-text decoder, the translated drafts, source and
cross-references (with their English fallbacks), the RENDERING stage with and without a language template, the
Word language and East Asian font, and the preview API."""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.schemas.v2 import (
    ArtifactKind,
    AssembledDocument,
    ClaimKind,
    RenderReport,
    ResolvedRef,
    Segment,
    SegmentKind,
    SegmentLocator,
    SegmentOrigin,
    SegmentSet,
    TranslationStatus,
)
from app.services.migration_v2.artifacts import ArtifactWriter
from app.services.translation.checks import check_segment
from app.services.translation.front_matter import extract_front_matter
from app.services.translation.render import (
    TRANSLATED,
    Fallbacks,
    east_asia_font,
    set_east_asia_font,
    set_language,
    translated_assembly,
    translated_drafts,
)
from app.services.translation.segments import build_segments
from app.services.translation.template_text import decode_into, encode_paragraph
from tests.test_translation_front_matter import language_template
from tests.test_translation_jobs import _env, _fake_translator, _wait_status, api, translating  # noqa: F401
from tests.test_v2_drafter import FakeChain, FakeFactory
from tests.test_v2_renderer import _STALE_TOC, tpl  # noqa: F401  (fixture)

W = nsdecls("w")


# ── The template-text decoder ─────────────────────────────────────────


def _mixed_paragraph():
    """'Press ' Save(bold) tab ' and see ' link(here) ' now' field(PAGE), with a bookmark around it all."""
    return parse_xml(
        f'<w:p {W} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<w:pPr><w:jc w:val="left"/></w:pPr><w:bookmarkStart w:id="7" w:name="label"/>'
        '<w:r><w:t xml:space="preserve">Press </w:t></w:r>'
        '<w:r><w:rPr><w:b/></w:rPr><w:t>Save</w:t></w:r>'
        '<w:r><w:tab/></w:r>'
        '<w:r><w:t xml:space="preserve"> and see </w:t></w:r>'
        '<w:hyperlink r:id="rId9"><w:r><w:t>here</w:t></w:r></w:hyperlink>'
        '<w:proofErr w:type="spellStart"/><w:r><w:t xml:space="preserve"> now and then</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>3</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        '<w:bookmarkEnd w:id="7"/></w:p>')


def _runs(p) -> list[tuple[str, bool]]:
    return [("".join(t.text or "" for t in r.iter(qn("w:t"))), r.find(f"{qn('w:rPr')}/{qn('w:b')}") is not None)
            for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None]


def test_decode_into_rebuilds_the_runs_with_their_formatting_objects_and_bookmarks():
    p = _mixed_paragraph()
    encoded = encode_paragraph(p)
    assert encoded.tagged == "Press ⟦1⟧Save⟦/1⟧⟦2⟧ and see ⟦3⟧ now and then⟦4⟧"
    # The translation moves the bold word and the link.
    assert decode_into(p, encoded, "Voir ⟦3⟧ et appuyer sur ⟦1⟧Enregistrer⟦/1⟧⟦2⟧ maintenant⟦4⟧")
    assert encode_paragraph(p).tagged == "Voir ⟦1⟧ et appuyer sur ⟦2⟧Enregistrer⟦/2⟧⟦3⟧ maintenant⟦4⟧"
    assert ("Enregistrer", True) in _runs(p) and ("Voir ", False) in _runs(p)
    children = [c.tag.split("}")[1] for c in p]
    assert children[:2] == ["pPr", "bookmarkStart"] and children[-1] == "bookmarkEnd"
    assert "proofErr" not in children and p.find(qn("w:hyperlink")).get(qn("r:id")) == "rId9"
    assert len(list(p.iter(qn("w:tab")))) == 1 and "PAGE" in "".join(i.text for i in p.iter(qn("w:instrText")))
    assert p.find(qn("w:pPr")).find(qn("w:jc")) is not None


@pytest.mark.parametrize("target", [
    "Appuyer ⟦1⟧Enregistrer⟦/1⟧⟦2⟧ voir ⟦3⟧",           # the field ⟦4⟧ is lost
    "Appuyer ⟦1⟧Enregistrer⟦/1⟧⟦2⟧ voir ⟦3⟧⟦3⟧⟦4⟧",     # the link doubled
    "Appuyer ⟦9⟧Enregistrer⟦/9⟧⟦2⟧ voir ⟦3⟧⟦4⟧",        # an unknown marker
])
def test_decode_into_leaves_the_paragraph_when_the_markers_do_not_fit(target):
    p = _mixed_paragraph()
    before = encode_paragraph(p).tagged
    assert not decode_into(p, encode_paragraph(p), target)
    assert encode_paragraph(p).tagged == before


# ── Translated inputs and their fallbacks ─────────────────────────────


def _segment(sid: str, kind: SegmentKind, source: str, target, **locator) -> Segment:
    return Segment(segment_id=sid, kind=kind, source_text=source, target_text=target,
                   origin=SegmentOrigin.LLM if target is not None else None, locator=SegmentLocator(**locator))


def test_a_reference_keeps_its_number_or_falls_back_to_english():
    refs = [ResolvedRef(claim_id=f"C{i}", target="SRC-4", source_phrase="section 4", text="see section 4.2",
                        number="4.2", number_at=12, bookmark="_Ref1") for i in range(3)]
    assembled = AssembledDocument(job_id="MIG-X", version=1, refs=refs)
    fallbacks = Fallbacks()
    out = translated_assembly(assembled, [
        _segment("SEG-1", SegmentKind.XREF, "see section ⟦1⟧", "voir la section ⟦1⟧", claim_id="C0", ref_index=0),
        _segment("SEG-2", SegmentKind.XREF, "see section ⟦1⟧", "voir la section", claim_id="C1", ref_index=1),
        _segment("SEG-3", SegmentKind.XREF, "see section ⟦1⟧", None, claim_id="C2", ref_index=2),
    ], fallbacks)
    assert (out.refs[0].text, out.refs[0].number_at) == ("voir la section 4.2", 16)
    assert out.refs[0].text[16:19] == "4.2" and out.refs[0].bookmark == "_Ref1"
    assert [r.text for r in out.refs[1:]] == ["see section 4.2"] * 2 and assembled.refs[0].text == "see section 4.2"
    assert fallbacks.kept == ["SEG-2", "SEG-3"] and "number ⟦1⟧ is lost" in fallbacks.warnings[0]


def test_an_edited_table_row_is_one_claim_segment_and_keeps_its_cells(tpl, tmp_path):  # noqa: F811
    from tests.test_v2_assembly import scenario

    source, drafts, _ = scenario(tmp_path)
    units = {u.unit_id: u for u in source.iter_units()}
    claim = next(c for d in drafts for s in d.slots for c in s.claims if c.kind == ClaimKind.TABLE_ROW)
    unit = units[claim.source_unit_ids[0]]
    edited = " | ".join(f"edited {c.text}" for c in unit.table_ref.cells)
    claim.text = edited
    segments = build_segments(tpl, drafts, source, None)
    row = [s for s in segments if s.locator.claim_id == claim.claim_id]
    assert [(s.kind, s.source_text) for s in row] == [(SegmentKind.CLAIM, edited)]
    assert not [s for s in segments if s.kind == SegmentKind.TABLE_CELL and s.locator.unit_id == unit.unit_id]

    translated = " | ".join(f"modifié {c.text}" for c in unit.table_ref.cells)
    lost = translated.replace(" | ", " ", 1)
    check = lambda text: {i.issue_id.rsplit("-", 1)[-1] for i in check_segment(  # noqa: E731
        row[0].model_copy(update={"target_text": text, "origin": SegmentOrigin.LLM}), "fr-FR")}
    assert "CELLS" not in check(translated) and "CELLS" in check(lost)
    out = translated_drafts(drafts, [row[0].model_copy(update={"target_text": translated, "origin": SegmentOrigin.LLM})],
                            Fallbacks())
    assert next(c for d in out for s in d.slots for c in s.claims if c.claim_id == claim.claim_id).text == translated
    assert claim.text == edited  # the snapshot's drafts are not changed


# ── The RENDERING stage ───────────────────────────────────────────────

_COVER = f'<w:p {W}><w:r><w:t>English cover title</w:t></w:r></w:p>'
_COVER_END = f'<w:p {W}><w:pPr><w:sectPr><w:type w:val="nextPage"/></w:sectPr></w:pPr></w:p>'


def _dress(tpl):  # noqa: F811
    """The fixture template with an English first page (ended by a section break), a header, a table of contents
    before its first heading, and a template paragraph with a bold word, a tab and a run language after it."""
    doc = Document(tpl.normalized_file)
    body = doc.element.body
    first_heading = next(p for p in doc.paragraphs if p.style.name.startswith("Heading"))
    first_heading._p.addprevious(parse_xml(_STALE_TOC))
    body.insert(0, parse_xml(_COVER_END))
    body.insert(0, parse_xml(_COVER))
    doc.sections[-1].header.paragraphs[0].text = "EN header"
    first_heading._p.addnext(parse_xml(
        f'<w:p {W}><w:r><w:t xml:space="preserve">Read the </w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>notes</w:t></w:r>'
        f'<w:r><w:tab/></w:r><w:r><w:rPr><w:lang w:val="en-US"/></w:rPr><w:t xml:space="preserve"> first.</w:t></w:r></w:p>'))
    doc.save(tpl.normalized_file)
    return tpl


def _texts(doc) -> list[str]:
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]


def _toc_titles(doc) -> list[str]:
    from tests.test_v2_renderer import _toc

    return [text for _, text, _ in _toc(doc)]


def _translated(env, job):
    report = env.orch.artifacts.latest(job, ArtifactKind.RENDER_REPORT, RenderReport, TRANSLATED)
    ref = job.latest_artifact(ArtifactKind.DOCX, TRANSLATED)
    return report, Document(io.BytesIO(ArtifactWriter.read_bytes(ref)))


def _render_event(env, job_id) -> dict:
    return next(e["detail"] for e in reversed(env.store.list_events(job_id)) if e["event"] == "render")


@pytest.mark.asyncio
async def test_a_translated_job_is_rendered_in_its_language_with_the_english_front_matter(tpl, tmp_path):  # noqa: F811
    env = _env(tmp_path, _dress(tpl), FakeFactory(TranslatorOutput=FakeChain(translating())))
    job = await env.orch.wait(env.orch.create(env.migration.job_id, "fr-FR").job_id)
    assert job.status == TranslationStatus.REVIEW_PENDING, job.error
    report, doc = _translated(env, job)
    assert report.problems == [] and report.mode.value == "final"

    segment_set = env.orch.artifacts.latest(job, ArtifactKind.SEGMENTS, SegmentSet)
    texts = _texts(doc)
    joined = " ".join(texts)
    for segment in segment_set.segments:
        if segment.kind != SegmentKind.XREF:
            probe = segment.source_text.splitlines()[0].split("⟦")[0].split("{{")[0].strip()
            assert any(t.startswith("FR ") and probe in t for t in texts), segment.segment_id
    assert "FR see section" in joined
    event = _render_event(env, job.job_id)
    assert event["kept_english"] == 0 and event["front_matter"] is None
    assert sum(event["applied"].values()) == len(segment_set.segments)

    # The template paragraph keeps its bold word; the TOC lists the translated headings.
    label = next(p for p in doc.element.body.iter(qn("w:p")) if "notes" in "".join(_texts_of(p)))
    assert _runs(label) == [("FR Read the ", False), ("notes", True), (" first.", False)]
    assert label.find(f"{qn('w:r')}/{qn('w:tab')}") is not None
    titles = _toc_titles(doc)
    assert titles and all(t.startswith("FR ") for t in titles) and "FR PURPOSE" in titles
    assert "OLD TEMPLATE CHAPTER" not in joined

    # The English first page, header and footer stay; the body and the defaults are French.
    assert texts[0] == "English cover title"
    assert {p.text for s in doc.sections for p in s.header.paragraphs if p.text} == {"EN header"}
    default = doc.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:lang')}")
    assert default.get(qn("w:val")) == "fr-FR" and default.get(qn("w:eastAsia")) != "fr-FR"  # only zh/ja set it
    assert {n.get(qn("w:val")) for n in doc.element.body.iter(qn("w:lang"))} <= {"fr-FR"}

    preview = env.orch.artifacts.latest(job, ArtifactKind.RENDER_REPORT, RenderReport, TRANSLATED)
    assert preview.job_id == job.job_id and preview.version == 1


def _texts_of(p) -> list[str]:
    return [t.text or "" for t in p.iter(qn("w:t"))]


@pytest.mark.asyncio
async def test_a_language_template_replaces_the_first_page_headers_and_footers(tpl, tmp_path):  # noqa: F811
    env = _env(tmp_path, _dress(tpl), FakeFactory(TranslatorOutput=FakeChain(translating(
        lambda key, text, repair: f"DE {text}"))))
    package, preview = tmp_path / "lt_pkg.docx", tmp_path / "lt_preview.docx"
    report = extract_front_matter(language_template(tmp_path / "lt.docx"), package, preview, template_id="t-de",
                                  version=1, language="de-DE")
    assert report.ok, report.problems
    env.templates.create_record("t-de", "German", "de-DE", front_matter_path=str(package), status="ready")
    job = await env.orch.wait(env.orch.create(env.migration.job_id, "de-DE", language_template=("t-de", 1)).job_id)
    assert job.status == TranslationStatus.REVIEW_PENDING, job.error

    report, doc = _translated(env, job)
    assert report.problems == []  # front_matter_problems: the first page, headers and footers are the package's
    texts = _texts(doc)
    assert texts[0] == "Deckblatt Titel" and "English cover title" not in texts
    headers = {p.text for s in doc.sections for h in (s.header, s.first_page_header) for p in h.paragraphs if p.text}
    assert "LT Kopf Inhalt" in headers and "EN header" not in headers
    assert "DE PURPOSE" in _toc_titles(doc)
    event = _render_event(env, job.job_id)
    assert event["front_matter"]["blocks_inserted"] > 0 and event["kept_english"] == 0
    # The first page keeps the language template's own text and language; the body is German.
    assert {n.get(qn("w:val")) for n in doc.element.body.iter(qn("w:lang"))} <= {"de-DE"}


@pytest.mark.asyncio
async def test_japanese_gets_the_east_asian_language_and_font(tpl, tmp_path):  # noqa: F811
    env = _env(tmp_path, tpl, FakeFactory(TranslatorOutput=FakeChain(translating(
        lambda key, text, repair: f"日本語の文書 {text}"))))
    env.settings.translation_east_asia_fonts = {"ja-JP": "MS Gothic"}
    job = await env.orch.wait(env.orch.create(env.migration.job_id, "ja-JP").job_id)
    assert job.status == TranslationStatus.REVIEW_PENDING, job.error
    report, doc = _translated(env, job)
    assert report.problems == []
    rpr = doc.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
    lang, fonts = rpr.find(qn("w:lang")), rpr.find(qn("w:rFonts"))
    assert (lang.get(qn("w:val")), lang.get(qn("w:eastAsia"))) == ("ja-JP", "ja-JP")
    assert fonts.get(qn("w:eastAsia")) == "MS Gothic" and fonts.get(qn("w:eastAsiaTheme")) is None
    assert _render_event(env, job.job_id)["east_asia_font"] == "MS Gothic"


def test_set_language_leaves_the_first_page_and_the_headers_alone():
    doc = Document()
    doc.add_paragraph().add_run("Cover")._r.get_or_add_rPr().append(parse_xml(f'<w:lang {W} w:val="en-US"/>'))
    doc.element.body.insert(1, parse_xml(_COVER_END))
    body = doc.add_paragraph()
    body.add_run("Body")._r.get_or_add_rPr().append(parse_xml(f'<w:lang {W} w:val="en-US"/>'))
    fonts = parse_xml(f'<w:rFonts {W} w:ascii="Arial" w:eastAsiaTheme="minorEastAsia"/>')
    body.runs[0]._r.get_or_add_rPr().insert(0, fonts)
    doc.sections[-1].header.paragraphs[0].add_run("Header")._r.get_or_add_rPr().append(
        parse_xml(f'<w:lang {W} w:val="en-US"/>'))
    set_language(doc, "zh-CN")
    set_east_asia_font(doc, "SimSun")
    langs = [r.find(f"{qn('w:rPr')}/{qn('w:lang')}") for r in doc.element.body.iter(qn("w:r"))]
    assert [(n.get(qn("w:val")), n.get(qn("w:eastAsia"))) for n in langs] == [("en-US", None), ("zh-CN", "zh-CN")]
    assert (fonts.get(qn("w:eastAsia")), fonts.get(qn("w:eastAsiaTheme")), fonts.get(qn("w:ascii"))) == (
        "SimSun", None, "Arial")
    header = doc.sections[-1].header.paragraphs[0].runs[0]._r
    assert header.find(f"{qn('w:rPr')}/{qn('w:lang')}").get(qn("w:val")) == "en-US"


def test_the_east_asian_font_comes_from_the_language_template_first(tmp_path):
    package = tmp_path / "pkg.docx"
    doc = Document()
    rpr = doc.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
    rpr.find(qn("w:rFonts")).set(qn("w:eastAsia"), "SimSun")
    doc.save(str(package))
    assert east_asia_font("zh-CN", package) == "SimSun"
    assert east_asia_font("zh-CN") == "Microsoft YaHei" and east_asia_font("ja-JP") == "Yu Gothic"
    assert east_asia_font("fr-FR", package) is None


@pytest.mark.asyncio
async def test_a_segment_without_a_translation_is_shown_in_english(tpl, tmp_path):  # noqa: F811
    env = _env(tmp_path, tpl, FakeFactory(TranslatorOutput=FakeChain(translating(
        lambda key, text, repair: "" if text == "Submit it." else f"FR {text}"))))
    job = await env.orch.wait(env.orch.create(env.migration.job_id, "fr-FR").job_id)
    assert job.status == TranslationStatus.REVIEW_PENDING, job.error
    report, doc = _translated(env, job)
    assert report.problems == [] and "Submit it." in _texts(doc) and "FR Open the record." in _texts(doc)
    event = _render_event(env, job.job_id)
    assert event["kept_english"] == 1
    assert any(w.endswith("no translation; shown in English") and w.startswith(event["kept_ids"][0])
               for w in report.warnings)


# ── API ───────────────────────────────────────────────────────────────


def test_the_preview_and_render_report_wait_for_the_render(api):  # noqa: F811
    client, base = api.client, "/api/v1/translations"
    job = client.post(base, json={"migration_job_id": api.migration.job_id, "languages": [{"language": "de-DE"}]}).json()[0]
    assert _wait_status(client, job["job_id"], "TRANSLATING")["status"] == "TRANSLATING"  # no LLM: it waits there
    for path in ("export/preview", "render-report"):
        response = client.get(f"{base}/{job['job_id']}/{path}")
        assert response.status_code == 404 and response.json()["error"]["code"] == "NO_RENDER"


@pytest.mark.parametrize("api", [_fake_translator()], indirect=True)
def test_the_preview_serves_the_translated_document(api):  # noqa: F811
    client, base = api.client, "/api/v1/translations"
    job = client.post(base, json={"migration_job_id": api.migration.job_id, "languages": [{"language": "de-DE"}]}).json()[0]
    assert _wait_status(client, job["job_id"], "REVIEW_PENDING")["status"] == "REVIEW_PENDING"
    preview = client.get(f"{base}/{job['job_id']}/export/preview")
    assert preview.status_code == 200 and preview.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml")
    assert "inline" in preview.headers["content-disposition"] and "de-DE" in preview.headers["content-disposition"]
    assert "DE PURPOSE" in _texts(Document(io.BytesIO(preview.content)))
    report = client.get(f"{base}/{job['job_id']}/render-report").json()
    assert not report.get("problems") and report["job_id"] == job["job_id"] and report["mode"] == "final"
