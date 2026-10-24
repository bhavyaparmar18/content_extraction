"""Phase T2: glossary files, term matching, versions and the glossary API."""

import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api import translations
from app.api.auth import require_user
from app.config.settings import Settings, get_settings
from app.core.exception_handlers import register_exception_handlers
from app.schemas.v2 import GlossaryEntry, MigrationJob, TranslationJob
from app.services.translation import glossary_service
from app.services.translation.glossary import (
    GlossaryInvalid,
    match_terms,
    parse_glossary_file,
    target_present,
    terms_in,
    to_csv,
    to_xlsx,
)
from app.stores.glossary_store import GlossaryStore
from app.stores.migration_store import MigrationStore
from app.stores.translation_store import TranslationStore
from tests.db_helpers import seed_job_inputs, seed_users


@pytest.fixture
def settings(tmp_path, monkeypatch):
    monkeypatch.setenv("SOP_DB_PATH", str(tmp_path / "data" / "app.db"))
    s = Settings(project_root=tmp_path)
    s.resolve_paths(tmp_path)
    s.ensure_directories()
    return s


def _xlsx(rows, style_rows=None, title="Glossary") -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    if style_rows:
        notes = book.create_sheet("Style notes")
        for row in style_rows:
            notes.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def E(source, target="", **kw) -> GlossaryEntry:
    return GlossaryEntry(source_term=source, target_term=target, **kw)


# ── Reading files ─────────────────────────────────────────────────────


def test_csv_with_alternative_headers_and_semicolons():
    data = "English;German;Comment;DNT\ndeviation;Abweichung;QA term;\nVeeva Vault;;;yes\n\n".encode()
    parsed = parse_glossary_file(data, "terms.csv", "de-DE")
    assert [(e.source_term, e.target_term, e.do_not_translate) for e in parsed.entries] == [
        ("deviation", "Abweichung", False), ("Veeva Vault", "Veeva Vault", True)]
    assert parsed.entries[0].note == "QA term" and parsed.style_notes is None


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16"])
def test_csv_encodings_from_excel(encoding):
    data = "source_term\ttarget_term\ndeviation\tαπόκλιση\n".encode(encoding)
    assert parse_glossary_file(data, "g.csv", "el-GR").entries[0].target_term == "απόκλιση"


def test_windows_1252_csv():
    data = "source,target\nbatch,lot de fabrication\nfacility,établissement\n".encode("cp1252")
    assert parse_glossary_file(data, "g.csv", "fr-FR").entries[1].target_term == "établissement"


def test_xlsx_with_a_style_notes_sheet_and_the_language_code_as_header():
    data = _xlsx([["Source", "ja-JP", "Case sensitive"], ["SOP", "標準作業手順書", True], ["deviation", "逸脱", None]],
                 style_rows=[["Use the polite form."], ["Keep product names in English."]])
    parsed = parse_glossary_file(data, "ja.xlsx", "ja-JP")
    assert [(e.source_term, e.target_term, e.case_sensitive) for e in parsed.entries] == [
        ("SOP", "標準作業手順書", True), ("deviation", "逸脱", False)]
    assert parsed.style_notes == "Use the polite form.\nKeep product names in English."


def test_every_bad_row_is_reported_and_nothing_is_kept():
    data = ("source_term,target_term,do_not_translate\n"
            "deviation,,\n"            # no target
            "batch,Charge,maybe\n"      # bad flag
            "Batch,Los,\n"              # duplicate (case-insensitive)
            "Vault,Tresor,yes\n").encode()  # do-not-translate with another target
    with pytest.raises(GlossaryInvalid) as err:
        parse_glossary_file(data, "g.csv", "de-DE")
    problems = err.value.problems
    assert any(p.startswith("row 2:") and "no target term" in p for p in problems)
    assert any(p.startswith("row 3:") and "yes or no" in p for p in problems)
    assert any(p.startswith("row 4:") and "already listed in row 3" in p for p in problems)
    assert any(p.startswith("row 5:") and "do-not-translate" in p for p in problems)


def test_missing_columns_and_wrong_files():
    with pytest.raises(GlossaryInvalid, match="no target_term column"):
        parse_glossary_file(b"source,comment\nx,y\n", "g.csv", "fr-FR")
    with pytest.raises(GlossaryInvalid, match="empty"):
        parse_glossary_file(b"\n\n", "g.csv", "fr-FR")
    with pytest.raises(GlossaryInvalid, match="not a readable"):
        parse_glossary_file(b"not a zip", "g.xlsx", "fr-FR")
    with pytest.raises(GlossaryInvalid, match="unsupported"):
        parse_glossary_file(b"x", "g.txt", "fr-FR")


def test_export_round_trips_through_csv_and_xlsx(settings):
    glossary = glossary_service.build(
        {"glossary_id": "G", "version": 1, "language": "de-DE", "name": "G"},
        [E("deviation", "Abweichung", note="QA"), E("Veeva Vault", do_not_translate=True, case_sensitive=True)],
        "Formal register.")
    again = parse_glossary_file(to_csv(glossary).encode(), "g.csv", "de-DE")
    assert again.entries == glossary.entries
    again = parse_glossary_file(to_xlsx(glossary), "g.xlsx", "de-DE")
    assert again.entries == glossary.entries and again.style_notes == "Formal register."


# ── Matching the English source ──────────────────────────────────────


def test_source_terms_match_whole_words_plurals_and_spacing():
    entries = [E("deviation", "x"), E("facility", "x"), E("batch record", "x"), E("SOP", "x", case_sensitive=True)]
    text = "Deviations in two facilities go into the batch\nrecord. Not a deviationist; sop is not SOPs."
    assert [(m.entry.source_term, text[m.start:m.end]) for m in match_terms(text, entries)] == [
        ("deviation", "Deviations"), ("facility", "facilities"), ("batch record", "batch\nrecord"), ("SOP", "SOPs")]


def test_longer_terms_win_and_each_term_is_listed_once():
    entries = [E("batch", "Charge"), E("batch record", "Chargenprotokoll")]
    text = "Check the batch record, then sign the batch record and the batch."
    assert [e.source_term for e in terms_in(text, entries)] == ["batch record", "batch"]
    assert [m.entry.source_term for m in match_terms(text, entries)] == ["batch record", "batch record", "batch"]


# ── Checking a translation ────────────────────────────────────────────


@pytest.mark.parametrize("language, entry, text, present", [
    ("de-DE", E("deviation", "Abweichung"), "Alle Abweichungen sind zu melden.", True),
    ("de-DE", E("purpose", "Zweck"), "Zwei Personen prüfen.", False),
    ("de-DE", E("purpose", "Zweck"), "Zu diesem Zwecke gilt:", True),
    ("de-DE", E("batch record", "Chargenprotokoll"), "im Chargenprotokolls", True),
    ("fr-FR", E("facility", "établissement"), "Chaque ETABLISSEMENTS doit…", True),
    ("el-GR", E("deviation", "απόκλιση"), "Όλες οι αποκλίσεις καταγράφονται.", True),
    ("el-GR", E("deviation", "απόκλιση"), "Όλα τα σφάλματα καταγράφονται.", False),
    ("zh-CN", E("deviation", "偏差"), "所有偏差必须记录。", True),
    ("ja-JP", E("batch record", "製造 記録"), "製造記録に記入する。", True),
    ("ja-JP", E("deviation", "逸脱"), "すべての問題を記録する。", False),
    ("it-IT", E("Veeva Vault", do_not_translate=True, case_sensitive=True), "Caricare in Veeva Vault.", True),
    ("it-IT", E("Veeva Vault", do_not_translate=True, case_sensitive=True), "Caricare in veeva vault.", False),
    ("pt-BR", E("Veeva Vault", do_not_translate=True), "Carregar no VEEVA VAULT.", True),
])
def test_target_terms_in_translations(language, entry, text, present):
    assert target_present(text, entry, language) is present


# ── Versions ──────────────────────────────────────────────────────────


def _create(settings, glossary_id="GLS-de", language="de-DE", entries=None):
    store = GlossaryStore(settings.db_path, settings)
    record = glossary_service.create_version(store, settings, glossary_id=glossary_id, name="German",
                                             language=language, entries=entries or [E("deviation", "Abweichung")])
    return store, record


def test_review_edits_in_place_and_approved_versions_are_never_changed(settings):
    store, record = _create(settings)
    assert (record["version"], record["status"], record["total_entries"]) == (1, "in_review", 1)
    record = glossary_service.save_entries(store, settings, record, [E("deviation", "Abweichung"), E("batch", "Charge")])
    assert (record["version"], record["total_entries"]) == (1, 2)

    approved = glossary_service.approve(store, record)
    edited = glossary_service.save_entries(store, settings, approved, [E("batch", "Los")], style_notes="Sie-Form")
    assert (edited["version"], edited["status"]) == (2, "in_review")
    assert [e.target_term for e in glossary_service.load(store.get_version("GLS-de", 1)).entries] == [
        "Abweichung", "Charge"]
    assert glossary_service.load(edited).style_notes == "Sie-Form"


def test_a_glossary_keeps_its_language_and_invalid_edits_are_refused(settings):
    store, record = _create(settings)
    with pytest.raises(glossary_service.GlossaryLanguageMismatch):
        _create(settings, language="fr-FR")
    with pytest.raises(GlossaryInvalid, match="more than once"):
        glossary_service.save_entries(store, settings, record, [E("x", "a"), E("X", "b")])
    assert glossary_service.load(store.get_record_by_id(record["id"])).entries[0].source_term == "deviation"


# ── API ───────────────────────────────────────────────────────────────


@pytest.fixture
def client(settings):
    seed_users(settings.db_path, "alice")
    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(translations.router)
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[require_user] = lambda: {"userId": "alice"}
    with TestClient(app) as c:
        yield c


def test_languages(client):
    codes = [lang["code"] for lang in client.get("/api/v1/translations/languages").json()]
    assert codes == ["zh-CN", "fr-FR", "de-DE", "it-IT", "ja-JP", "pt-BR", "es-ES", "el-GR"]


def test_upload_review_approve_export_and_reupload(client, settings):
    files = {"file": ("German glossary.csv", b"source,target\ndeviation,Abweichung\n", "text/csv")}
    record = client.post("/api/v1/translations/glossaries", data={"language": "de-DE"}, files=files).json()
    assert (record["glossary_id"], record["name"], record["version"], record["status"], record["created_by"]) == (
        "German_glossary", "German glossary", 1, "in_review", "alice")
    assert (settings.glossary_upload_dir / "German_glossary_v1.csv").exists()

    base = "/api/v1/translations/glossaries/German_glossary"
    body = {"entries": [{"source_term": "deviation", "target_term": "Abweichung"},
                        {"source_term": "batch", "target_term": "Charge"}], "style_notes": "Sie"}
    assert client.put(f"{base}/entries", json=body).json()["total_entries"] == 2
    assert client.post(f"{base}/approve").json()["status"] == "approved"
    entries = client.get(f"{base}/entries").json()
    assert entries["style_notes"] == "Sie" and len(entries["entries"]) == 2

    csv_text = client.get(f"{base}/export", params={"format": "csv"}).content.decode("utf-8-sig")
    assert csv_text.splitlines()[2] == "batch,Charge,,no,no"
    book = load_workbook(io.BytesIO(client.get(f"{base}/export").content))
    assert book["Glossary"]["A3"].value == "batch" and book["Style notes"]["A1"].value == "Sie"

    again = client.post("/api/v1/translations/glossaries", data={"language": "de-DE"}, files=files).json()
    assert again["version"] == 2
    assert [r["version"] for r in client.get("/api/v1/translations/glossaries",
                                             params={"language": "de-DE"}).json()] == [2, 1]
    assert client.get(base, params={"version": 1}).json()["status"] == "approved"


def test_upload_errors(client):
    url = "/api/v1/translations/glossaries"
    bad = client.post(url, data={"language": "de-DE"}, files={"file": ("g.csv", b"source,target\nx,\n", "text/csv")})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "GLOSSARY_INVALID"
    assert "row 2" in bad.json()["error"]["fieldErrors"][0]["message"]
    assert client.post(url, data={"language": "pt-PT", "name": "x"}).json()["error"]["code"] == "UNSUPPORTED_LANGUAGE"
    wrong = client.post(url, data={"language": "de-DE"}, files={"file": ("g.txt", b"x", "text/plain")})
    assert wrong.json()["error"]["code"] == "UNSUPPORTED_GLOSSARY_FILE_TYPE"
    assert client.post(url, data={"language": "de-DE"}).json()["error"]["code"] == "GLOSSARY_NAME_REQUIRED"
    assert client.get(f"{url}/nope").status_code == 404

    empty = client.post(url, data={"language": "fr-FR", "name": "French terms"}).json()
    assert (empty["glossary_id"], empty["total_entries"]) == ("French_terms", 0)
    clash = client.post(url, data={"language": "de-DE", "name": "French terms"})
    assert clash.status_code == 409 and clash.json()["error"]["code"] == "GLOSSARY_LANGUAGE_MISMATCH"


def test_a_glossary_a_translation_job_uses_cannot_be_deleted(client, settings):
    url = "/api/v1/translations/glossaries"
    client.post(url, data={"language": "fr-FR", "name": "FR"})
    client.post(url, data={"language": "fr-FR", "name": "FR"})
    sop_id = seed_job_inputs(settings.db_path, settings)
    MigrationStore(settings.db_path, settings).create_job(
        MigrationJob(job_id="MIG-1", sop_record_id=sop_id, template_id="tpl", template_version=1))
    TranslationStore(settings.db_path, settings).create_job(TranslationJob(
        job_id="TRN-1", migration_job_id="MIG-1", language="fr-FR", glossary_id="FR", glossary_version=1))

    blocked = client.delete(f"{url}/FR", params={"version": 1})
    assert blocked.status_code == 409 and blocked.json()["error"]["code"] == "RECORD_IN_USE"
    assert client.delete(f"{url}/FR", params={"version": 2}).json() == {
        "deleted": True, "glossary_id": "FR", "version": 2}
