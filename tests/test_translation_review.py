"""Phase T7: review of a translation job. Segment edits (that segment or every segment with the same English text),
find and replace, issue resolutions, the render going out of date and back, the approval gate and the approved
export; then the whole flow through the API."""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.schemas.v2 import (
    ArtifactKind,
    Glossary,
    GlossaryEntry,
    QualityReport,
    RenderReport,
    Segment,
    SegmentKind,
    SegmentLocator,
    SegmentSet,
    TranslationStatus,
)
from app.services.translation import review_service
from app.services.translation.checks import repair_candidates
from app.services.translation.render import TRANSLATED
from app.services.translation.review_service import Occurrence, ReviewError
from app.services.translation.stages import CHECKS, current_segments, translation_versions
from app.services.translation.translator import TERMS
from tests.test_translation_jobs import _env, _fake_translator, _glossary, _wait_status, api, translating  # noqa: F401
from tests.db_helpers import seed_users
from tests.test_v2_drafter import FakeChain, FakeFactory
from tests.test_v2_renderer import tpl  # noqa: F401  (fixture)

REVIEWER = "reviewer-1"


def _factory(transform):
    return FakeFactory(TranslatorOutput=FakeChain(translating(transform)))


def _drops_numbers(key, text, repair):
    """German-ish, but two texts lose their number (also on repair): open NUMBERS gate issues."""
    return f"DE {text}".replace("5 days", "days").replace("Figure 1", "Figure")


async def _in_review(tpl, tmp_path, transform=lambda key, text, repair: f"DE {text}", language="de-DE",
                     glossary=False):  # noqa: F811
    env = _env(tmp_path, tpl, _factory(transform))
    seed_users(env.e.store.db_path, REVIEWER)  # approved_by is a user
    named = ("g-de", _glossary(env)["version"]) if glossary else None
    job = await env.orch.wait(env.orch.create(env.migration.job_id, language, named).job_id, 60)
    assert job.status == TranslationStatus.REVIEW_PENDING
    env.job_id = job.job_id
    return env


def _ctx(env, actor: str = REVIEWER):
    return env.orch.context(env.store.get_job(env.job_id), actor)


def _segments(env) -> dict[str, Segment]:
    ctx = _ctx(env)
    return {s.segment_id: s for s in current_segments(ctx, ctx.require(ArtifactKind.SEGMENTS, SegmentSet)).segments}


def _report(env) -> QualityReport:
    return _ctx(env).latest(ArtifactKind.QUALITY_REPORT, QualityReport, CHECKS)


def _events(env, name: str) -> list[dict]:
    return [e for e in env.store.list_events(env.job_id) if e["event"] == name]


# ── Segment edits ─────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_an_edit_goes_to_every_segment_with_the_same_english_unless_asked_otherwise(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path)
    before = translation_versions(env.store.get_job(env.job_id))
    # "Role" is two table headers of TGT-4 and a template table header.
    result = review_service.edit_segment(_ctx(env), "SEG-0011", "  Rolle ")
    assert sorted(s["segment_id"] for s in result["segments"]) == ["SEG-0011", "SEG-0015", "SEG-0040"]
    after = translation_versions(env.store.get_job(env.job_id))
    assert {s for s in after if after[s] != before[s]} == {"TGT-4", "template"}  # only the scopes it touched
    segments = _segments(env)
    assert {segments[s].target_text for s in ("SEG-0011", "SEG-0015", "SEG-0040")} == {"Rolle"}
    assert {segments[s].origin.value for s in ("SEG-0011", "SEG-0015", "SEG-0040")} == {"human"}
    assert result["render_current"] is False and result["validation"]["version"] == 2

    review_service.edit_segment(_ctx(env), "SEG-0040", "ROLLE", "this")
    segments = _segments(env)
    assert (segments["SEG-0040"].target_text, segments["SEG-0011"].target_text) == ("ROLLE", "Rolle")
    # Its siblings are all human edits now: they are not overwritten.
    result = review_service.edit_segment(_ctx(env), "SEG-0015", "Funktion")
    assert [s["segment_id"] for s in result["segments"]] == ["SEG-0015"]
    assert _segments(env)["SEG-0011"].target_text == "Rolle"

    event = _events(env, "segment_edited")[0]
    assert event["actor"] == REVIEWER and event["detail"]["applied_to"] == ["SEG-0011", "SEG-0015", "SEG-0040"]
    assert event["detail"]["before"] == "DE Role" and event["detail"]["after"] == "Rolle"
    written = env.store.get_job(env.job_id).latest_artifact(ArtifactKind.TRANSLATION, "TGT-4")
    assert written.created_by == REVIEWER


@pytest.mark.asyncio
async def test_an_edit_the_document_could_not_show_is_refused_and_nothing_is_written(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path)
    count = len(env.store.get_job(env.job_id).artifacts)
    for text in ("   ", "DE Record every deviation within 5 days."):  # empty; the {{ref:}} token lost
        with pytest.raises(ReviewError) as refused:
            review_service.edit_segment(_ctx(env), "SEG-0021", text)
        assert (refused.value.status, refused.value.code) == (422, "SEGMENT_INVALID") and refused.value.details
    with pytest.raises(ReviewError) as missing:
        review_service.edit_segment(_ctx(env), "SEG-9999", "x")
    assert missing.value.code == "SEGMENT_NOT_FOUND"
    assert len(env.store.get_job(env.job_id).artifacts) == count

    # Markers, a cross-reference's number and table cell separators must stay too.
    ctx = _ctx(env)
    tagged = Segment(segment_id="SEG-9001", kind=SegmentKind.TEMPLATE_TEXT, source_text="⟦1⟧Note:⟦/1⟧ read this",
                     locator=SegmentLocator(block_path="body/p[0]"))
    xref = Segment(segment_id="SEG-9002", kind=SegmentKind.XREF, source_text="see section ⟦1⟧",
                   locator=SegmentLocator(claim_id="C-1", ref_index=0))
    row = Segment(segment_id="SEG-9003", kind=SegmentKind.CLAIM, source_text="QA | Quality Assurance",
                  locator=SegmentLocator(section_id="TGT-3", slot_id="S-1", claim_id="C-2"))
    assert review_service.structural_problems(ctx, tagged, "Hinweis: lies das")
    assert not review_service.structural_problems(ctx, tagged, "Lies das ⟦1⟧Hinweis:⟦/1⟧")
    assert review_service.structural_problems(ctx, xref, "siehe Abschnitt")
    assert review_service.structural_problems(ctx, row, "QA Qualitätssicherung")
    assert not review_service.structural_problems(ctx, row, "QA | Qualitätssicherung")


@pytest.mark.asyncio
async def test_an_edit_reopens_its_own_issues_and_keeps_the_other_resolutions(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path, _drops_numbers)
    open_gates = {i.issue_id for i in _report(env).open_gate_issues}
    assert open_gates == {"TQ-SEG-0021-NUMBERS", "TQ-SEG-0030-NUMBERS"}
    for issue_id in sorted(open_gates):
        review_service.resolve_issue(_ctx(env), issue_id, "checked with the author")
    assert _report(env).gates_passed

    # An edit that still lacks the number is accepted (not structural), and reopens only its own issue.
    result = review_service.edit_segment(_ctx(env), "SEG-0021", "DE Abweichungen binnen Tagen erfassen ({{ref:SRC-4.2}}).")
    issues = {i.issue_id: i for i in _report(env).issues}
    assert not issues["TQ-SEG-0021-NUMBERS"].resolved and issues["TQ-SEG-0030-NUMBERS"].resolved
    assert [i["issue_id"] for i in result["issues"]] == ["TQ-SEG-0021-NUMBERS"]
    assert "SEG-0021" not in repair_candidates(_segments(env).values(), _report(env))  # a human edit is never repaired


# ── Find and replace ──────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_find_matches_whole_words_case_insensitively_outside_markers(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path, glossary=True)
    preview = review_service.find_matches(_ctx(env), "role")
    assert [m["segment_id"] for m in preview["matches"]] == ["SEG-0011", "SEG-0015", "SEG-0040"]  # not "ROLES"
    assert preview["matches"][0]["match"] == "Role" and preview["matches"][0]["context_before"] == "DE "
    assert preview["longer_words"] == {"roles": 1, "ROLES": 1} and "longer word" in preview["warnings"][-1]
    assert preview["translation_versions"] == translation_versions(env.store.get_job(env.job_id))
    assert review_service.find_matches(_ctx(env), "role", whole_word=False)["total"] == 5  # + roles, ROLES
    assert review_service.find_matches(_ctx(env), "role", case_sensitive=True)["total"] == 0
    assert [m["segment_id"] for m in review_service.find_matches(_ctx(env), "role", section="template")["matches"]] == [
        "SEG-0040"]
    # Never inside a {{ref:}} token: "SRC" is only in one.
    assert review_service.find_matches(_ctx(env), "SRC")["total"] == 0
    assert review_service.find_matches(_ctx(env), "Abweichung")["glossary_terms"] == ["deviation"]
    for bad in ("", "{{ref:SRC-4.2}}", "⟦1⟧"):
        with pytest.raises(ReviewError) as refused:
            review_service.find_matches(_ctx(env), bad)
        assert refused.value.code == "REPLACE_INVALID"


@pytest.mark.asyncio
async def test_chinese_and_japanese_match_inside_words(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path, lambda key, text, repair: f"確認 {text}", language="ja-JP")
    assert review_service.find_matches(_ctx(env), "認")["total"] == len(
        [s for s in _segments(env).values() if s.target_text and "確認" in s.target_text])


@pytest.mark.asyncio
async def test_replace_one_occurrence_a_chosen_set_or_all(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path)
    versions = review_service.find_matches(_ctx(env), "DE")["translation_versions"]
    one = review_service.replace(_ctx(env), "DE", "", [Occurrence("SEG-0014", 0)], versions)
    assert one["replaced"] == 1 and [s["segment_id"] for s in one["segments"]] == ["SEG-0014"]
    segments = _segments(env)
    assert segments["SEG-0014"].target_text == " Owns the process" and segments["SEG-0013"].target_text == "DE Process Owner"
    assert segments["SEG-0014"].origin.value == "human"

    with pytest.raises(ReviewError) as stale:  # the translations changed since that preview
        review_service.replace(_ctx(env), "DE", "", "all", versions)
    assert stale.value.code == "PREVIEW_OUT_OF_DATE"
    versions = translation_versions(env.store.get_job(env.job_id))
    with pytest.raises(ReviewError) as unknown:
        review_service.replace(_ctx(env), "DE", "", [Occurrence("SEG-0014", 0)], versions)
    assert unknown.value.code == "REPLACE_INVALID"

    chosen = review_service.replace(_ctx(env), "Process", "Prozess", [Occurrence("SEG-0013", 0),
                                                                     Occurrence("SEG-0014", 0)], versions)
    assert chosen["replaced"] == 2
    every = review_service.replace(_ctx(env), "DE", "GER", "all", translation_versions(env.store.get_job(env.job_id)))
    assert every["replaced"] == len(every["segments"]) > 30 and not every["skipped"]
    assert not review_service.find_matches(_ctx(env), "DE")["total"]
    with pytest.raises(ReviewError) as none:
        review_service.replace(_ctx(env), "DE", "x", "all", translation_versions(env.store.get_job(env.job_id)))
    assert none.value.code == "NO_MATCHES"
    event = _events(env, "segments_replaced")[-1]["detail"]
    assert event["targets"] == "all" and event["find"] == "DE" and event["count"] == every["replaced"]


@pytest.mark.asyncio
async def test_replacing_all_of_a_document_terms_translation_agrees_the_new_one(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path)
    ctx = _ctx(env)
    agreed = ctx.next_version(ArtifactKind.GLOSSARY, TERMS)
    ctx.write(ArtifactKind.GLOSSARY, Glossary(
        glossary_id=f"{env.job_id}-terms", version=agreed, language="de-DE",
        name="Document terms", entries=[GlossaryEntry(source_term="process owner", target_term="DE Process Owner")]),
        scope=TERMS)
    preview = review_service.find_matches(_ctx(env), "DE Process Owner")
    assert preview["document_terms"] == ["process owner"] and preview["warnings"]
    result = review_service.replace(_ctx(env), "DE Process Owner", "Prozesseigner", "all",
                                    preview["translation_versions"])
    assert result["term_updated"] == ["process owner"]
    terms = _ctx(env).latest(ArtifactKind.GLOSSARY, Glossary, TERMS)
    assert terms.version == agreed + 1 and terms.entries[0].target_term == "Prozesseigner"
    assert not [i for i in _report(env).issues if i.issue_id.startswith("TQ-SEG-0013-TERM")]


# ── Issues, render, approval, export ──────────────────────────────────


@pytest.mark.asyncio
async def test_resolving_an_issue_saves_the_next_report(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path, _drops_numbers)
    before = _report(env)
    result = review_service.resolve_issue(_ctx(env), "TQ-SEG-0021-NUMBERS", "  5 days is in the next sentence ")
    report = _report(env)
    assert report.version == before.version + 1 and len(report.open_gate_issues) == 1
    issue = next(i for i in report.issues if i.issue_id == "TQ-SEG-0021-NUMBERS")
    assert issue.resolved and issue.resolution_note == f"{REVIEWER}: 5 days is in the next sentence"
    assert [b["issue_id"] for b in result["blockers"]] == ["TQ-SEG-0030-NUMBERS"]
    assert _events(env, "issue_resolved")[0]["detail"]["gate"] == "numerical_change"
    for issue_id, code in (("TQ-SEG-0021-NUMBERS", "ISSUE_RESOLVED"), ("TQ-NOPE", "ISSUE_NOT_FOUND")):
        with pytest.raises(ReviewError) as refused:
            review_service.resolve_issue(_ctx(env), issue_id, "again")
        assert refused.value.code == code


@pytest.mark.asyncio
async def test_the_render_goes_out_of_date_with_an_edit_and_back_with_a_new_render(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path)
    assert review_service.render_state(_ctx(env)).current
    assert "translation_versions" in _events(env, "render")[-1]["detail"]
    # A render recorded before T7 (no versions in its event) is current while no translation is newer.
    docx = env.store.get_job(env.job_id).latest_artifact(ArtifactKind.DOCX, TRANSLATED)
    env.store.add_event(env.job_id, "render", "system", {"docx_version": docx.version})
    assert review_service.render_state(_ctx(env)).current

    review_service.edit_segment(_ctx(env), "SEG-0014", "Besitzt den Prozess")
    assert not review_service.render_state(_ctx(env)).current
    assert [b.code for b in review_service.approval_blockers(_ctx(env))] == ["RENDER_OUT_OF_DATE"]

    env.orch.rerender(env.job_id, REVIEWER)
    with pytest.raises(ReviewError) as running:  # no edit while it renders
        review_service.require_review(env.store.get_job(env.job_id), env.orch.is_running(env.job_id))
    assert running.value.code == "JOB_RUNNING"
    job = await env.orch.wait(env.job_id, 60)
    assert job.status == TranslationStatus.REVIEW_PENDING
    state = review_service.render_state(_ctx(env))
    assert state.current and state.docx.version == docx.version + 1
    assert _events(env, "render_requested")[0]["actor"] == REVIEWER


@pytest.mark.asyncio
async def test_approval_waits_for_a_clean_current_render_and_no_open_gate_or_critical_issue(tpl, tmp_path):  # noqa: F811
    env = await _in_review(tpl, tmp_path, _drops_numbers)
    codes = lambda: [b.code for b in review_service.approval_blockers(_ctx(env))]  # noqa: E731
    assert codes() == ["OPEN_GATE", "OPEN_GATE"]
    with pytest.raises(ReviewError) as blocked:
        review_service.approve(_ctx(env))
    assert blocked.value.code == "APPROVAL_BLOCKED" and len(blocked.value.details) == 2

    # Fix one, accept the other: the gate is passed, but the fix needs a new render.
    review_service.edit_segment(_ctx(env), "SEG-0030", "Abbildung 1: Bewertungsablauf")
    review_service.resolve_issue(_ctx(env), "TQ-SEG-0021-NUMBERS", "the period is given in the next sentence")
    assert codes() == ["RENDER_OUT_OF_DATE"]
    env.orch.rerender(env.job_id, REVIEWER)
    await env.orch.wait(env.job_id, 60)
    assert codes() == []

    # A render with post-render problems blocks too.
    ctx = _ctx(env)
    report = ctx.latest(ArtifactKind.RENDER_REPORT, RenderReport, TRANSLATED)
    ctx.write(ArtifactKind.RENDER_REPORT, report.model_copy(update={
        "problems": ["header differs"], "version": ctx.next_version(ArtifactKind.RENDER_REPORT, TRANSLATED)}),
        scope=TRANSLATED)
    assert codes() == ["RENDER_PROBLEMS"]
    ctx.write(ArtifactKind.RENDER_REPORT, report.model_copy(update={
        "version": ctx.next_version(ArtifactKind.RENDER_REPORT, TRANSLATED)}), scope=TRANSLATED)

    with pytest.raises(ReviewError) as early:
        review_service.approved_export(_ctx(env))
    assert early.value.code == "NOT_APPROVED"
    job = review_service.approve(_ctx(env), "fine")
    assert job.status == TranslationStatus.APPROVED and job.approved_by == REVIEWER
    approved = _events(env, "approved")[0]["detail"]
    assert approved["note"] == "fine" and approved["docx_version"] == 2 and approved["quality_report_version"] >= 3

    with pytest.raises(ReviewError) as again:
        review_service.require_review(env.store.get_job(env.job_id), False)
    assert again.value.code == "NOT_IN_REVIEW"
    data, name = review_service.approved_export(_ctx(env))
    ref = env.store.get_job(env.job_id).latest_artifact(ArtifactKind.DOCX, TRANSLATED)
    assert data == env.orch.artifacts.read_bytes(ref) and name == f"{env.job_id}_de-DE_v{ref.version}.docx"
    assert _events(env, "export")[0]["detail"] == {"docx_version": ref.version, "sha256": ref.sha256}


# ── Through the API ───────────────────────────────────────────────────


def _doc_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


@pytest.mark.parametrize("api", [_fake_translator()], indirect=True)
def test_edit_render_approve_and_export_through_the_api(api):
    client, base = api.client, "/api/v1/translations"
    job_id = client.post(base, json={"migration_job_id": api.migration.job_id,
                                     "languages": [{"language": "de-DE"}]}).json()[0]["job_id"]
    job = _wait_status(client, job_id, "REVIEW_PENDING")
    assert job["review"]["can_approve"] and job["review"]["render_current"]
    early = client.get(f"{base}/{job_id}/export/word")
    assert early.status_code == 409 and early.json()["error"]["code"] == "NOT_APPROVED"

    edited = client.patch(f"{base}/{job_id}/segments/SEG-0014", json={"target_text": "Besitzt den Prozess"})
    assert edited.status_code == 200, edited.text
    assert edited.json()["segments"][0]["origin"] == "human" and edited.json()["render_current"] is False
    bad = client.patch(f"{base}/{job_id}/segments/SEG-0021", json={"target_text": "Ohne Verweis"})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "SEGMENT_INVALID"

    preview = client.post(f"{base}/{job_id}/replace/preview", json={"find": "Role"}).json()
    assert preview["total"] == 3
    replaced = client.post(f"{base}/{job_id}/replace", json={
        "find": "Role", "replace": "Rolle", "targets": [{"segment_id": "SEG-0040", "occurrence": 0}],
        "expected_versions": preview["translation_versions"]})
    assert replaced.status_code == 200 and replaced.json()["replaced"] == 1
    stale = client.post(f"{base}/{job_id}/replace", json={"find": "Role", "replace": "Rolle", "targets": "all",
                                                          "expected_versions": preview["translation_versions"]})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "PREVIEW_OUT_OF_DATE"

    job = client.get(f"{base}/{job_id}").json()
    assert [b["code"] for b in job["review"]["blockers"]] == ["RENDER_OUT_OF_DATE"]
    blocked = client.post(f"{base}/{job_id}/approve")
    assert blocked.status_code == 409 and blocked.json()["error"]["code"] == "APPROVAL_BLOCKED"
    assert blocked.json()["error"]["fieldErrors"]

    rendering = client.post(f"{base}/{job_id}/render")
    assert rendering.status_code == 200 and rendering.json()["status"] == "RENDERING"
    job = _wait_status(client, job_id, "REVIEW_PENDING")
    assert job["review"]["can_approve"]
    resolved = client.post(f"{base}/{job_id}/issues/TQ-NOPE/resolve", json={"note": "x"})
    assert resolved.status_code == 404 and resolved.json()["error"]["code"] == "ISSUE_NOT_FOUND"

    approved = client.post(f"{base}/{job_id}/approve", json={"note": "Reviewed"})
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "APPROVED" and approved.json()["approved_by"] == "reviewer-1"
    exported = client.get(f"{base}/{job_id}/export/word")
    assert exported.status_code == 200 and "attachment" in exported.headers["content-disposition"]
    text = _doc_text(exported.content)
    assert "Besitzt den Prozess" in text and "Rolle" in text

    after = client.patch(f"{base}/{job_id}/segments/SEG-0013", json={"target_text": "Prozesseigner"})
    assert after.status_code == 409 and after.json()["error"]["code"] == "NOT_IN_REVIEW"
    again = client.post(f"{base}/{job_id}/render")
    assert again.status_code == 409 and again.json()["error"]["code"] == "NOT_IN_REVIEW"
    audit = [e["event"] for e in client.get(f"{base}/{job_id}/audit").json()]
    for event in ("segment_edited", "segments_replaced", "render_requested", "approved", "export"):
        assert event in audit
