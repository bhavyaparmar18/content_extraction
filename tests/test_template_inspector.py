"""Unit tests for TemplateInspector."""

import pytest
from pathlib import Path
from docx import Document
from docx.shared import RGBColor

from app.services.migration.template_inspector import TemplateInspector


@pytest.fixture
def sample_template_docx(tmp_path) -> Path:
    """Create a mock .docx template for testing."""
    docx_path = tmp_path / "test_template.docx"
    doc = Document()

    # Preamble / cover page table
    t0 = doc.add_table(rows=2, cols=2)
    t0.cell(0, 0).text = "Title"
    t0.cell(0, 1).text = "${vault:title__v}"
    t0.cell(1, 0).text = "Document Number"
    t0.cell(1, 1).text = "${vault:document_number__v}"

    # Heading 1: Table of Contents
    doc.add_heading("Table of Content", level=1)

    # Heading 1: 1 PURPOSE
    doc.add_heading("1 PURPOSE", level=1)
    p1 = doc.add_paragraph("This Directive/SOP/Work Instruction/Guidance:")
    # Add blue instruction text
    p2 = doc.add_paragraph()
    r2 = p2.add_run("Brief description of what the document is about? (Instruction in blue)")
    r2.font.color.rgb = RGBColor(0, 112, 192)

    # Heading 1: 3 DEFINITIONS & ABBREVIATIONS
    doc.add_heading("3 DEFINITIONS & ABBREVIATIONS", level=1)
    t1 = doc.add_table(rows=2, cols=2)
    t1.cell(0, 0).text = "Term"
    t1.cell(0, 1).text = "Definition/Explanation"
    t1.cell(1, 0).text = "[Term 1]"
    t1.cell(1, 1).text = "[Definition 1]"

    # Optional section: 9 DISTRIBUTION
    doc.add_heading("9 DISTRIBUTION OF CONTROLLED PRINTS/COPIES (OPTIONAL)", level=1)
    p3 = doc.add_paragraph()
    r3 = p3.add_run("Use this chapter if distribution must be defined. (Blue instruction)")
    r3.font.color.rgb = RGBColor(0, 0, 255)

    doc.save(str(docx_path))
    return docx_path


def test_template_inspector_basic(sample_template_docx):
    inspector = TemplateInspector()
    profile = inspector.inspect(sample_template_docx)

    assert profile.filename == "test_template.docx"
    assert profile.total_paragraphs > 0
    assert profile.total_tables == 2

    # Check that blue instruction paragraphs are flagged
    blue_paras = [p for p in profile.paragraphs if p.is_blue_instruction]
    assert len(blue_paras) >= 2

    # Check table placeholders
    t0_info = profile.tables[0]
    assert t0_info.has_placeholder_text is True

    # Check preceding heading resolution
    t1_info = profile.tables[1]
    assert t1_info.preceding_heading is not None
    assert "DEFINITIONS" in t1_info.preceding_heading
