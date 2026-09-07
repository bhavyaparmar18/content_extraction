"""Integration tests for DocxMigrator pipeline with mocked LLM."""

import pytest
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock
from docx import Document
from docx.shared import RGBColor

from app.config.settings import Settings
from app.schemas.migration import (
    DocxMigrationOutput,
    MigrationMetadata,
    MigrationSection,
    MigrationElement,
    MigrationTableCell,
    MigrationIconRef,
)
from app.services.llm.chain_factory import ChainFactory
from app.services.migration.docx_migrator import DocxMigrator
from app.services.migration.schemas import (
    MigrationPlan,
    SectionPlan,
    ElementPlacement,
    CalloutStyleDef,
    PlaceholderTablePlan,
)


@pytest.fixture
def mock_chain_factory():
    factory = MagicMock(spec=ChainFactory)

    # Mock migration plan returned by LLM planner
    mock_plan = MigrationPlan(
        template_name="test_template.docx",
        document_type_detected="Guidance",
        font_family="Arial",
        section_plans=[
            SectionPlan(
                template_section_heading="1 PURPOSE",
                template_paragraph_indices_to_delete=[1],
                elements=[
                    ElementPlacement(
                        source_section_title="1 PURPOSE",
                        source_element_index=0,
                        source_element_type="paragraph",
                        target_section_heading="1 PURPOSE",
                        placement_order=0,
                        action="insert_paragraph",
                        embed_icons_inline=True,
                    )
                ],
            ),
            SectionPlan(
                template_section_heading="6 PRINCIPLES FOR DOCUMENT WRITING",
                is_unmapped_source=True,
                insertion_after_section="1 PURPOSE",
                elements=[
                    ElementPlacement(
                        source_section_title="6 PRINCIPLES FOR DOCUMENT WRITING",
                        source_element_index=0,
                        source_element_type="heading",
                        target_section_heading="6 PRINCIPLES FOR DOCUMENT WRITING",
                        placement_order=0,
                        action="insert_heading",
                        heading_level=2,
                    ),
                    ElementPlacement(
                        source_section_title="6 PRINCIPLES FOR DOCUMENT WRITING",
                        source_element_index=1,
                        source_element_type="paragraph",
                        target_section_heading="6 PRINCIPLES FOR DOCUMENT WRITING",
                        placement_order=1,
                        action="insert_callout",
                        callout_type="executive_summary",
                    ),
                ],
            ),
            SectionPlan(
                template_section_heading="9 DISTRIBUTION (OPTIONAL)",
                has_source_content=False,
                fallback_action="delete_section",
            ),
        ],
        callout_styles=[
            CalloutStyleDef(
                callout_type="executive_summary",
                display_name="Executive Summary",
                background_color_hex="#D9E1F2",
                left_border_color_hex="#2F5597",
            )
        ],
        overall_confidence=0.95,
    )

    mock_structured_planner = MagicMock()
    mock_structured_planner.ainvoke = AsyncMock(return_value=mock_plan)
    factory.create_structured_planner.return_value = mock_structured_planner

    return factory


@pytest.fixture
def sample_template(tmp_path) -> Path:
    doc = Document()
    doc.add_heading("Table of Content", level=1)
    doc.add_heading("1 PURPOSE", level=1)
    p_blue = doc.add_paragraph()
    r = p_blue.add_run("Blue instruction to delete.")
    r.font.color.rgb = RGBColor(0, 112, 192)

    doc.add_heading("9 DISTRIBUTION (OPTIONAL)", level=1)
    doc.add_paragraph("Distribution info.")

    tpl_path = tmp_path / "template.docx"
    doc.save(str(tpl_path))
    return tpl_path


@pytest.fixture
def sample_extracted() -> DocxMigrationOutput:
    return DocxMigrationOutput(
        document_id="test_doc_001",
        metadata=MigrationMetadata(
            document_id="test_doc_001",
            title="Good Writing Practice",
        ),
        sections=[
            MigrationSection(
                title="0 PREAMBLE",
                elements=[MigrationElement(element_type="paragraph", text="Cover Page Info")],
            ),
            MigrationSection(
                title="1 PURPOSE",
                elements=[
                    MigrationElement(
                        element_type="paragraph",
                        text="This document establishes writing guidance.",
                    )
                ],
            ),
            MigrationSection(
                title="6 PRINCIPLES FOR DOCUMENT WRITING",
                elements=[
                    MigrationElement(element_type="heading", level=2, text="6.1 LANGUAGE"),
                    MigrationElement(
                        element_type="paragraph",
                        text="Executive Summary: Keep sentences clear and active.",
                    ),
                ],
            ),
        ],
    )


@pytest.mark.asyncio
async def test_docx_migrator_end_to_end(tmp_path, mock_chain_factory, sample_template, sample_extracted):
    settings = Settings()
    migrator = DocxMigrator(settings, mock_chain_factory)

    output_path = tmp_path / "migrated_output.docx"

    result = await migrator.migrate(
        extracted=sample_extracted,
        template_path=sample_template,
        output_path=output_path,
    )

    assert output_path.exists()
    assert result.plan.overall_confidence == 0.95
    assert result.qa_report.status in ("pass", "pass_with_warnings")
    assert result.qa_report.blue_text_remaining == 0

    # Verify generated document contents
    migrated_doc = Document(str(output_path))
    texts = [p.text for p in migrated_doc.paragraphs]

    # Verify Purpose content was inserted
    assert any("This document establishes writing guidance" in t for t in texts)

    # Verify unmapped section was inserted with number prefix cleanly stripped
    assert any("LANGUAGE" in t for t in texts)
    lang_heading = next(p for p in migrated_doc.paragraphs if "LANGUAGE" in p.text)
    assert lang_heading.style.name == "Heading 2"
    assert not any("6.1 LANGUAGE" in t for t in texts)

    # Verify blue text was deleted
    assert all("Blue instruction" not in t for t in texts)

    # Verify optional deleted section is gone
    assert all("DISTRIBUTION" not in t for t in texts)


@pytest.mark.asyncio
async def test_unmapped_subsection_merging(tmp_path, mock_chain_factory, sample_template):
    """Verify that an unmapped subsection like '6.1 Language' is merged into parent section '6 Principles' rather than creating a new top-level section."""
    settings = Settings()
    migrator = DocxMigrator(settings, mock_chain_factory)
    output_path = tmp_path / "migrated_subsections.docx"

    extracted = DocxMigrationOutput(
        document_id="test_doc_sub",
        sections=[
            MigrationSection(
                title="6 PRINCIPLES FOR DOCUMENT WRITING",
                elements=[
                    MigrationElement(element_type="paragraph", text="Core principles intro."),
                    MigrationElement(element_type="list", items=["1. Active voice:", "2. Brevity:"]),
                ],
            ),
            # An unmapped subsection
            MigrationSection(
                title="6.1 SUBSECTION DETAILS",
                elements=[
                    MigrationElement(element_type="heading", level=2, text="6.1.1 Nested Rule"),
                    MigrationElement(element_type="paragraph", text="Details of the rule."),
                ],
            ),
        ],
    )

    result = await migrator.migrate(
        extracted=extracted,
        template_path=sample_template,
        output_path=output_path,
    )

    migrated_doc = Document(str(output_path))
    # Check that '6.1 SUBSECTION DETAILS' did NOT create a top-level Heading 1
    h1_texts = [p.text for p in migrated_doc.paragraphs if p.style.name == "Heading 1"]
    assert not any("SUBSECTION DETAILS" in t for t in h1_texts)

    # Check that content from 6.1 was merged into the document and previous section number was stripped
    all_texts = [p.text for p in migrated_doc.paragraphs] + [
        c.text for t in migrated_doc.tables for r in t.rows for c in r.cells
    ]
    assert any("Nested Rule" in t for t in all_texts)
    assert not any("6.1.1 Nested Rule" in t for t in all_texts)
    assert any("Details of the rule." in t for t in all_texts)
    assert any("1. Active voice:" in t for t in all_texts)


@pytest.mark.asyncio
async def test_reclassified_heading_as_list(tmp_path, mock_chain_factory, sample_template):
    """Verify that a heading with bullet/list format is reclassified to insert_list and its text is not dropped."""
    settings = Settings()
    migrator = DocxMigrator(settings, mock_chain_factory)
    output_path = tmp_path / "migrated_reclassified.docx"

    extracted = DocxMigrationOutput(
        document_id="test_doc_reclassified",
        sections=[
            MigrationSection(
                title="1 PURPOSE",
                elements=[
                    MigrationElement(element_type="paragraph", text="Purpose statement."),
                    MigrationElement(element_type="heading", text="• Always wear PPE:"),
                ],
            ),
        ],
    )

    result = await migrator.migrate(
        extracted=extracted,
        template_path=sample_template,
        output_path=output_path,
    )

    migrated_doc = Document(str(output_path))
    all_texts = [p.text for p in migrated_doc.paragraphs]

    # Verify that 'Always wear PPE:' is not lost and is rendered as a list item
    assert any("Always wear PPE:" in t for t in all_texts)
    assert any("•\tAlways wear PPE:" in t for t in all_texts)

