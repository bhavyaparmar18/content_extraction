"""Unit tests for TableMigrator, TOCBuilder, and InstructionCleaner."""

import pytest
from pathlib import Path
from docx import Document
from docx.shared import RGBColor

from app.schemas.migration import (
    DocxMigrationOutput,
    MigrationSection,
    MigrationElement,
    MigrationTableCell,
)
from app.services.migration.table_migrator import TableMigrator
from app.services.migration.toc_builder import TOCBuilder
from app.services.migration.instruction_cleaner import InstructionCleaner
from app.services.migration.schemas import PlaceholderTablePlan, MigrationPlan, SectionPlan


def test_table_migrator_insert_and_populate(tmp_path):
    migrator = TableMigrator()
    doc = Document()

    # Insert a new table with merged cells
    elem = MigrationElement(
        element_type="table",
        num_rows=2,
        num_cols=2,
        cells=[
            MigrationTableCell(row_index=0, col_index=0, text="H1", is_header=True),
            MigrationTableCell(row_index=0, col_index=1, text="H2", is_header=True),
            MigrationTableCell(row_index=1, col_index=0, text="D1", col_span=2),
            MigrationTableCell(row_index=1, col_index=1, text="D2"),
        ],
    )
    t = migrator.insert_table(doc, elem, font_family="Arial")
    assert len(t.rows) == 2
    assert len(t.columns) == 2

    # Test populate placeholder
    plan = PlaceholderTablePlan(
        table_index=0,
        table_purpose="Definitions",
        parent_section_heading="3 DEFINITIONS",
        action="populate",
        source_section_title="3 DEFINITIONS",
        sort_alphabetically=True,
    )
    extracted = DocxMigrationOutput(
        document_id="test",
        sections=[
            MigrationSection(
                title="3 DEFINITIONS",
                elements=[elem],
            )
        ],
    )
    migrator.populate(doc, plan, extracted)
    assert len(t.rows) >= 2
    assert "D2" in t.rows[1].cells[0].text

    # Test delete table
    migrator.delete_table(doc, 0)
    assert len(doc.tables) == 0


def test_toc_builder_and_verify(tmp_path):
    toc_builder = TOCBuilder()
    doc = Document()

    doc.add_heading("1 PURPOSE", level=1)
    doc.add_paragraph("Body")
    doc.add_heading("2 APPLICABILITY", level=1)

    # Paragraph that looks like a heading but has Normal style
    bad_p = doc.add_paragraph("3 DEFINITIONS")
    bad_p.style = doc.styles["Normal"]

    toc_builder.insert_toc(doc)

    issues = toc_builder.verify_heading_styles(doc)
    assert len(issues) >= 1
    assert "3 DEFINITIONS" in issues[0]

    out_file = tmp_path / "toc_test.docx"
    doc.save(str(out_file))
    assert out_file.exists()


def test_instruction_cleaner(tmp_path):
    cleaner = InstructionCleaner()
    doc = Document()

    doc.add_heading("1 PURPOSE", level=1)
    p_black = doc.add_paragraph("This is black text.")
    p_blue = doc.add_paragraph()
    r_blue = p_blue.add_run("This is blue instruction text.")
    r_blue.font.color.rgb = RGBColor(0, 112, 192)

    doc.add_heading("9 DISTRIBUTION (OPTIONAL)", level=1)
    doc.add_paragraph("Optional info.")

    plan = MigrationPlan(
        template_name="test",
        section_plans=[
            SectionPlan(
                template_section_heading="1 PURPOSE",
                template_paragraph_indices_to_delete=[1],  # index of p_blue
            ),
            SectionPlan(
                template_section_heading="9 DISTRIBUTION (OPTIONAL)",
                has_source_content=False,
                fallback_action="delete_section",
            ),
        ],
    )

    cleaner.clean(doc, plan)

    # Ensure blue text is gone
    assert all("blue instruction" not in p.text for p in doc.paragraphs)
    # Ensure deleted section is gone
    assert all("DISTRIBUTION" not in p.text for p in doc.paragraphs)
