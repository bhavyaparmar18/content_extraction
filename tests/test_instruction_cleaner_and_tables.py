"""Tests for enhanced instruction cleaner and template table removal."""

import pytest
from pathlib import Path
from docx import Document
from docx.shared import RGBColor

from app.services.migration.instruction_cleaner import InstructionCleaner
from app.services.migration.schemas import MigrationPlan, SectionPlan, PlaceholderTablePlan
from app.services.migration.docx_migrator import DocxMigrator
from app.schemas.migration import DocxMigrationOutput, MigrationSection, MigrationElement


def test_instruction_cleaner_deletes_blue_instruction_tables():
    cleaner = InstructionCleaner()
    doc = Document()

    # Table 0: Vault token table (should be kept)
    t0 = doc.add_table(rows=1, cols=2)
    t0.rows[0].cells[0].text = "Scope"
    t0.rows[0].cells[1].text = "${vault:document_scope__c}"

    # Table 1: 1-row table with bright blue 0075FF instruction text (should be deleted)
    t1 = doc.add_table(rows=1, cols=2)
    t1.rows[0].cells[0].text = ""
    p_blue = t1.rows[0].cells[1].paragraphs[0]
    r_blue = p_blue.add_run("Brief description of what the document is about? What process is described in the document?")
    r_blue.font.color.rgb = RGBColor(0x00, 0x75, 0xFF)

    # Table 2: 1-row table with instruction keywords (should be deleted)
    t2 = doc.add_table(rows=1, cols=2)
    t2.rows[0].cells[0].text = ""
    t2.rows[0].cells[1].paragraphs[0].add_run(
        "Executive Summary/Introduction (short description of subject). It is provided with blue colour."
    )

    # Table 3: Normal content table with black text (should be kept)
    t3 = doc.add_table(rows=2, cols=2)
    t3.rows[0].cells[0].text = "Header 1"
    t3.rows[0].cells[1].text = "Header 2"
    t3.rows[1].cells[0].text = "Data 1"
    t3.rows[1].cells[1].text = "Data 2"

    assert len(doc.tables) == 4

    plan = MigrationPlan(template_name="test.docx")
    cleaner.clean(doc, plan)

    # Tables 1 and 2 should have been deleted; Tables 0 and 3 should remain
    assert len(doc.tables) == 2
    remaining_texts = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert any("${vault:document_scope__c}" in txt for txt in remaining_texts)
    assert any("Data 1" in txt for txt in remaining_texts)
    assert not any("Brief description of what" in txt for txt in remaining_texts)
    assert not any("Executive Summary/Introduction" in txt for txt in remaining_texts)


def test_instruction_cleaner_detects_0075ff_body_paragraphs():
    cleaner = InstructionCleaner()
    doc = Document()

    p1 = doc.add_paragraph()
    r1 = p1.add_run("Instructional text in 0075FF.")
    r1.font.color.rgb = RGBColor(0x00, 0x75, 0xFF)

    p2 = doc.add_paragraph()
    p2.add_run("Normal black content.")

    assert len(doc.paragraphs) == 2

    plan = MigrationPlan(template_name="test.docx")
    cleaner.clean(doc, plan)

    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0].text == "Normal black content."
