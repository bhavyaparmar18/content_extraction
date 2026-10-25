"""Phase T3: a language template's first page, headers and footers, extracted and swapped into a document."""

import hashlib
import io
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image

from app.api import translations
from app.api.auth import require_user
from app.config.settings import Settings, get_settings
from app.core.exception_handlers import register_exception_handlers
from app.schemas.v2 import FrontMatterReport
from app.services.translation.front_matter import (
    FrontMatterError,
    PREVIEW_TEXT,
    apply_front_matter,
    effective_refs,
    extract_front_matter,
    find_cover,
    front_matter_problems,
    section_props,
)
from tests.db_helpers import seed_users

ROOT = Path(__file__).resolve().parent.parent
REAL_TEMPLATE = ROOT / "data" / "template_output" / "Template_Main_GP_Docs_v1_normalized.docx"
REAL_FINAL = ROOT / "data" / "migrations" / "MIG-20261019-7BF76968" / "docx_final_v1.docx"
WORD_SCRIPT = ROOT / "scripts" / "word_resave.ps1"


def _png(rgb) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), rgb).save(buf, "PNG")
    return buf.getvalue()


LT_COVER, LT_FOOTER, LT_BODY = _png((200, 0, 0)), _png((0, 200, 0)), _png((0, 0, 200))
EN_COVER, EN_HEADER, EN_BODY = _png((10, 10, 10)), _png((20, 20, 20)), _png((30, 30, 30))


def _sha(data: bytes) -> str:
    return hashlib.sha1(data).hexdigest()


def _picture(paragraph, data: bytes) -> None:
    paragraph.add_run().add_picture(io.BytesIO(data), width=Inches(0.3))


def _style(doc, style_id: str, name: str) -> None:
    doc.styles.element.append(parse_xml(
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="{style_id}"><w:name w:val="{name}"/>'
        f'<w:basedOn w:val="Normal"/><w:rPr><w:b/></w:rPr></w:style>'))


def _set_style(paragraph, style_id: str) -> None:
    paragraph._p.get_or_add_pPr().style = style_id


def _hyperlink(paragraph, url: str, text: str) -> None:
    rId = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rId)
    run, t = OxmlElement("w:r"), OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _page_field(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def _numbered(doc, paragraph) -> str:
    """Give *paragraph* a list of its own; the list's numId."""
    numbering = doc.part.numbering_part.element
    numbering.insert(0, parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90"><w:lvl w:ilvl="0"><w:start w:val="1"/>'
        f'<w:numFmt w:val="upperRoman"/><w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="90"/></w:numPr>'))
    return "90"


def _end_first_page(doc, how: str) -> None:
    if how == "section":
        doc.add_section(WD_SECTION.NEW_PAGE)
    elif how == "page_break":
        doc.add_paragraph("Ende").add_run().add_break(WD_BREAK.PAGE)
    else:  # the next paragraph starts a new page
        p = doc.add_paragraph("Inhalt beginnt")
        p.paragraph_format.page_break_before = True


def language_template(path: Path, *, end: str = "section", even_odd: bool = False, numbered: bool = False,
                      cover_text: str = "Deckblatt Titel") -> Path:
    """A German template shaped like the company template: a cover section with three header and footer types,
    and a body section with its own headers that inherits the cover's footers."""
    doc = Document()
    _style(doc, "CoverTitel", "Cover title")
    _style(doc, "Kopfzeile", "header")  # a German Word's header style id; the name is the built-in one
    title = doc.add_paragraph(cover_text)
    _set_style(title, "CoverTitel")
    _picture(doc.add_paragraph(), LT_COVER)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Titel", "${vault:title__v}"
    if numbered:
        _numbered(doc, doc.add_paragraph("Punkt eins"))
    _end_first_page(doc, end)
    if end != "next_paragraph":
        doc.add_paragraph("Vorlagentext, der nicht übernommen wird")
    _picture(doc.add_paragraph(), LT_BODY)

    cover = doc.sections[0]
    cover.different_first_page_header_footer = True
    header = cover.header.paragraphs[0]
    header.text = "LT Kopf Standard "
    _set_style(header, "Kopfzeile")
    _page_field(header)
    _hyperlink(header, "https://example.com/de", "Link")
    cover.first_page_header.paragraphs[0].text = "LT Kopf erste Seite"
    footer = cover.first_page_footer.paragraphs[0]
    footer.text = "LT Fuss erste Seite"
    _picture(footer, LT_FOOTER)
    cover.footer.paragraphs[0].text = "LT Fuss"
    cover.even_page_header.paragraphs[0].text = "LT gerade Kopf"
    cover.even_page_footer.paragraphs[0].text = "LT gerade Fuss"
    if end == "section":
        body = doc.sections[1]
        body.header.is_linked_to_previous = False
        body.header.paragraphs[0].text = "LT Kopf Inhalt"
        body.different_first_page_header_footer = True
        body.first_page_header.is_linked_to_previous = False
        body.first_page_header.paragraphs[0].text = "LT Kopf Inhalt erste"
    if even_odd:
        doc.settings.odd_and_even_pages_header_footer = True
    doc.save(str(path))
    return path


def english(path: Path, *, end: str = "section", even_odd: bool = False) -> Path:
    """A rendered English document: cover, a body section, a landscape section and a final portrait section."""
    doc = Document()
    doc.add_paragraph("English cover title")
    _picture(doc.add_paragraph(), EN_COVER)
    _end_first_page(doc, end)
    doc.add_heading("PURPOSE", 1)
    doc.add_paragraph("English body text.")
    _picture(doc.add_paragraph(), EN_BODY)
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Landscape table text")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Final portrait text")
    landscape = doc.sections[-2]
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    first = doc.sections[0]
    first.header.paragraphs[0].text = "EN header"
    _picture(first.header.paragraphs[0], EN_HEADER)
    first.footer.paragraphs[0].text = "EN footer"
    if end == "section":
        doc.sections[1].header.is_linked_to_previous = False
        doc.sections[1].header.paragraphs[0].text = "EN body header"
    if even_odd:
        doc.settings.odd_and_even_pages_header_footer = True
    doc.save(str(path))
    return path


def _extract(tmp_path, template: Path, language: str = "de-DE") -> tuple[FrontMatterReport, Path, Path]:
    package, preview = tmp_path / "pkg.docx", tmp_path / "preview.docx"
    report = extract_front_matter(template, package, preview, template_id="lt", version=1, language=language)
    return report, package, preview


def _swap(tmp_path, template: Path, target: Path, language: str = "de-DE"):
    report, package, _ = _extract(tmp_path, template, language)
    assert report.ok, report.problems
    doc = Document(str(target))
    swap = apply_front_matter(doc, package, language)
    out = tmp_path / "swapped.docx"
    doc.save(str(out))
    return Document(str(out)), swap, package, out


def _body_text(doc) -> list[str]:
    return ["".join(t.text or "" for t in el.iter(qn("w:t"))) for el in doc.element.body]


def _image_hashes(part, root) -> set[str]:
    return {_sha(part.related_parts[b.get(qn("r:embed"))].blob) for b in root.iter(qn("a:blip"))}


def _all_parts_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist()
                        if n.startswith("word/header") or n.startswith("word/footer"))


# ── Finding the first page ────────────────────────────────────────────


@pytest.mark.parametrize("end, blocks, kind", [
    ("section", 4, "section_break"),        # title, picture, table, the section-break paragraph
    ("page_break", 4, "page_break"),        # ... and the paragraph holding the break
    ("next_paragraph", 3, "page_break"),    # the next paragraph starts a new page
])
def test_the_first_page_ends_at_the_first_break(tmp_path, end, blocks, kind):
    doc = Document(str(language_template(tmp_path / "lt.docx", end=end)))
    cover = find_cover(doc.element.body)
    assert (len(cover.blocks), cover.end) == (blocks, kind)
    assert "Deckblatt Titel" in "".join(t.text for t in cover.blocks[0].iter(qn("w:t")))
    assert (cover.next_block is not None) == (end == "next_paragraph")


def test_a_page_break_before_a_later_section_break_ends_the_first_page():
    doc = Document()
    doc.add_paragraph("Cover").add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Table of contents page")
    doc.add_section(WD_SECTION.NEW_PAGE)
    assert find_cover(doc.element.body).end == "page_break"
    assert len(find_cover(doc.element.body).blocks) == 1


def test_text_after_the_page_break_is_a_warning_and_no_break_is_an_error():
    doc = Document()
    p = doc.add_paragraph("Cover")
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("still here")
    assert "stays with the first page" in find_cover(doc.element.body).warnings[0]
    with pytest.raises(FrontMatterError, match="no first page"):
        find_cover(Document().element.body)


def test_a_continuous_section_break_cannot_end_the_first_page():
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_section(WD_SECTION.CONTINUOUS)
    with pytest.raises(FrontMatterError, match="continuous"):
        find_cover(doc.element.body)


def test_a_section_without_a_reference_inherits_the_previous_sections(tmp_path):
    doc = Document(str(language_template(tmp_path / "lt.docx")))
    cover, body = effective_refs(section_props(doc.element.body))
    assert body[("footer", "default")] == cover[("footer", "default")]  # footers are inherited
    assert body[("header", "default")] != cover[("header", "default")]  # headers are the body section's own
    assert set(body) == set(cover) == {(k, t) for k in ("header", "footer") for t in ("default", "first", "even")}


# ── Extraction ────────────────────────────────────────────────────────


def test_extraction_reports_the_front_matter_and_keeps_only_it(tmp_path):
    report, package, preview = _extract(tmp_path, language_template(tmp_path / "lt.docx"))
    assert report.ok and report.warnings == []
    assert (report.cover_end, report.cover_blocks, report.sections) == ("section_break", 4, 2)
    assert report.title_page_cover and report.title_page_body and not report.even_and_odd
    assert report.cover_text.startswith("Deckblatt Titel")
    assert {(p.kind, p.type) for p in report.body_parts} == {(p.kind, p.type) for p in report.cover_parts}
    body = {(p.kind, p.type): p for p in report.body_parts}
    assert body[("header", "default")].text == "LT Kopf Inhalt"
    assert body[("footer", "first")].images == 1
    assert {"CoverTitel", "Kopfzeile"} <= set(report.styles)
    assert report.images == 2  # the cover picture and the first-page footer's; not the body's

    doc = Document(str(package))
    text = " ".join(_body_text(doc))
    assert "Deckblatt Titel" in text and "Vorlagentext" not in text
    with zipfile.ZipFile(package) as z:
        media = {_sha(z.read(n)) for n in z.namelist() if n.startswith("word/media/")}
    assert media == {_sha(LT_COVER), _sha(LT_FOOTER)}
    sects = section_props(doc.element.body)
    assert len(sects) == 2 and len(sects[1].findall(qn("w:footerReference"))) == 3  # inherited, written out

    shown = " ".join(_body_text(Document(str(preview))))
    assert PREVIEW_TEXT[0] in shown and PREVIEW_TEXT[1] in shown


def test_a_template_without_a_page_break_has_a_problem_and_no_package(tmp_path):
    doc = Document()
    doc.add_paragraph("only text")
    doc.save(str(tmp_path / "flat.docx"))
    report, package, preview = _extract(tmp_path, tmp_path / "flat.docx")
    assert not report.ok and "no first page" in report.problems[0]
    assert not package.exists() and not preview.exists()


def test_objects_that_cannot_be_copied_and_tracked_changes_are_problems(tmp_path):
    path = language_template(tmp_path / "lt.docx")
    doc = Document(str(path))
    header = doc.sections[0].header.part
    ole = Part(PackURI("/word/embeddings/oleObject1.bin"), "application/vnd.openxmlformats-officedocument.oleObject",
               b"ole", doc.part.package)
    header.relate_to(ole, RT.OLE_OBJECT)
    first = doc.element.body[0]
    first.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="1" w:author="a"><w:r><w:t>neu</w:t></w:r></w:ins>'))
    doc.save(str(path))
    report, package, _ = _extract(tmp_path, path)
    assert any("oleObject" in p for p in report.problems)
    assert any("tracked changes" in p for p in report.problems)
    assert package.exists()  # the preview still shows what was found


def test_a_chinese_template_without_chinese_text_is_a_warning(tmp_path):
    report, _, _ = _extract(tmp_path, language_template(tmp_path / "lt.docx"), language="zh-CN")
    assert report.ok and any("Chinese or Japanese" in w for w in report.warnings)
    report, _, _ = _extract(tmp_path, language_template(tmp_path / "zh.docx", cover_text="封面标题"), language="zh-CN")
    assert not any("Chinese or Japanese" in w for w in report.warnings)


def test_comments_on_the_first_page_are_removed_with_a_warning(tmp_path):
    path = language_template(tmp_path / "lt.docx")
    doc = Document(str(path))
    p = doc.element.body[0]
    p.insert(1, parse_xml(f'<w:commentRangeStart {nsdecls("w")} w:id="0"/>'))
    p.append(parse_xml(f'<w:commentRangeEnd {nsdecls("w")} w:id="0"/>'))
    p.append(parse_xml(f'<w:r {nsdecls("w")}><w:commentReference w:id="0"/></w:r>'))
    doc.save(str(path))
    report, package, _ = _extract(tmp_path, path)
    assert report.ok and any("comments" in w for w in report.warnings)
    assert not list(Document(str(package)).element.body.iter(qn("w:commentReference")))


# ── The swap ──────────────────────────────────────────────────────────


def test_the_swap_replaces_the_first_page_headers_and_footers(tmp_path):
    target = english(tmp_path / "en.docx")
    before = Document(str(target))
    english_body = _body_text(before)[3:]  # after the cover (title, picture, section break)
    doc, swap, package, out = _swap(tmp_path, language_template(tmp_path / "lt.docx"), target)

    text = _body_text(doc)
    assert text[0] == "Deckblatt Titel" and "English cover title" not in " ".join(text)
    assert text[4:] == english_body  # the body is untouched and in order
    assert (swap.blocks_removed, swap.blocks_inserted, swap.sections_rewired) == (3, 4, 4)

    body = doc.element.body
    cover, rest = list(body)[:4], list(body)[4:]
    assert _image_hashes(doc.part, cover[1]) == {_sha(LT_COVER)}
    assert {_sha(EN_BODY)} == set().union(*(_image_hashes(doc.part, el) for el in rest))

    sections = doc.sections
    assert sections[0].first_page_header.paragraphs[0].text == "LT Kopf erste Seite"
    assert sections[0].header.paragraphs[0].text.startswith("LT Kopf Standard")
    assert all(s.header.paragraphs[0].text == "LT Kopf Inhalt" for s in sections[1:])
    assert all(s.first_page_header.paragraphs[0].text == "LT Kopf Inhalt erste" for s in sections[1:])
    assert all(s.footer.paragraphs[0].text == "LT Fuss" for s in sections)
    assert all(s.different_first_page_header_footer for s in sections)
    footer = sections[0].first_page_footer
    assert _image_hashes(footer.part, footer._element) == {_sha(LT_FOOTER)}
    link = sections[0].header.part.rels
    assert any(r.is_external and r.target_ref == "https://example.com/de" for r in link.values())
    assert sections[2].orientation == WD_ORIENT.LANDSCAPE  # the English page setup is kept
    assert "EN header" not in _all_parts_text(out) and "EN body header" not in _all_parts_text(out)
    assert len(swap.parts_dropped) == 3 and len(swap.parts_added) == 8

    styles = {s.get(qn("w:styleId")): s for s in doc.styles.element.findall(qn("w:style"))}
    assert styles["Kopfzeile"].find(qn("w:name")).get(qn("w:val")) == "header (de-DE)"
    assert "CoverTitel" in styles and set(swap.styles_added) == {"CoverTitel", "Kopfzeile"}

    ids = [d.get("id") for part in [doc.part] + [s.header.part for s in sections[:1]] for d in
           part.element.iter(qn("wp:docPr"))]
    assert len(ids) == len(set(ids))
    assert front_matter_problems(doc, package) == []


def test_the_check_names_what_differs_on_an_unswapped_document(tmp_path):
    _, package, _ = _extract(tmp_path, language_template(tmp_path / "lt.docx"))
    problems = front_matter_problems(Document(str(english(tmp_path / "en.docx"))), package)
    assert any("first page differs" in p for p in problems)
    assert any("section 2 (body)" in p for p in problems)


@pytest.mark.parametrize("lt_end, en_end, sections", [
    ("page_break", "section", 3),        # the English cover section goes; the cover joins the first body section
    ("section", "next_paragraph", 4),    # the language cover brings its section
])
def test_first_pages_ending_differently_swap_cleanly(tmp_path, lt_end, en_end, sections):
    doc, swap, package, _ = _swap(tmp_path, language_template(tmp_path / "lt.docx", end=lt_end),
                                  english(tmp_path / "en.docx", end=en_end))
    assert len(doc.sections) == sections
    assert front_matter_problems(doc, package) == []
    assert "English cover title" not in " ".join(_body_text(doc))
    if en_end == "next_paragraph":  # the language cover ends the page itself: no blank page
        assert not list(doc.element.body.iter(qn("w:pageBreakBefore")))
    if lt_end == "page_break":
        assert all(s.first_page_header.paragraphs[0].text == "LT Kopf erste Seite" for s in doc.sections)


def test_a_cover_ended_by_the_next_paragraph_carries_its_own_page_break(tmp_path):
    doc, _, package, _ = _swap(tmp_path, language_template(tmp_path / "lt.docx", end="next_paragraph"),
                               english(tmp_path / "en.docx"))
    cover = find_cover(doc.element.body)
    assert cover.end == "page_break" and list(cover.blocks[-1].iter(qn("w:br")))
    assert "Inhalt beginnt" not in " ".join(_body_text(doc))
    assert front_matter_problems(doc, package) == []


@pytest.mark.parametrize("lt_even, en_even", [(True, False), (False, True)])
def test_even_and_odd_pages_follow_the_language_template(tmp_path, lt_even, en_even):
    doc, _, package, _ = _swap(tmp_path, language_template(tmp_path / "lt.docx", even_odd=lt_even),
                               english(tmp_path / "en.docx", even_odd=en_even))
    assert doc.settings.odd_and_even_pages_header_footer is lt_even
    assert front_matter_problems(doc, package) == []


def test_a_list_on_the_first_page_gets_a_numbering_definition_of_its_own(tmp_path):
    doc, _, package, _ = _swap(tmp_path, language_template(tmp_path / "lt.docx", numbered=True),
                               english(tmp_path / "en.docx"))
    numbering = doc.part.numbering_part.element
    known = {n.get(qn("w:numId")): n for n in numbering.findall(qn("w:num"))}
    used = [n.get(qn("w:val")) for n in doc.element.body.iter(qn("w:numId"))]
    assert used and all(u in known for u in used)
    abstract = known[used[0]].find(qn("w:abstractNumId")).get(qn("w:val"))
    lvl = next(a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract)
    assert lvl.find(f"{qn('w:lvl')}/{qn('w:numFmt')}").get(qn("w:val")) == "upperRoman"
    assert front_matter_problems(doc, package) == []


def test_bookmarks_on_the_language_first_page_are_renumbered_or_dropped(tmp_path):
    path = language_template(tmp_path / "lt.docx")
    doc = Document(str(path))
    title = doc.element.body[0]
    for name, i in (("_Cover", 0), ("PURPOSE_MARK", 1)):
        title.insert(1, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{i}" w:name="{name}"/>'))
        title.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{i}"/>'))
    doc.save(str(path))
    target = english(tmp_path / "en.docx")
    en = Document(str(target))
    heading = en.element.body[3]
    heading.insert(1, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="0" w:name="PURPOSE_MARK"/>'))
    heading.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="0"/>'))
    en.save(str(target))

    doc, swap, package, _ = _swap(tmp_path, path, target)
    starts = list(doc.element.body.iter(qn("w:bookmarkStart")))
    assert sorted(b.get(qn("w:name")) for b in starts) == ["PURPOSE_MARK", "_Cover"]
    assert len({b.get(qn("w:id")) for b in starts}) == 2
    assert any("PURPOSE_MARK" in w for w in swap.warnings)


def test_a_document_without_a_first_page_cannot_take_the_front_matter(tmp_path):
    _, package, _ = _extract(tmp_path, language_template(tmp_path / "lt.docx"))
    doc = Document()
    doc.add_paragraph("no breaks")
    with pytest.raises(FrontMatterError):
        apply_front_matter(doc, package, "de-DE")


# ── The real company template ─────────────────────────────────────────


@pytest.mark.skipif(not (REAL_TEMPLATE.exists() and REAL_FINAL.exists()), reason="sample files not available")
def test_the_company_template_front_matter_swaps_into_a_rendered_migration(tmp_path):
    report, package, _ = _extract(tmp_path, REAL_TEMPLATE)
    assert report.ok and report.cover_end == "section_break" and report.sections == 2
    doc = Document(str(REAL_FINAL))
    body_before = _body_text(doc)[6:]
    apply_front_matter(doc, package, "de-DE")
    doc.save(str(tmp_path / "swapped.docx"))
    doc = Document(str(tmp_path / "swapped.docx"))
    assert _body_text(doc)[7:] == body_before
    assert front_matter_problems(doc, package) == []


@pytest.mark.word
@pytest.mark.skipif(os.environ.get("SOP_RUN_WORD_TESTS") != "1", reason="set SOP_RUN_WORD_TESTS=1 to run")
@pytest.mark.skipif(sys.platform != "win32" or shutil.which("powershell") is None, reason="needs Windows")
def test_the_swapped_document_opens_in_word(tmp_path):
    """Word rewrites the parts on saving, so the texts are compared, not the bytes."""
    doc, _, _, out = _swap(tmp_path, language_template(tmp_path / "lt.docx"), english(tmp_path / "en.docx"))
    resaved = tmp_path / "resaved.docx"
    result = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(WORD_SCRIPT),
         "-Source", str(out), "-Target", str(resaved)],
        capture_output=True, text=True, timeout=180,
    )
    assert result.returncode == 0 and resaved.exists(), result.stdout + result.stderr
    after = Document(str(resaved))
    assert [s.header.paragraphs[0].text for s in after.sections] == [s.header.paragraphs[0].text for s in doc.sections]
    assert after.sections[0].first_page_header.paragraphs[0].text == "LT Kopf erste Seite"
    assert _body_text(after)[0] == "Deckblatt Titel"


# ── API ───────────────────────────────────────────────────────────────


@pytest.fixture
def settings(tmp_path, monkeypatch):
    monkeypatch.setenv("SOP_DB_PATH", str(tmp_path / "data" / "app.db"))
    s = Settings(project_root=tmp_path)
    s.resolve_paths(tmp_path)
    s.ensure_directories()
    return s


@pytest.fixture
def client(settings):
    seed_users(settings.db_path, "alice")
    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(translations.router)
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[require_user] = lambda: {"userId": "alice"}
    with TestClient(app) as c:
        yield c


API = "/api/v1/translations/language-templates"
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _upload(client, data: bytes, language="de-DE", filename="Vorlage DE.docx", **form):
    return client.post(API, data={"language": language, **form}, files={"file": (filename, data, _DOCX)})


def test_upload_extracts_and_an_admin_marks_it_ready(client, settings, tmp_path):
    data = language_template(tmp_path / "lt.docx").read_bytes()
    created = _upload(client, data, name="German template")
    assert created.status_code == 200, created.text
    record = created.json()
    assert record["status"] == "extracted" and record["version"] == 1 and record["language"] == "de-DE"
    assert record["template_id"] == "German_template" and record["created_by"] == "alice"
    assert record["report"]["cover_end"] == "section_break" and record["report"]["problems"] == []
    assert Path(record["upload_path"]).read_bytes() == data

    preview = client.get(f"{API}/German_template/preview")
    assert preview.status_code == 200 and preview.headers["content-type"] == _DOCX
    assert PREVIEW_TEXT[0] in " ".join(_body_text(Document(io.BytesIO(preview.content))))
    assert client.get(f"{API}/German_template/report").json()["cover_blocks"] == 4

    ready = client.post(f"{API}/German_template/ready")
    assert ready.status_code == 200 and ready.json()["status"] == "ready"
    assert [r["version"] for r in client.get(API, params={"language": "de-DE", "status": "ready"}).json()] == [1]
    assert client.get(API, params={"language": "fr-FR"}).json() == []

    again = _upload(client, data, template_id="German_template").json()
    assert again["version"] == 2 and again["status"] == "extracted"
    assert client.get(f"{API}/German_template").json()["version"] == 2
    assert client.get(f"{API}/German_template", params={"version": 1}).json()["status"] == "ready"


def test_a_template_with_problems_cannot_be_marked_ready(client, tmp_path):
    path = language_template(tmp_path / "lt.docx")
    doc = Document(str(path))
    doc.element.body[0].append(parse_xml(f'<w:del {nsdecls("w")} w:id="1" w:author="a"><w:r><w:delText>x'
                                         f'</w:delText></w:r></w:del>'))
    doc.save(str(path))
    record = _upload(client, path.read_bytes()).json()
    assert record["status"] == "extracted" and record["report"]["problems"]
    refused = client.post(f"{API}/{record['id']}/ready")
    assert refused.status_code == 409 and refused.json()["error"]["code"] == "LANGUAGE_TEMPLATE_NOT_READY"
    assert "tracked changes" in refused.json()["error"]["fieldErrors"][0]["message"]


def test_a_file_that_is_not_usable_is_saved_as_failed(client, tmp_path):
    broken = _upload(client, b"not a zip", filename="broken.docx").json()
    assert broken["status"] == "failed" and "cannot be opened" in broken["error"]
    flat = Document()
    flat.add_paragraph("no first page")
    buf = io.BytesIO()
    flat.save(buf)
    failed = _upload(client, buf.getvalue(), filename="flat.docx").json()
    assert failed["status"] == "failed" and "no first page" in failed["error"]
    assert client.post(f"{API}/{failed['id']}/ready").status_code == 409
    assert client.get(f"{API}/{failed['id']}/preview").status_code == 404


def test_upload_errors(client, tmp_path):
    data = language_template(tmp_path / "lt.docx").read_bytes()
    pdf = _upload(client, b"%PDF", filename="x.pdf")
    assert pdf.status_code == 400 and pdf.json()["error"]["code"] == "UNSUPPORTED_LANGUAGE_TEMPLATE_FILE_TYPE"
    assert _upload(client, b"", filename="x.docx").json()["error"]["code"] == "EMPTY_LANGUAGE_TEMPLATE_FILE"
    assert _upload(client, data, language="xx").json()["error"]["code"] == "UNSUPPORTED_LANGUAGE"
    assert _upload(client, data, template_id="T").status_code == 200
    clash = _upload(client, data, language="fr-FR", template_id="T")
    assert clash.status_code == 409 and clash.json()["error"]["code"] == "LANGUAGE_TEMPLATE_LANGUAGE_MISMATCH"
    assert client.get(f"{API}/nope").json()["error"]["code"] == "LANGUAGE_TEMPLATE_NOT_FOUND"


def test_delete_removes_the_version_and_its_files(client, tmp_path):
    record = _upload(client, language_template(tmp_path / "lt.docx").read_bytes()).json()
    files = [Path(record[k]) for k in ("upload_path", "front_matter_path", "preview_path", "report_path")]
    assert all(f.exists() for f in files)
    deleted = client.delete(f"{API}/{record['id']}")
    assert deleted.status_code == 200 and deleted.json()["deleted"] is True
    assert not any(f.exists() for f in files)
    assert client.get(f"{API}/{record['id']}").status_code == 404
