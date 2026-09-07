"""Main orchestrator for the SOP Content-to-Template DOCX Migration Engine."""

from __future__ import annotations

import asyncio
import re
from pathlib import Path
from typing import Optional, Any
from docx import Document
from docx.shared import Pt
from loguru import logger

from app.config.settings import Settings
from app.schemas.migration import DocxMigrationOutput, MigrationElement, MigrationSection
from app.services.llm.chain_factory import ChainFactory
from app.services.llm.rate_limiter import LLMRateLimiter
from app.services.migration.schemas import (
    MigrationPlan,
    SectionPlan,
    ElementPlacement,
    PlaceholderTablePlan,
    CalloutStyleDef,
    MigrationResult,
)
from app.services.migration.template_inspector import TemplateInspector
from app.services.migration.content_summarizer import ContentSummarizer
from app.services.migration.llm_section_summarizer import LLMSectionSummarizer
from app.services.migration.section_aligner import SectionAligner
from app.services.migration.docx_styler import DocxStyler
from app.services.migration.table_migrator import TableMigrator
from app.services.migration.callout_builder import CalloutBuilder
from app.services.migration.toc_builder import TOCBuilder
from app.services.migration.instruction_cleaner import InstructionCleaner
from app.services.migration.migration_validator import MigrationValidator


class DocxMigrator:
    """Orchestrates the full 5-phase migration pipeline."""

    def __init__(self, settings: Settings, chain_factory: ChainFactory):
        self.settings = settings
        self.chain_factory = chain_factory
        self.inspector = TemplateInspector()
        self.toc_builder = TOCBuilder()
        self.styler = DocxStyler()
        self.table_migrator = TableMigrator()
        self.callout_builder = CalloutBuilder()
        self.cleaner = InstructionCleaner()
        self.validator = MigrationValidator()

        # Shared rate limiter
        self.rate_limiter = LLMRateLimiter(
            max_concurrent=settings.llm_max_concurrent,
            min_delay_seconds=settings.llm_min_delay_seconds,
            max_retries=settings.llm_max_retries,
            base_backoff_seconds=settings.llm_base_backoff_seconds,
        )

        # Mode A or Mode B content summarizer
        if settings.use_llm_section_summarizer:
            self.summarizer = LLMSectionSummarizer(chain_factory, self.rate_limiter)
        else:
            self.summarizer = ContentSummarizer()

        self.aligner = SectionAligner(chain_factory, self.rate_limiter)

    async def migrate(
        self,
        extracted: DocxMigrationOutput,
        template_path: Path | str,
        output_path: Path | str,
    ) -> MigrationResult:
        """Execute full end-to-end migration into the target template."""
        t_path = Path(template_path)
        o_path = Path(output_path)
        o_path.parent.mkdir(parents=True, exist_ok=True)

        logger.info(f"Starting migration: doc '{extracted.document_id}' into template '{t_path.name}'")

        # ── Phase 1: Pre-process ─────────────────────────────────
        logger.info("Phase 1: Inspecting template structure")
        template_profile = self.inspector.inspect(t_path)

        logger.info(f"Phase 1: Summarizing content (LLM Mode={self.settings.use_llm_section_summarizer})")
        if asyncio.iscoroutinefunction(getattr(self.summarizer, "summarize", None)):
            content_summary = await self.summarizer.summarize(
                extracted, skip_preamble=self.settings.skip_preamble_migration
            )
        else:
            content_summary = self.summarizer.summarize(
                extracted, skip_preamble=self.settings.skip_preamble_migration
            )

        # ── Phase 2: LLM Migration Planning ──────────────────────
        logger.info("Phase 2: Calling LLM Migration Planner")
        plan = await self.aligner.create_migration_plan(template_profile, content_summary)
        logger.info(
            f"Phase 2: Plan created — {len(plan.section_plans)} sections, "
            f"confidence={plan.overall_confidence:.2f}, warnings={len(plan.warnings)}"
        )

        # ── Phase 3: Programmatic Execution ──────────────────────
        logger.info("Phase 3: Building output .docx from template")
        doc = Document(str(t_path))

        # 1. Clean blue instructions and deleted sections from template
        # (Table cleaning is deferred until after placeholder tables are populated by index)
        self.cleaner.clean(doc, plan, clean_tables=False)

        # 2. Self-healing check for any shifted icons in extracted sections
        self._sanitize_shifted_icons(extracted)

        # 3. Locate template's real Document History table and vault token tables
        doc_hist_table = next(
            (t for t in template_profile.tables if getattr(t, "is_document_history_table", False)),
            None
        )
        vault_table_indices = {
            t.index for t in template_profile.tables if getattr(t, "is_vault_token_table", False)
        }

        # Populate placeholder tables
        populated_table_indices = set()
        self._populated_tbl_elements: dict[tuple[str, int], Any] = {}
        for tp in plan.placeholder_tables:
            if tp.action == "populate":
                # Check if Document History was mapped to the wrong table (e.g. cover page vault tables)
                if tp.source_section_title and "document history" in tp.source_section_title.lower():
                    if doc_hist_table and tp.table_index != doc_hist_table.index:
                        logger.info(
                            f"Redirecting Document History from table {tp.table_index} to "
                            f"real Document History table {doc_hist_table.index}"
                        )
                        tp.table_index = doc_hist_table.index
                        if doc_hist_table.index in plan.tables_to_delete:
                            plan.tables_to_delete.remove(doc_hist_table.index)

                # Protect cover-page vault tables from being populated
                if tp.table_index in vault_table_indices:
                    logger.warning(
                        f"Skipping population of cover-page vault table {tp.table_index} "
                        f"('{tp.table_purpose}')"
                    )
                    continue

                if tp.table_index in populated_table_indices:
                    continue

                self.table_migrator.populate(doc, tp, extracted)
                populated_table_indices.add(tp.table_index)
                if tp.table_index < len(doc.tables):
                    self._populated_tbl_elements[(tp.source_section_title, tp.source_element_index)] = doc.tables[tp.table_index]._element

        # If Document History is present in source but wasn't populated, auto-populate it
        if doc_hist_table and not any(
            tp.table_index == doc_hist_table.index and tp.action == "populate"
            for tp in plan.placeholder_tables
        ):
            hist_sec = next((s for s in extracted.sections if "document history" in s.title.lower()), None)
            if hist_sec:
                hist_elem_idx = next(
                    (i for i, el in enumerate(hist_sec.elements) if el.element_type == "table"),
                    None
                )
                if hist_elem_idx is not None:
                    logger.info(f"Auto-populating real Document History table {doc_hist_table.index}")
                    tp_hist = PlaceholderTablePlan(
                        table_index=doc_hist_table.index,
                        table_purpose="Document History",
                        parent_section_heading="DOCUMENT HISTORY",
                        action="populate",
                        source_section_title=hist_sec.title,
                        source_element_index=hist_elem_idx,
                    )
                    if doc_hist_table.index in plan.tables_to_delete:
                        plan.tables_to_delete.remove(doc_hist_table.index)
                    self.table_migrator.populate(doc, tp_hist, extracted)
                    if doc_hist_table.index < len(doc.tables):
                        self._populated_tbl_elements[(tp_hist.source_section_title, tp_hist.source_element_index)] = doc.tables[doc_hist_table.index]._element

        # Delete unused template tables
        delete_indices = set(plan.tables_to_delete)
        for tp in plan.placeholder_tables:
            if tp.action == "delete":
                delete_indices.add(tp.table_index)

        # Do NOT delete cover page vault tables or real Document History table
        delete_indices -= vault_table_indices
        if doc_hist_table:
            delete_indices.discard(doc_hist_table.index)

        populated_indices = {tp.table_index for tp in plan.placeholder_tables if tp.action == "populate"}
        if doc_hist_table:
            populated_indices.add(doc_hist_table.index)

        for t_info in template_profile.tables:
            # Every template table that is not a cover-page vault table
            # and was not populated with source data must be deleted.
            if t_info.index not in populated_indices and t_info.index not in vault_table_indices:
                delete_indices.add(t_info.index)

        for idx in sorted(delete_indices, reverse=True):
            self.table_migrator.delete_table(doc, idx)

        # Ensure all template sections from template inspection exist in plan
        existing_template_headings = {
            sp.template_section_heading.strip().lower() for sp in plan.section_plans
        }
        for tp in template_profile.paragraphs:
            if tp.heading_level is not None and tp.text.strip():
                norm_th = tp.text.strip().lower()
                if norm_th not in existing_template_headings:
                    new_sp = SectionPlan(
                        template_section_heading=tp.text.strip(),
                        template_heading_level=tp.heading_level,
                        source_sections_mapped=[],
                        elements=[],
                        has_source_content=False,
                        is_unmapped_source=False,
                        fallback_action=None,
                    )
                    plan.section_plans.append(new_sp)
                    existing_template_headings.add(norm_th)

        # Auto-align any matching template sections that the LLM left unmapped
        self._align_unmapped_template_sections(plan, extracted)

        # Consolidate any subsection plans that were mistakenly created as top-level unmapped sections
        merged_sp_indices = set()
        for idx, sp in enumerate(plan.section_plans):
            if sp.is_unmapped_source:
                m_sub = re.match(r"^(\d+)\.(\d+)", sp.template_section_heading.strip())
                if m_sub:
                    parent_num = m_sub.group(1)
                    parent_sp = next(
                        (other for other in plan.section_plans if other != sp and any(re.match(rf"^{parent_num}(?:\s|$)", s.strip()) for s in other.source_sections_mapped)),
                        None
                    )
                    if parent_sp is not None:
                        logger.info(f"Consolidating unmapped subsection plan '{sp.template_section_heading}' into parent plan '{parent_sp.template_section_heading}'")
                        for src in sp.source_sections_mapped:
                            if src not in parent_sp.source_sections_mapped:
                                parent_sp.source_sections_mapped.append(src)
                        parent_sp.elements.extend(sp.elements)
                        merged_sp_indices.add(idx)

        if merged_sp_indices:
            plan.section_plans = [sp for i, sp in enumerate(plan.section_plans) if i not in merged_sp_indices]

        # Process mapped sections
        mapped_plans = [sp for sp in plan.section_plans if not sp.is_unmapped_source]
        unmapped_plans = [sp for sp in plan.section_plans if sp.is_unmapped_source]

        # Detect any remaining source sections not in any section plan
        mapped_source_titles = set()
        for sp in plan.section_plans:
            mapped_source_titles.update(sp.source_sections_mapped)
            for p in sp.elements:
                if p.source_section_title:
                    mapped_source_titles.add(p.source_section_title)

        for sec in extracted.sections:
            if self.settings.skip_preamble_migration and sec.title.strip().startswith("0 "):
                continue
            if sec.title not in mapped_source_titles:
                # If this unmapped section is a subsection of an existing section, merge it into parent
                m_sub = re.match(r"^(\d+)\.(\d+)", sec.title.strip())
                if m_sub:
                    parent_num = m_sub.group(1)
                    parent_sp = next(
                        (sp for sp in plan.section_plans if any(re.match(rf"^{parent_num}(?:\s|$)", s.strip()) for s in sp.source_sections_mapped)),
                        None
                    )
                    if parent_sp is not None:
                        logger.info(f"Subsection '{sec.title}' belongs to parent section '{parent_sp.template_section_heading}', merging source content.")
                        parent_sp.source_sections_mapped.append(sec.title)
                        mapped_source_titles.add(sec.title)
                        continue

                logger.info(f"Auto-inserting unmapped source section: '{sec.title}'")
                unmapped_sp = SectionPlan(
                    template_section_heading=sec.title,
                    template_heading_level=1,
                    source_sections_mapped=[sec.title],
                    elements=[],
                    has_source_content=True,
                    is_unmapped_source=True,
                    insertion_after_section="PROCESS",
                )
                unmapped_plans.append(unmapped_sp)
                plan.section_plans.append(unmapped_sp)

        for sp in mapped_plans:
            self._execute_section_plan(doc, sp, extracted, plan)

        # Process unmapped sections sequentially
        running_anchor = None
        for sp in unmapped_plans:
            running_anchor = self._insert_unmapped_section(
                doc, sp, extracted, plan, anchor=running_anchor
            )

        # ── Final Scaffolding Cleanup Pass ───────────────────────
        logger.info("Phase 3: Running final cleanup pass on remaining template scaffolding")
        protected_tbl_elements = set(getattr(self, "_populated_tbl_elements", {}).values())
        self.cleaner.clean(doc, plan, protected_tbl_elements=protected_tbl_elements)

        # ── Phase 4: Insert TOC & Polish Formatting ──────────────
        logger.info("Phase 4: Inserting Table of Contents field")
        self.toc_builder.insert_toc(doc, plan)

        # Ensure no two tables are directly adjacent in the DOM (prevents Word from auto-joining them)
        self._separate_adjacent_tables(doc)

        # Save document
        doc.save(str(o_path))
        logger.info(f"Phase 3-4: Saved migrated document to '{o_path}'")

        # ── Phase 5: Validate ────────────────────────────────────
        logger.info("Phase 5: Validating migrated document")
        qa_report = self.validator.validate(o_path, extracted, plan)
        qa_report.summarizer_mode = content_summary.summarizer_mode
        logger.info(f"Phase 5: QA validation complete — status='{qa_report.status}'")

        return MigrationResult(
            output_path=str(o_path),
            plan=plan,
            qa_report=qa_report,
        )

    def _align_unmapped_template_sections(
        self, plan: MigrationPlan, extracted: DocxMigrationOutput
    ):
        """Map source sections directly to matching template headings if LLM left them unmapped."""
        import re
        for sp in plan.section_plans:
            if not sp.source_sections_mapped or not sp.has_source_content:
                norm_template = re.sub(r"^\d+[\.\s]*", "", sp.template_section_heading.strip().lower()).strip()
                if not norm_template or norm_template == "general information":
                    continue
                for sec in extracted.sections:
                    if self.settings.skip_preamble_migration and sec.title.strip().startswith("0 "):
                        continue
                    # Decimal subsections (e.g. 7.1) must not match top-level template sections
                    if re.match(r"^\d+\.\d+", sec.title.strip()):
                        continue
                    norm_src = re.sub(r"^\d+[\.\s]*", "", sec.title.strip().lower()).strip()
                    if norm_template and (norm_src == norm_template or norm_template in norm_src or norm_src in norm_template):
                        already_claimed = any(sec.title in other.source_sections_mapped for other in plan.section_plans if other != sp)
                        if not already_claimed:
                            logger.info(f"Auto-aligning source section '{sec.title}' to template section '{sp.template_section_heading}'")
                            sp.source_sections_mapped.append(sec.title)
                            sp.has_source_content = True
                            sp.fallback_action = None
                            break

    def _execute_section_plan(
        self,
        doc: Document,
        section_plan: SectionPlan,
        extracted: DocxMigrationOutput,
        plan: MigrationPlan,
    ):
        """Execute one SectionPlan for a mapped section, guaranteeing all source elements are placed."""
        anchor = self._find_heading_anchor(doc, section_plan.template_section_heading)

        if not section_plan.has_source_content or not section_plan.source_sections_mapped:
            if not section_plan.elements:
                if section_plan.fallback_action == "insert_none":
                    self._insert_fallback_text(doc, anchor, "(None)")
                elif section_plan.fallback_action == "insert_na":
                    self._insert_fallback_text(doc, anchor, "N/A")
                return

        # Collect all source elements that belong to this section plan
        all_placements = self._build_complete_placements(section_plan, extracted)

        for placement in all_placements:
            new_oxml = self._execute_element(doc, placement, extracted, plan, anchor=anchor, section_plan=section_plan)
            if new_oxml is not None:
                anchor = new_oxml

    @staticmethod
    def _clean_heading_number_prefix(text: str) -> str:
        """Strip leading numbering prefix (e.g. '6.1 ', '6.5.1 ', '1. ', '6 ') from heading text."""
        if not text:
            return ""
        stripped = text.strip()
        # 1. Decimal subsections: e.g. '6.1 ', '6.1.1 ', '6.1 - ', '6.1: '
        cleaned = re.sub(r"^\d{1,2}(?:\.\d{1,2})+[.)\-:\s]*\s*", "", stripped).strip()
        if cleaned and cleaned != stripped:
            return cleaned
        # 2. Numbered with dot: e.g. '1. ', '6. '
        cleaned = re.sub(r"^\d{1,2}\.[)\-:\s]*\s*", "", stripped).strip()
        if cleaned and cleaned != stripped:
            return cleaned
        # 3. Major section number followed by space and capital letter: e.g. '6 PRINCIPLES'
        cleaned = re.sub(r"^\d{1,2}\s+(?=[A-Z])", "", stripped).strip()
        if cleaned and cleaned != stripped:
            return cleaned
        return stripped

    def _insert_unmapped_section(
        self,
        doc: Document,
        section_plan: SectionPlan,
        extracted: DocxMigrationOutput,
        plan: MigrationPlan,
        anchor: Optional[Any] = None,
    ) -> Optional[Any]:
        """Insert an unmapped source section at a logical contextual position."""
        if anchor is None:
            if section_plan.insertion_after_section:
                anchor = self._find_heading_anchor(doc, section_plan.insertion_after_section)
            if anchor is None and doc.paragraphs:
                anchor = doc.paragraphs[-1]._element

        lvl = section_plan.template_heading_level or 1
        font_size = getattr(plan, f"font_size_heading{lvl}_pt", 14.0)

        heading_text = self._clean_heading_number_prefix(section_plan.template_section_heading)
        heading = self.styler.insert_heading(
            doc=doc,
            text=heading_text,
            level=lvl,
            style_name=f"Heading {lvl}",
            font_family=plan.font_family,
            font_size_pt=font_size,
        )

        if anchor is not None and heading._element != anchor:
            anchor.addnext(heading._element)
            anchor = heading._element

        all_placements = self._build_complete_placements(section_plan, extracted)

        for placement in all_placements:
            new_oxml = self._execute_element(doc, placement, extracted, plan, anchor=anchor, section_plan=section_plan)
            if new_oxml is not None:
                anchor = new_oxml

        return anchor

    def _build_complete_placements(
        self,
        section_plan: SectionPlan,
        extracted: DocxMigrationOutput,
    ) -> list[ElementPlacement]:
        """Ensure 100% of source elements from mapped sections are scheduled for placement."""
        import re
        placements: list[ElementPlacement] = []
        explicit_by_sec_and_idx: dict[tuple[str, int], ElementPlacement] = {}

        for p in section_plan.elements:
            explicit_by_sec_and_idx[(p.source_section_title, p.source_element_index)] = p

        # Iterate over all mapped source sections in order
        for src_title in section_plan.source_sections_mapped:
            src_sec = next((s for s in extracted.sections if s.title == src_title), None)
            if src_sec is None:
                continue

            # Normalise source section title for comparison (strip leading numbers)
            norm_src_title = re.sub(r"^\d+[\.\ ]*", "", src_title.strip()).strip().lower()

            for idx, elem in enumerate(src_sec.elements):
                if (src_title, idx) in explicit_by_sec_and_idx:
                    explicit_p = explicit_by_sec_and_idx[(src_title, idx)]
                    # Still honour the skip for explicit section-title headings
                    if explicit_p.action == "skip":
                        continue
                    if elem.icons:
                        explicit_p.embed_icons_inline = True
                    placements.append(explicit_p)
                    continue

                action = "insert_paragraph"
                if elem.element_type == "heading":
                    # If heading text looks like a list item (bullet or numbered item ending with ':'), reclassify as list
                    h_text = (elem.text or "").strip()
                    if (
                        re.match(r"^[\u2022\u2023\u25E6\u2043\u2219\u25AA\u25AB\u25CF\u25CB\u25A0\u25A1\u2013\u2014○●◆◇■□▪▫–—•‣⁃]", h_text)
                        or (re.match(r"^(?:\d{1,3}[.)]\s|[a-zA-Z][.)]\s)", h_text) and h_text.endswith(":"))
                    ):
                        action = "insert_list"
                    else:
                        action = "insert_heading"
                elif elem.element_type == "list":
                    action = "insert_list"
                elif elem.element_type == "table":
                    if (src_title, idx) in getattr(self, "_populated_tbl_elements", {}):
                        action = "populate_placeholder"
                    else:
                        action = "insert_table"
                elif elem.element_type == "image":
                    action = "insert_image"

                # Skip source headings that ARE the section title itself — the
                # template already provides the correct heading name and number.
                if action == "insert_heading" and elem.text:
                    norm_elem = re.sub(r"^\d+[\.\ ]*", "", (elem.text or "").strip()).strip().lower()
                    if norm_elem == norm_src_title or norm_elem in norm_src_title or norm_src_title in norm_elem:
                        logger.debug(
                            f"Skipping source section-title heading '{elem.text}' "
                            f"(template heading '{section_plan.template_section_heading}' already present)"
                        )
                        continue

                placements.append(
                    ElementPlacement(
                        source_section_title=src_title,
                        source_element_index=idx,
                        source_element_type=elem.element_type,
                        target_section_heading=section_plan.template_section_heading,
                        placement_order=idx,
                        action=action,
                        heading_level=elem.level or 2,
                        embed_icons_inline=True if elem.icons else False,
                    )
                )

        # Append any remaining explicit placements not covered above
        for p in section_plan.elements:
            if p not in placements:
                placements.append(p)

        return placements


    def _execute_element(
        self,
        doc: Document,
        placement: ElementPlacement,
        extracted: DocxMigrationOutput,
        plan: MigrationPlan,
        anchor: Optional[Any] = None,
        section_plan: Optional[SectionPlan] = None,
    ) -> Optional[Any]:
        """Insert element and anchor it sequentially in the document."""
        if placement.action == "skip":
            return None

        if placement.action == "populate_placeholder":
            key = (placement.source_section_title, placement.source_element_index)
            if hasattr(self, "_populated_tbl_elements") and key in self._populated_tbl_elements:
                return self._populated_tbl_elements[key]
            return None

        source_elem = self._find_source_element(extracted, placement)
        if source_elem is None:
            logger.debug(f"Source element not found: {placement.source_section_title}[{placement.source_element_index}]")
            return None

        created_items: list[Any] = []

        match placement.action:
            case "insert_heading":
                parent_lvl = getattr(section_plan, "template_heading_level", 1) or 1
                raw_lvl = placement.heading_level or source_elem.level or (parent_lvl + 1 if parent_lvl else 2)
                lvl = max(parent_lvl + 1, raw_lvl) if parent_lvl else raw_lvl
                lvl = min(6, lvl)
                font_size = getattr(plan, f"font_size_heading{lvl}_pt", 12.0)
                icon_paths = [ref.path for ref in source_elem.icons] if (placement.embed_icons_inline and source_elem.icons) else []
                heading_text = self._clean_heading_number_prefix(source_elem.text or "")
                h = self.styler.insert_heading(
                    doc=doc,
                    text=heading_text,
                    level=lvl,
                    style_name=placement.heading_style or f"Heading {lvl}",
                    font_family=plan.font_family,
                    font_size_pt=font_size,
                    icon_paths=icon_paths,
                    icon_size_pt=24.0,
                )
                created_items = [h]

            case "insert_paragraph":
                icon_paths = [ref.path for ref in source_elem.icons] if source_elem.icons else []
                para = self.styler.insert_paragraph(
                    doc=doc,
                    text=source_elem.text or "",
                    font_family=plan.font_family,
                    font_size_pt=plan.font_size_body_pt,
                    icon_paths=icon_paths,
                    icon_size_pt=24.0,
                )
                created_items = [para]

            case "insert_list":
                icon_paths = [ref.path for ref in source_elem.icons] if (placement.embed_icons_inline and source_elem.icons) else []
                raw_items = source_elem.items or ([source_elem.text] if source_elem.text else [])
                paras = self.styler.insert_list(
                    doc=doc,
                    items=raw_items,
                    font_family=plan.font_family,
                    font_size_pt=plan.font_size_body_pt,
                    icon_paths=icon_paths,
                    icon_size_pt=20.0,
                )
                created_items = paras

            case "insert_table":
                t = self.table_migrator.insert_table(
                    doc=doc,
                    source_element=source_elem,
                    font_family=plan.font_family,
                )
                if t is not None:
                    created_items = [t]

            case "insert_image":
                if source_elem.image_path:
                    created_items = self.styler.insert_image(
                        doc=doc,
                        image_path=source_elem.image_path,
                        caption=source_elem.title,
                    )

            case "insert_callout":
                style = self._resolve_callout_style(plan, placement)
                c = self.callout_builder.build(
                    doc=doc,
                    source_element=source_elem,
                    style=style,
                    font_family=plan.font_family,
                )
                if c is not None:
                    created_items = [c]

        last_anchor = anchor
        for item in created_items:
            oxml_elm = getattr(item, "_element", None)
            if oxml_elm is not None:
                # If current item is a table and last_anchor is also a table,
                # Word automatically merges adjacent <w:tbl> elements.
                # To prevent this, insert a separator paragraph between them.
                if (
                    last_anchor is not None
                    and oxml_elm.tag.endswith("tbl")
                    and last_anchor.tag.endswith("tbl")
                ):
                    spacer = doc.add_paragraph()
                    spacer.paragraph_format.space_before = Pt(6.0)
                    spacer.paragraph_format.space_after = Pt(6.0)
                    last_anchor.addnext(spacer._element)
                    last_anchor = spacer._element

                if last_anchor is not None and oxml_elm != last_anchor:
                    last_anchor.addnext(oxml_elm)
                last_anchor = oxml_elm

        return last_anchor

    def _resolve_callout_style(
        self, plan: MigrationPlan, placement: ElementPlacement
    ) -> Optional[CalloutStyleDef]:
        """Resolve CalloutStyleDef from plan or placement colors."""
        if placement.callout_type:
            for s in plan.callout_styles:
                if s.callout_type == placement.callout_type:
                    return s

        if placement.callout_background_hex:
            return CalloutStyleDef(
                callout_type=placement.callout_type or "custom",
                display_name=placement.callout_type or "Custom Callout",
                background_color_hex=placement.callout_background_hex,
                left_border_color_hex=placement.callout_border_hex or "#2F5597",
            )

        return None

    def _find_source_element(
        self, extracted: DocxMigrationOutput, placement: ElementPlacement
    ) -> Optional[MigrationElement]:
        """Look up the exact source element from extracted JSON."""
        for sec in extracted.sections:
            if sec.title == placement.source_section_title:
                if 0 <= placement.source_element_index < len(sec.elements):
                    return sec.elements[placement.source_element_index]
        return None

    def _find_heading_anchor(self, doc: Document, heading_text: str) -> Optional[Any]:
        """Locate the OXML element for a section heading in the document."""
        import re
        norm_h = re.sub(r"^\d+[\.\s]*", "", heading_text.strip().lower()).strip()
        for para in doc.paragraphs:
            p_text = para.text.strip().lower()
            p_norm = re.sub(r"^\d+[\.\s]*", "", p_text).strip()
            if p_norm == norm_h or (norm_h and (norm_h == p_text or norm_h in p_text and len(p_text) < 80)):
                return para._element
        return None

    def _insert_fallback_text(self, doc: Document, anchor: Optional[Any], fallback_str: str):
        """Insert fallback text paragraph anchored after the heading."""
        new_p = doc.add_paragraph(fallback_str)
        if new_p.runs:
            new_p.runs[0].font.name = "Arial"
            new_p.runs[0].font.size = Pt(10.0)
        if anchor is not None and new_p._element != anchor:
            anchor.addnext(new_p._element)

    def _separate_adjacent_tables(self, doc: Document):
        """Ensure no two <w:tbl> elements are directly adjacent in the DOM, which causes Word to merge them."""
        body = doc._body._element
        tbl_elements = body.xpath("./w:tbl")
        for tbl in tbl_elements:
            next_elm = tbl.getnext()
            if next_elm is not None and next_elm.tag.endswith("tbl"):
                # Insert empty spacing paragraph between them
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6.0)
                p.paragraph_format.space_after = Pt(6.0)
                tbl.addnext(p._element)

    def _sanitize_shifted_icons(self, extracted: DocxMigrationOutput):
        """Self-healing check for legacy or drifting icon extraction.
        
        If an introductory sentence ending with ':' has an icon, and subsequent
        items have icons while the last item has no icon, shift the icons down
        to match their true semantic target.
        """
        for sec in extracted.sections:
            elems = sec.elements
            if len(elems) < 2:
                continue

            i = 0
            while i < len(elems) - 1:
                e0 = elems[i]
                if (
                    e0.element_type == "paragraph"
                    and e0.text
                    and e0.text.strip().endswith(":")
                    and e0.icons
                ):
                    group = [e0]
                    j = i + 1
                    while j < len(elems) and elems[j].element_type == "paragraph" and not (elems[j].text or "").strip().endswith(":"):
                        group.append(elems[j])
                        j += 1

                    if len(group) >= 3 and not group[-1].icons and all(g.icons for g in group[:-1]):
                        logger.info(f"Self-healing shifted icons in section '{sec.title}' under '{e0.text[:30]}...'")
                        collected_icons = [list(g.icons) for g in group]
                        group[0].icons = []
                        for k in range(1, len(group)):
                            group[k].icons = collected_icons[k - 1]
                    i = j
                else:
                    i += 1
