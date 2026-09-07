"""Docx styling service for typography, headings, paragraphs, lists, and images."""

from __future__ import annotations

from pathlib import Path
from typing import Optional
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from loguru import logger


import re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class DocxStyler:
    """Low-level python-docx / OXML element insertion and styling."""

    ORDERED_PREFIX_RE = re.compile(
        r"^(?:\(?\d{1,3}[.)]|\(?[a-zA-Z][.)]|\(?[ivxIVX]{1,5}[.)])\s+"
    )
    UNORDERED_PREFIX_RE = re.compile(
        r"^[\u2022\u2023\u25E6\u2043\u2219\u25AA\u25AB\u25CF\u25CB\u25A0\u25A1\u2013\u2014○●◆◇■□▪▫–—•‣⁃\*\-]\s*"
    )

    @staticmethod
    def _vertically_center_paragraph_content(para):
        """Align inline images/icons and text along the line's horizontal midline."""
        pPr = para._element.get_or_add_pPr()
        text_align = pPr.find(qn("w:textAlignment"))
        if text_align is None:
            text_align = OxmlElement("w:textAlignment")
            pPr.append(text_align)
        text_align.set(qn("w:val"), "center")

    def insert_heading(
        self,
        doc: Document,
        text: str,
        level: int = 1,
        style_name: Optional[str] = None,
        font_family: str = "Arial",
        font_size_pt: float = 14.0,
        icon_paths: Optional[list[str | Path]] = None,
        icon_size_pt: float = 24.0,
    ):
        """Insert a native Word heading with specified level, style, and optional left-aligned icons."""
        lvl = max(1, min(level, 9))
        heading = doc.add_heading(level=lvl)

        if style_name and style_name in [s.name for s in doc.styles]:
            heading.style = doc.styles[style_name]
        elif f"Heading {lvl}" in [s.name for s in doc.styles]:
            heading.style = doc.styles[f"Heading {lvl}"]

        heading.paragraph_format.space_before = Pt(10.0)
        heading.paragraph_format.space_after = Pt(4.0)
        heading.paragraph_format.keep_with_next = True

        # 1. Embed icon on the LEFT if present
        if icon_paths:
            self._vertically_center_paragraph_content(heading)
            for ipath in icon_paths:
                p = Path(ipath)
                if p.exists():
                    try:
                        icon_run = heading.add_run()
                        icon_run.add_picture(str(p), width=Pt(icon_size_pt), height=Pt(icon_size_pt))
                        heading.add_run("\u2003\u2002")
                    except Exception as exc:
                        logger.warning(f"Failed to embed icon '{p}' in heading: {exc}")

        # 2. Add heading text
        text_run = heading.add_run(text)
        text_run.font.name = font_family
        text_run.font.size = Pt(font_size_pt)

        return heading

    def insert_paragraph(
        self,
        doc: Document,
        text: str,
        font_family: str = "Arial",
        font_size_pt: float = 10.0,
        icon_paths: Optional[list[str | Path]] = None,
        icon_size_pt: float = 24.0,
    ):
        """Insert a body paragraph with left-aligned icons (if any) placed before text with generous spacing."""
        para = doc.add_paragraph()
        if "Normal" in [s.name for s in doc.styles]:
            para.style = doc.styles["Normal"]

        para.paragraph_format.space_after = Pt(4.0)
        para.paragraph_format.line_spacing = 1.15

        # 1. Embed icon on the LEFT first
        if icon_paths:
            self._vertically_center_paragraph_content(para)
            for ipath in icon_paths:
                p = Path(ipath)
                if p.exists():
                    try:
                        icon_run = para.add_run()
                        icon_run.add_picture(str(p), width=Pt(icon_size_pt), height=Pt(icon_size_pt))
                        para.add_run("\u2003\u2002")
                    except Exception as exc:
                        logger.warning(f"Failed to embed inline icon '{p}': {exc}")

        # 2. Add body text
        text_run = para.add_run(text)
        text_run.font.name = font_family
        text_run.font.size = Pt(font_size_pt)

        return para

    def insert_list(
        self,
        doc: Document,
        items: list[str],
        font_family: str = "Arial",
        font_size_pt: float = 10.0,
        style: Optional[str] = None,
        is_ordered: Optional[bool] = None,
        icon_paths: Optional[list[str | Path]] = None,
        icon_size_pt: float = 20.0,
    ) -> list[Any]:
        """Insert a bulleted or numbered list with clean hanging indents and stripped prefix markers."""
        if not items:
            return []

        # Determine if ordered or unordered (if not explicitly given)
        if is_ordered is None:
            if style and "number" in style.lower():
                is_ordered = True
            elif style and "bullet" in style.lower():
                is_ordered = False
            else:
                ordered_matches = sum(1 for it in items if self.ORDERED_PREFIX_RE.match((it or "").strip()))
                is_ordered = ordered_matches >= max(1, len(items) // 2)

        paras = []
        valid_items = [it for it in items if (it or "").strip()]
        total_items = len(valid_items)

        for idx, raw_item in enumerate(valid_items):
            item_text = raw_item.strip()
            para = doc.add_paragraph()

            # Detect sub-level (e.g. leading tabs/spaces or lettered item a), b))
            sub_level = 0
            if is_ordered and re.match(r"^\(?[a-zA-Z][.)]", item_text):
                sub_level = 1
            elif raw_item.startswith("\t") or raw_item.startswith("    "):
                sub_level = 1

            # Strip prefixes to prevent double bullets/numbers
            if is_ordered:
                cleaned_text = self.ORDERED_PREFIX_RE.sub("", item_text).strip()
                prefix_label = f"{idx + 1}.\t" if sub_level == 0 else f"{chr(97 + (idx % 26))}.\t"
            else:
                cleaned_text = self.UNORDERED_PREFIX_RE.sub("", item_text).strip()
                prefix_label = "•\t"

            # Hanging indent layout
            base_indent = 0.25 * (sub_level + 1)
            para.paragraph_format.left_indent = Inches(base_indent)
            para.paragraph_format.first_line_indent = Inches(-0.25)
            para.paragraph_format.space_before = Pt(0.0)
            para.paragraph_format.space_after = Pt(6.0 if idx == total_items - 1 else 2.0)
            para.paragraph_format.line_spacing = 1.15

            # If icons provided, embed on first list item and center vertically
            if idx == 0 and icon_paths:
                self._vertically_center_paragraph_content(para)
                for ipath in icon_paths:
                    p = Path(ipath)
                    if p.exists():
                        try:
                            icon_run = para.add_run()
                            icon_run.add_picture(str(p), width=Pt(icon_size_pt), height=Pt(icon_size_pt))
                            para.add_run("\u2003")
                        except Exception as exc:
                            logger.warning(f"Failed to embed icon in list item: {exc}")

            lbl_run = para.add_run(prefix_label)
            lbl_run.font.name = font_family
            lbl_run.font.size = Pt(font_size_pt)
            lbl_run.bold = is_ordered

            text_run = para.add_run(cleaned_text)
            text_run.font.name = font_family
            text_run.font.size = Pt(font_size_pt)
            paras.append(para)

        return paras

    def embed_inline_icon(
        self,
        paragraph,
        icon_path: str | Path,
        size_pt: float = 24.0,
    ):
        """Embed an icon image at the beginning of a paragraph with generous spacing."""
        path = Path(icon_path)
        if not path.exists():
            logger.warning(f"Icon file not found for inline embedding: {path}")
            return

        try:
            self._vertically_center_paragraph_content(paragraph)
            # If paragraph has runs, insert at start
            if paragraph.runs:
                first_run = paragraph.runs[0]
                icon_run = paragraph.add_run()
                icon_run.add_picture(str(path), width=Pt(size_pt), height=Pt(size_pt))
                spacer = paragraph.add_run("\u2003\u2002")
                # Move icon_run and spacer before first_run in OXML
                first_run._r.addprevious(icon_run._r)
                first_run._r.addprevious(spacer._r)
            else:
                icon_run = paragraph.add_run()
                icon_run.add_picture(str(path), width=Pt(size_pt), height=Pt(size_pt))
                paragraph.add_run("\u2003\u2002")
        except Exception as exc:
            logger.warning(f"Failed to embed inline icon '{path}': {exc}")

    def insert_image(
        self,
        doc: Document,
        image_path: str | Path,
        caption: Optional[str] = None,
        max_width_inches: float = 5.5,
    ) -> list[Any]:
        """Insert a centered image with optional caption below."""
        path = Path(image_path)
        if not path.exists():
            logger.warning(f"Image file not found for insertion: {path}")
            return []

        created = []
        try:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(6.0)
            para.paragraph_format.space_after = Pt(4.0)
            run = para.add_run()
            run.add_picture(str(path), width=Inches(max_width_inches))
            created.append(para)

            if caption:
                cap_para = doc.add_paragraph(caption)
                cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_para.paragraph_format.space_after = Pt(6.0)
                for r in cap_para.runs:
                    r.italic = True
                    r.font.name = "Arial"
                    r.font.size = Pt(9.0)
                created.append(cap_para)
        except Exception as exc:
            logger.warning(f"Failed to insert image '{path}': {exc}")

        return created
