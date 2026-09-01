from pathlib import Path
from typing import Optional
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from loguru import logger

from app.schemas.migration import MigrationElement
from app.services.migration.schemas import CalloutStyleDef


class CalloutBuilder:
    """Builds styled callout boxes as single-cell Word tables with OXML shading and left borders."""

    DEFAULT_CALLOUT_STYLES = {
        "executive_summary": CalloutStyleDef(
            callout_type="executive_summary",
            display_name="Executive Summary / Introduction",
            background_color_hex="#D9E1F2",
            left_border_color_hex="#2F5597",
            icon_description="Document / Summary icon",
        ),
        "explanation": CalloutStyleDef(
            callout_type="explanation",
            display_name="Explanation",
            background_color_hex="#E2EFDA",
            left_border_color_hex="#385723",
            icon_description="Question mark icon",
        ),
        "attention": CalloutStyleDef(
            callout_type="attention",
            display_name="Attention / Warning",
            background_color_hex="#FCE4D6",
            left_border_color_hex="#C65911",
            icon_description="Warning triangle icon",
        ),
        "key_takeaway": CalloutStyleDef(
            callout_type="key_takeaway",
            display_name="Key Takeaway",
            background_color_hex="#FFF2CC",
            left_border_color_hex="#BF8F00",
            icon_description="Eye / Focus icon",
        ),
    }

    def build(
        self,
        doc: Document,
        source_element: Optional[MigrationElement],
        style: Optional[CalloutStyleDef] = None,
        font_family: str = "Arial",
    ):
        """Construct a single-cell shaded callout box in the Word document with left-aligned icon."""
        if style is None:
            style = self.DEFAULT_CALLOUT_STYLES.get("executive_summary")

        bg_hex = style.background_color_hex if style else "#D9E1F2"
        border_hex = style.left_border_color_hex if style else "#2F5597"

        # Create single-cell table
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)

        # Center table
        tbl_pr = table._tbl.tblPr
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        tbl_pr.append(jc)

        # Apply OXML shading & borders
        tc_pr = cell._tc.get_or_add_tcPr()
        self._set_cell_shading(tc_pr, bg_hex)
        self._set_left_border(tc_pr, border_hex, width_pt=3.5)
        self._set_cell_margins(tc_pr, top=140, bottom=140, left=180, right=180)

        # Extract text and icons from source element
        para = cell.paragraphs[0]
        para.text = ""
        para.paragraph_format.line_spacing = 1.15
        para.paragraph_format.space_after = Pt(2.0)

        # 1. Embed icon on the LEFT if available (at 26pt legible size)
        icon_path = None
        if source_element:
            if source_element.icons:
                icon_path = source_element.icons[0].path
            elif source_element.cells:
                for c in source_element.cells:
                    if c.icon_path:
                        icon_path = c.icon_path
                        break

        if icon_path and Path(icon_path).exists():
            try:
                icon_run = para.add_run()
                icon_run.add_picture(str(icon_path), width=Pt(26.0), height=Pt(26.0))
                para.add_run("\u2003\u2003")
            except Exception as exc:
                logger.warning(f"Failed to embed icon in callout: {exc}")

        # 2. Add text
        text = self._extract_callout_text(source_element)
        text_run = para.add_run(text)
        text_run.font.name = font_family
        text_run.font.size = Pt(10.0)

        return table

    def _extract_callout_text(self, elem: Optional[MigrationElement]) -> str:
        """Extract plain text from table cells, list items, or paragraph."""
        if elem is None:
            return ""

        if elem.element_type == "paragraph" and elem.text:
            return elem.text.strip()

        if elem.element_type == "list" and elem.items:
            return "\n".join(elem.items)

        if elem.element_type == "table" and elem.cells:
            parts = [c.text.strip() for c in elem.cells if c.text.strip()]
            return " — ".join(parts)

        return elem.text or ""

    def _set_cell_shading(self, tc_pr, hex_color: str):
        """Apply background shading to a table cell via OXML w:shd element."""
        clean_hex = hex_color.lstrip("#").upper()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), clean_hex)
        tc_pr.append(shd)

    def _set_left_border(self, tc_pr, hex_color: str, width_pt: float = 3.5):
        """Set a thick left border and nil top/bottom/right borders."""
        clean_hex = hex_color.lstrip("#").upper()
        sz_val = str(int(width_pt * 8))

        borders = OxmlElement("w:tcBorders")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), sz_val)
        left.set(qn("w:space"), "0")
        left.set(qn("w:color"), clean_hex)
        borders.append(left)

        for side in ["top", "right", "bottom"]:
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "nil")
            borders.append(el)

        tc_pr.append(borders)

    def _set_cell_margins(self, tc_pr, top=140, bottom=140, left=180, right=180):
        """Set internal cell padding (margins) in dxa."""
        tc_mar = OxmlElement("w:tcMar")
        for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), str(val))
            m.set(qn("w:type"), "dxa")
            tc_mar.append(m)
        tc_pr.append(tc_mar)
