"""Template inspector service for extracting structural profiles from .docx templates."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Optional
from docx import Document
from loguru import logger

from app.services.migration.schemas import (
    TemplateRawProfile,
    TemplateParagraphInfo,
    TemplateTableInfo,
)


class TemplateInspector:
    """Parses a .docx template into a structured TemplateRawProfile."""

    BLUE_HEX_VALUES = {
        "0000FF", "0070C0", "4472C4", "2E74B5", "5B9BD5", "00B0F0",
        "1F4E78", "2F5597", "41719C", "002060", "1B365D", "0072CE",
        "0075FF", "00A2E8", "004B87", "0080FF", "0099FF", "0066CC",
        "1B75BB", "005696", "3399FF", "003366", "2E75B6",
    }

    PLACEHOLDER_PATTERN = re.compile(
        r"(\$\{[^}]+\}|<<[^>]+>>|\[Insert\s+[^\]]+\]|\[Role\s*\d+\]|e\.g\.,\s*xxx-AD|\$\{vault:[^}]+\}|\bn\.0\b|BI-VQD-\d+)",
        re.IGNORECASE,
    )

    INSTRUCTION_KEYWORDS = (
        "brief description of what",
        "brief description of the intention",
        "explain which target roles",
        "executive summary/introduction",
        "explanation section",
        "attention section",
        "key-take-away section",
        "embedding of files is not possible",
        "to ensure effective communication",
        "for global and (cross-) divisional gp docs",
    )

    def inspect(self, template_path: Path | str) -> TemplateRawProfile:
        """Inspect a .docx template and extract its structural profile."""
        path = Path(template_path)
        if not path.exists():
            raise FileNotFoundError(f"Template file not found: {path}")

        doc = Document(str(path))
        logger.info(f"Inspecting template '{path.name}' ({len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables)")

        paragraphs_info: list[TemplateParagraphInfo] = []
        heading_history: list[tuple[int, str]] = []  # (paragraph_index, heading_text)

        for i, para in enumerate(doc.paragraphs):
            text = (para.text or "").strip()
            style_name = para.style.name if para.style else "Normal"
            heading_level = self._detect_heading_level(para, style_name, text)
            is_blue, color_hex = self._detect_color(para)
            is_bold = any(getattr(r, "bold", False) for r in para.runs)

            if heading_level is not None and text:
                heading_history.append((i, text))

            paragraphs_info.append(
                TemplateParagraphInfo(
                    index=i,
                    text=text,
                    style_name=style_name,
                    is_blue_instruction=is_blue,
                    is_bold=is_bold,
                    heading_level=heading_level,
                    font_color_hex=color_hex,
                )
            )

        tables_info: list[TemplateTableInfo] = []
        for t_idx, table in enumerate(doc.tables):
            num_rows = len(table.rows)
            num_cols = len(table.columns) if num_rows > 0 else 0

            # Extract header row
            header_texts: list[str] = []
            if num_rows > 0:
                header_texts = [cell.text.strip() for cell in table.rows[0].cells if cell.text.strip()]

            # Sample cells & placeholder check
            sample_cells: list[str] = []
            has_placeholder = False
            is_vault = False
            has_blue_in_table = False
            has_instruction_phrase = False

            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    c_text = cell.text.strip()
                    if c_text and len(sample_cells) < 10:
                        sample_cells.append(c_text)
                    if "${vault:" in c_text.lower():
                        is_vault = True
                    if self.PLACEHOLDER_PATTERN.search(c_text):
                        has_placeholder = True
                    c_lower = c_text.lower()
                    if any(kw in c_lower for kw in self.INSTRUCTION_KEYWORDS):
                        has_instruction_phrase = True

                    # Check cell paragraphs for blue instruction text
                    for p in cell.paragraphs:
                        is_blue, _ = self._detect_color(p)
                        if is_blue:
                            has_blue_in_table = True

            # Find nearest preceding heading
            preceding_heading = self._find_preceding_heading(doc, table, heading_history)

            # Detect Document History table
            hdr_str = " ".join(header_texts).lower()
            is_doc_history = ("version" in hdr_str and "author" in hdr_str) or (
                preceding_heading is not None and "document history" in preceding_heading.lower() and not is_vault
            )
            if is_doc_history:
                has_placeholder = True

            is_instructional = (has_blue_in_table or has_instruction_phrase) and not is_vault
            if is_instructional:
                has_placeholder = True

            tables_info.append(
                TemplateTableInfo(
                    index=t_idx,
                    num_rows=num_rows,
                    num_cols=num_cols,
                    preceding_heading=preceding_heading,
                    header_row_text=header_texts,
                    has_placeholder_text=has_placeholder,
                    sample_cells=sample_cells,
                    is_vault_token_table=is_vault,
                    is_document_history_table=is_doc_history,
                    is_instructional_table=is_instructional,
                )
            )

        # Header / footer extraction
        header_text = ""
        footer_text = ""
        if doc.sections:
            sec0 = doc.sections[0]
            if sec0.header:
                header_text = "\n".join(p.text.strip() for p in sec0.header.paragraphs if p.text.strip())
            if sec0.footer:
                footer_text = "\n".join(p.text.strip() for p in sec0.footer.paragraphs if p.text.strip())

        return TemplateRawProfile(
            filename=path.name,
            total_paragraphs=len(paragraphs_info),
            total_tables=len(tables_info),
            paragraphs=paragraphs_info,
            tables=tables_info,
            header_text=header_text,
            footer_text=footer_text,
        )

    def _detect_heading_level(self, para, style_name: str, text: str) -> Optional[int]:
        """Detect heading level from style or text patterns."""
        if style_name.startswith("Heading"):
            try:
                lvl = int(style_name.replace("Heading", "").strip())
                return lvl
            except ValueError:
                return 1

        # Check numbered heading pattern: e.g., "1 PURPOSE", "6.1 LANGUAGE"
        match = re.match(r"^(\d+(?:\.\d+)*)\s+[A-Z]", text)
        if match:
            parts = match.group(1).split(".")
            return len(parts)

        return None

    def _detect_color(self, para) -> tuple[bool, Optional[str]]:
        """Detect if paragraph contains blue instruction text and return color hex."""
        # Check style name first
        if para.style and para.style.name:
            st_name = para.style.name.lower()
            if "instruction" in st_name:
                return True, "STYLE_INSTRUCTION"

        WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

        def _check_hex_or_rgb(hex_str: str) -> bool:
            clean = hex_str.upper().lstrip("#")
            if clean in self.BLUE_HEX_VALUES:
                return True
            if len(clean) == 6:
                try:
                    r = int(clean[0:2], 16)
                    g = int(clean[2:4], 16)
                    b = int(clean[4:6], 16)
                    if (b > 130 and b > r + 30) or (b > 140 and r < 100):
                        return True
                except ValueError:
                    pass
            return False

        for run in para.runs:
            # 1. Check python-docx RGBColor
            if run.font and run.font.color and run.font.color.rgb:
                rgb_str = str(run.font.color.rgb).upper()
                if _check_hex_or_rgb(rgb_str):
                    return True, rgb_str

            # 2. Check OXML w:rPr/w:color
            r_pr = run._element.find(f"{{{WNS}}}rPr")
            if r_pr is not None:
                color_el = r_pr.find(f"{{{WNS}}}color")
                if color_el is not None:
                    val = (
                        color_el.get(f"{{{WNS}}}val", "")
                        or color_el.get("w:val", "")
                        or color_el.get("val", "")
                    ).upper()
                    if _check_hex_or_rgb(val):
                        return True, val

                    theme_val = (
                        color_el.get(f"{{{WNS}}}themeColor", "")
                        or color_el.get("w:themeColor", "")
                        or color_el.get("themeColor", "")
                    ).lower()
                    if theme_val in {"accent1", "accent2", "accent5", "hyperlink"}:
                        return True, theme_val

        # 3. Check paragraph-level rPr
        p_pr = para._element.find(f"{{{WNS}}}pPr")
        if p_pr is not None:
            r_pr = p_pr.find(f"{{{WNS}}}rPr")
            if r_pr is not None:
                color_el = r_pr.find(f"{{{WNS}}}color")
                if color_el is not None:
                    val = (
                        color_el.get(f"{{{WNS}}}val", "")
                        or color_el.get("w:val", "")
                        or color_el.get("val", "")
                    ).upper()
                    if _check_hex_or_rgb(val):
                        return True, val

        return False, None

    def _find_preceding_heading(
        self, doc: Document, table, heading_history: list[tuple[int, str]]
    ) -> Optional[str]:
        """Find the nearest heading text preceding a table in the document body."""
        if not heading_history:
            return None

        # Determine table position in body relative to paragraphs
        try:
            tbl_elm = table._element
            body_elms = list(doc._body._element)
            if tbl_elm in body_elms:
                tbl_pos = body_elms.index(tbl_elm)
                # Look backwards in body_elms for a heading paragraph
                for i in range(tbl_pos - 1, -1, -1):
                    elm = body_elms[i]
                    if elm.tag.endswith("p"):
                        p_text = "".join(elm.itertext()).strip()
                        if p_text:
                            for _, h_text in reversed(heading_history):
                                if h_text == p_text or h_text in p_text:
                                    return h_text
                # If no preceding heading found above table position, return None
                return None
        except Exception:
            pass

        return None
