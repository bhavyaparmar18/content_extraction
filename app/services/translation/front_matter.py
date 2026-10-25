"""Language template front matter: a language's first page, headers and footers (Phase T3).

A language template is a full ``.docx``, but only its first page, headers and
footers are used. They replace the English ones in a translated document as they
are: they are never translated. A job without a language template keeps the
English front matter.

- ``extract_front_matter`` cuts the template down to a *package*: a ``.docx``
  with only the first page, the headers and footers of the first page's section
  and the first body section (with their images), and the template's styles and
  numbering. A translation job snapshots the package. The preview is the package
  with two placeholder body pages.
- ``apply_front_matter`` puts a package's first page, headers and footers into a
  rendered document in place of its own.
- ``front_matter_problems`` checks a document against a package.

The first page runs from the start of the body through the first section break,
or through the first page break when that comes earlier. Headers and footers
follow Word's inheritance: a section without a reference of one type uses the
previous section's. Header and footer parts are copied with their own
relationship ids, so their XML stays byte-identical to the package's; the first
page's relationships are created again in the target document.
"""

from __future__ import annotations

import copy
import hashlib
import json
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Iterable, Optional

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.oxml.ns import qn
from docx.parts.hdrftr import FooterPart, HeaderPart
from docx.parts.image import ImagePart
from lxml import etree

from app.schemas.v2.translation import CoverEnd, FrontMatterPart, FrontMatterReport, FrontMatterSwap
from app.services.migration_v2.render.xml import (
    W_P,
    W_R,
    W_SDT,
    insert_child,
    next_bookmark_id,
    remove,
    renumber_drawings,
    w_el,
)
from app.services.parser.ooxml import is_toc_sdt
from app.services.translation.languages import get_language

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
O_RELID = "{urn:schemas-microsoft-com:office:office}relid"
W_SECTPR = qn("w:sectPr")
REF_TAGS = {qn("w:headerReference"): "header", qn("w:footerReference"): "footer"}
REF_TYPES = ("default", "first", "even")
_STORY = {"header": (RT.HEADER, HeaderPart, "/word/header%d.xml"),
          "footer": (RT.FOOTER, FooterPart, "/word/footer%d.xml")}
_COPYABLE = {RT.IMAGE, RT.HYPERLINK}
# Relationships of a part to its content: dropped once nothing in the part refers to them.
_CONTENT_RELS = {RT.IMAGE, RT.HYPERLINK, RT.HEADER, RT.FOOTER, RT.OLE_OBJECT, RT.CHART, RT.PACKAGE, RT.CONTROL,
                 RT.DIAGRAM_DATA, RT.DIAGRAM_LAYOUT, RT.DIAGRAM_QUICK_STYLE, RT.DIAGRAM_COLORS, RT.A_F_CHUNK}
# sectPr children in schema order.
_SECTPR_ORDER = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
                 "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
                 "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"]
# settings children that come after evenAndOddHeaders.
_AFTER_EVEN_ODD = {
    "bookFoldRevPrinting", "bookFoldPrinting", "bookFoldPrintingSheets", "drawingGridHorizontalSpacing",
    "drawingGridVerticalSpacing", "displayHorizontalDrawingGridEvery", "displayVerticalDrawingGridEvery",
    "doNotUseMarginsForDrawingGridOrigin", "drawingGridHorizontalOrigin", "drawingGridVerticalOrigin",
    "doNotShadeFormData", "noPunctuationKerning", "characterSpacingControl", "printTwoOnOne",
    "strictFirstAndLastChars", "noLineBreaksAfter", "noLineBreaksBefore", "savePreviewPicture",
    "doNotValidateAgainstSchema", "saveInvalidXml", "ignoreMixedContent", "alwaysShowPlaceholderText",
    "doNotDemarcateInvalidXml", "saveXmlDataOnly", "useXSLTWhenSaving", "saveThroughXslt", "showXMLTags",
    "alwaysMergeEmptyNamespace", "updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars",
    "rsids", "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
    "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
    "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator",
}
_TRACKED = tuple(qn(f"w:{t}") for t in ("ins", "del", "moveFrom", "moveTo"))
_COMMENT_MARKS = (qn("w:commentRangeStart"), qn("w:commentRangeEnd"))
_EXCERPT = 300
PREVIEW_TEXT = ("Translated document text appears here (the first page of the body).",
                "Translated document text continues here.")

Refs = dict[tuple[str, str], str]  # (kind, type) → rId


class FrontMatterError(ValueError):
    """A document whose first page cannot be found, or front matter that cannot be copied."""


# ── The first page ────────────────────────────────────────────────────


@dataclass
class Cover:
    blocks: list                      # body blocks of the first page, in order
    end: CoverEnd
    next_block: Optional[object] = None  # the block whose "page break before" ends the first page
    warnings: list[str] = field(default_factory=list)


def _blocks(body) -> list:
    return [el for el in body if el.tag != W_SECTPR and isinstance(el.tag, str)]


def _section_break(el):
    """The section properties *el* ends a section with, or None."""
    paragraphs = [el] if el.tag == W_P else list(el.iter(W_P)) if el.tag == W_SDT else []
    for p in paragraphs:
        sect = p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}")
        if sect is not None:
            return sect
    return None


def _starts_new_page(sect, sects: list) -> bool:
    """A section break starts a new page unless the section after it starts continuously (a section's
    ``w:type`` is how that section begins)."""
    following = sects[sects.index(sect) + 1] if sect in sects and sects.index(sect) + 1 < len(sects) else None
    kind = following.find(qn("w:type")) if following is not None else None
    return kind is None or kind.get(qn("w:val")) != "continuous"


def _page_break(el):
    return next((br for br in el.iter(qn("w:br")) if br.get(qn("w:type")) == "page"), None)


def _break_before(el) -> bool:
    node = el.find(f"{qn('w:pPr')}/{qn('w:pageBreakBefore')}") if el.tag == W_P else None
    return node is not None and node.get(qn("w:val")) not in ("0", "false", "off")


def _text_after(el, br) -> bool:
    seen = False
    for node in el.iter(qn("w:br"), qn("w:t")):
        if node is br:
            seen = True
        elif seen and node.tag == qn("w:t") and (node.text or "").strip():
            return True
    return False


def find_cover(body) -> Cover:
    """The first page: the body blocks through the first section break or page break, whichever comes first."""
    blocks = _blocks(body)
    sects = section_props(body)
    cover = None
    for i, el in enumerate(blocks):
        if i > 0 and _break_before(el):
            cover = Cover(blocks[:i], "page_break", next_block=el)
            break
        sect = _section_break(el)
        if sect is not None and not _starts_new_page(sect, sects):
            raise FrontMatterError("the first page has a continuous section break; only a section break that starts "
                                   "a new page can end the first page")
        if sect is not None:
            cover = Cover(blocks[:i + 1], "section_break")
            break
        br = _page_break(el)
        if br is not None:
            cover = Cover(blocks[:i + 1], "page_break")
            if _text_after(el, br):
                cover.warnings.append("text after the page break that ends the first page stays with the first page")
            break
    if cover is None:
        raise FrontMatterError("no first page: the document has no section break or page break")
    if any(el.tag == W_SDT and is_toc_sdt(el) for el in cover.blocks):
        cover.warnings.append("the table of contents is on the first page")
    return cover


# ── Sections, headers and footers ─────────────────────────────────────


def section_props(body) -> list:
    """The body's section properties in document order, one per section; the last is the body's own."""
    return [s for s in body.iter(W_SECTPR) if s.getparent().tag != qn("w:sectPrChange")]


def own_refs(sect) -> Refs:
    return {(REF_TAGS[el.tag], el.get(qn("w:type")) or "default"): el.get(qn("r:id"))
            for el in sect if el.tag in REF_TAGS}


def effective_refs(sects: list) -> list[Refs]:
    """Each section's headers and footers as Word shows them: a type a section does not name is the previous
    section's."""
    out, current = [], {}
    for sect in sects:
        current = {**current, **own_refs(sect)}
        out.append(dict(current))
    return out


def _on(el) -> bool:
    return el is not None and el.get(qn("w:val")) not in ("0", "false", "off")


def title_page(sect) -> bool:
    return _on(sect.find(qn("w:titlePg")))


def set_refs(sect, refs: Refs, first_page: bool) -> None:
    """Write *refs* as the section's own header and footer references, and set its different first page."""
    for el in [c for c in sect if c.tag in REF_TAGS]:
        sect.remove(el)
    order = sorted(refs, key=lambda k: (k[0] != "header", REF_TYPES.index(k[1])))
    for i, key in enumerate(order):
        el = w_el(f"{key[0]}Reference", type=key[1])
        el.set(qn("r:id"), refs[key])
        sect.insert(i, el)
    for el in sect.findall(qn("w:titlePg")):
        sect.remove(el)
    if first_page:
        insert_child(sect, w_el("titlePg"), _SECTPR_ORDER)


def even_and_odd(doc) -> bool:
    return _on(doc.settings.element.find(qn("w:evenAndOddHeaders")))


def _set_even_and_odd(doc, on: bool) -> None:
    settings = doc.settings.element
    for el in settings.findall(qn("w:evenAndOddHeaders")):
        settings.remove(el)
    if not on:
        return
    node = w_el("evenAndOddHeaders")
    follower = next((c for c in settings if isinstance(c.tag, str) and etree.QName(c).localname in _AFTER_EVEN_ODD),
                    None)
    if follower is not None:
        follower.addprevious(node)
    else:
        settings.append(node)


# ── Relationships ─────────────────────────────────────────────────────


def _rel_attrs(roots) -> Iterable[tuple]:
    """(element, attribute, rId) of every relationship reference under *roots*."""
    for root in ([roots] if not isinstance(roots, (list, tuple)) else roots):
        for el in root.iter(tag=etree.Element):
            for name, value in el.attrib.items():
                if name.startswith("{" + R_NS + "}") or name == O_RELID:
                    yield el, name, value


def _rel_name(reltype: str) -> str:
    return reltype.rsplit("/", 1)[-1]


def _drop_unused_rels(part) -> list[str]:
    """Drop the part's content relationships that nothing in its XML refers to; the partnames of the dropped
    headers and footers."""
    used = {rId for _, _, rId in _rel_attrs(part.element)}
    dropped = []
    for rId, rel in list(part.rels.items()):
        if rId in used or rel.reltype not in _CONTENT_RELS:
            continue
        if not rel.is_external and rel.reltype in (RT.HEADER, RT.FOOTER):
            dropped.append(str(rel.target_part.partname))
        del part.rels[rId]
        part.rels.related_parts.pop(rId, None)
    return dropped


class _Images:
    """Image parts of a target package, matched by content so an image is stored once."""

    def __init__(self, package):
        self.package = package
        parts = list(package.iter_parts())
        self.by_hash = {hashlib.sha1(p.blob).hexdigest(): p for p in parts if isinstance(p, ImagePart)}
        self.names = {str(p.partname) for p in parts}

    def get(self, source) -> ImagePart:
        digest = hashlib.sha1(source.blob).hexdigest()
        part = self.by_hash.get(digest)
        if part is None:
            # Named here: a new image is not reachable from the document until its part is related.
            ext = Path(str(source.partname)).suffix or ".bin"
            n = 1
            while f"/word/media/image{n}{ext}" in self.names:
                n += 1
            partname = PackURI(f"/word/media/image{n}{ext}")
            self.names.add(str(partname))
            part = ImagePart.load(partname, source.content_type, source.blob, self.package)
            self.by_hash[digest] = part
        return part


def _copy_rel(source_part, rId: str, target_part, images: _Images) -> str:
    rel = source_part.rels.get(rId)
    if rel is None:
        raise FrontMatterError(f"the front matter refers to a missing relationship {rId}")
    if rel.is_external:
        return target_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
    if rel.reltype == RT.IMAGE:
        return target_part.relate_to(images.get(rel.target_part), RT.IMAGE)
    raise FrontMatterError(f"a {_rel_name(rel.reltype)} object in the front matter cannot be copied")


def _copy_story(source, kind: str, doc_part, images: _Images) -> str:
    """Copy a header or footer part, with its own relationship ids, into the document; the new rId."""
    reltype, cls, template = _STORY[kind]
    package = doc_part.package
    part = cls.load(package.next_partname(template), source.content_type, source.blob, package)
    for rId, rel in source.rels.items():
        if rel.is_external:
            part.load_rel(rel.reltype, rel.target_ref, rId, is_external=True)
        elif rel.reltype == RT.IMAGE:
            part.load_rel(RT.IMAGE, images.get(rel.target_part), rId)
        else:
            raise FrontMatterError(f"a {_rel_name(rel.reltype)} object in a {kind} cannot be copied")
    return doc_part.relate_to(part, reltype)


def _identity(part, rId: str) -> str:
    rel = part.rels.get(rId)
    if rel is None:
        return "missing"
    if rel.is_external:
        return f"url:{rel.target_ref}"
    return f"sha1:{hashlib.sha1(rel.target_part.blob).hexdigest()}"


# ── Styles and numbering ──────────────────────────────────────────────


def _style_ids(roots) -> set[str]:
    ids = set()
    for root in roots:
        for tag in ("w:pStyle", "w:rStyle", "w:tblStyle"):
            ids |= {el.get(qn("w:val")) for el in root.iter(qn(tag))}
    return ids - {None}


def _style_name(style) -> Optional[str]:
    name = style.find(qn("w:name"))
    return name.get(qn("w:val")) if name is not None else None


def _definition(style) -> bytes:
    style = copy.deepcopy(style)
    for el in style.findall(qn("w:rsid")):
        style.remove(el)
    return etree.tostring(style, method="c14n", exclusive=True)


def _styles_in_use(styles: dict, roots) -> set[str]:
    """The styles *roots* use, the default styles, and every style those are based on or linked to."""
    pending = list(_style_ids(roots)) + [sid for sid, s in styles.items() if _on_attr(s, "w:default")]
    used: set[str] = set()
    while pending:
        style_id = pending.pop()
        if style_id in used or style_id not in styles:
            continue
        used.add(style_id)
        for tag in ("w:basedOn", "w:link"):
            ref = styles[style_id].find(qn(tag))
            if ref is not None and ref.get(qn("w:val")):
                pending.append(ref.get(qn("w:val")))
    return used


def _on_attr(el, name: str) -> bool:
    return el.get(qn(name)) in ("1", "true", "on")


def _copy_styles(source_doc, doc, roots, body_roots, language: str, swap: FrontMatterSwap) -> list:
    """Give the document the styles *roots* (the first page, headers and footers) use, with the styles they are
    based on; the new style elements.

    A style the document lacks is copied. A style it defines differently takes the language template's definition
    when nothing in *body_roots* (the rest of the document) depends on it, as with the Header and Footer styles; a
    style the body uses keeps the document's definition, with a warning.
    """
    source = {s.get(qn("w:styleId")): s for s in source_doc.styles.element.findall(qn("w:style"))}
    styles = doc.styles.element
    target = {s.get(qn("w:styleId")): s for s in styles.findall(qn("w:style"))}
    body_styles = _styles_in_use(target, body_roots)
    names = {(_style_name(s) or "").lower() for s in target.values()}
    pending, seen, copies, differing = sorted(_style_ids(roots)), set(), [], []
    while pending:
        style_id = pending.pop(0)
        if style_id in seen or style_id not in source:
            continue  # a style neither file defines: Word uses the default
        seen.add(style_id)
        new = copy.deepcopy(source[style_id])
        if style_id in target:
            old = target[style_id]
            if _definition(old) == _definition(new):
                continue
            if style_id in body_styles:
                differing.append(style_id)
                continue
            new.attrib.pop(qn("w:default"), None)
            if old.get(qn("w:default")) is not None:
                new.set(qn("w:default"), old.get(qn("w:default")))
            old.addprevious(new)
            styles.remove(old)
            swap.styles_replaced.append(style_id)
        else:
            new.attrib.pop(qn("w:default"), None)
            name = _style_name(new)
            if name and name.lower() in names:  # e.g. German "Kopfzeile", named "header", next to English "Header"
                name = f"{name} ({language})"
                new.find(qn("w:name")).set(qn("w:val"), name)
            styles.append(new)
            names.add((name or "").lower())
            swap.styles_added.append(style_id)
        target[style_id] = new
        copies.append(new)
        for tag in ("w:basedOn", "w:link", "w:next"):
            ref = new.find(qn(tag))
            if ref is not None and ref.get(qn("w:val")):
                pending.append(ref.get(qn("w:val")))
    if differing:
        swap.warnings.append("styles the document's body also uses keep the document's definition, although the "
                             "language template's differ: " + ", ".join(sorted(differing)))
    return copies


def _num_uses(roots) -> list:
    return [el for root in roots for el in root.iter(qn("w:numId")) if el.getparent().tag == qn("w:numPr")]


def _copy_numbering(source_doc, doc, roots, swap: FrontMatterSwap) -> None:
    """Copy the lists *roots* use as new numbering definitions, and point them at the copies."""
    uses = [el for el in _num_uses(roots) if el.get(qn("w:val")) not in (None, "0")]
    if not uses:
        return
    source = source_doc.part.numbering_part.element
    target = doc.part.numbering_part.element
    nums = {n.get(qn("w:numId")): n for n in source.findall(qn("w:num"))}
    abstracts = {a.get(qn("w:abstractNumId")): a for a in source.findall(qn("w:abstractNum"))}

    def ids(tag, attr):
        return [int(e.get(qn(attr))) for e in target.findall(qn(tag)) if (e.get(qn(attr)) or "").isdigit()]

    next_num = max(ids("w:num", "w:numId"), default=0) + 1
    next_abstract = max(ids("w:abstractNum", "w:abstractNumId"), default=-1) + 1
    mapping: dict[str, Optional[str]] = {}
    for use in uses:
        old = use.get(qn("w:val"))
        if old not in mapping:
            num = nums.get(old)
            link = num.find(qn("w:abstractNumId")) if num is not None else None
            abstract = abstracts.get(link.get(qn("w:val"))) if link is not None else None
            if abstract is None:
                swap.warnings.append(f"list {old} of the language template has no definition; its paragraphs lose "
                                     "their numbering")
                mapping[old] = None
            else:
                new_abstract = copy.deepcopy(abstract)
                new_abstract.set(qn("w:abstractNumId"), str(next_abstract))
                for nsid in new_abstract.findall(qn("w:nsid")):
                    new_abstract.remove(nsid)
                new_num = copy.deepcopy(num)
                new_num.set(qn("w:numId"), str(next_num))
                new_num.find(qn("w:abstractNumId")).set(qn("w:val"), str(next_abstract))
                first_num = target.find(qn("w:num"))
                if first_num is not None:
                    first_num.addprevious(new_abstract)
                else:
                    target.append(new_abstract)
                existing = target.findall(qn("w:num"))
                (existing[-1].addnext(new_num) if existing else new_abstract.addnext(new_num))
                mapping[old] = str(next_num)
                next_num += 1
                next_abstract += 1
        if mapping[old] is None:
            remove(use.getparent())
        else:
            use.set(qn("w:val"), mapping[old])


def _renumber_bookmarks(blocks, existing_names: set, next_id: int, swap: FrontMatterSwap) -> None:
    ids: dict[str, str] = {}
    dropped = []
    for root in blocks:
        for el in list(root.iter(qn("w:bookmarkStart"))):
            name = el.get(qn("w:name"))
            if name in existing_names:
                dropped.append(name)
                remove(el)
                continue
            ids[el.get(qn("w:id"))] = str(next_id)
            el.set(qn("w:id"), str(next_id))
            next_id += 1
    for root in blocks:
        for el in list(root.iter(qn("w:bookmarkEnd"))):
            new = ids.get(el.get(qn("w:id")))
            if new is None:
                remove(el)
            else:
                el.set(qn("w:id"), new)
    if dropped:
        swap.warnings.append(f"bookmarks on the language template's first page already exist in the document and "
                             f"were dropped: {', '.join(sorted(set(dropped)))}")


# ── Extraction ────────────────────────────────────────────────────────


def _text(roots) -> str:
    return re.sub(r"\s+", " ", " ".join(t.text or "" for root in roots for t in root.iter(qn("w:t")))).strip()


def _strip_comments(roots) -> int:
    count = 0
    for root in roots:
        for el in list(root.iter(*_COMMENT_MARKS)):
            remove(el)
            count += 1
        for ref in list(root.iter(qn("w:commentReference"))):
            run = ref.getparent()
            remove(run if run is not None and run.tag == W_R else ref)
            count += 1
    return count


def _end_with_page_break(blocks: list) -> None:
    """The first page ended because the next block starts a new page; the package's first page carries the break
    itself, since that block is not copied."""
    last = blocks[-1]
    run = w_el("r")
    run.append(w_el("br", type="page"))
    if last.tag == W_P:
        last.append(run)
    else:
        p = w_el("p")
        p.append(run)
        last.addnext(p)
        blocks.append(p)


def _docvariables(roots) -> list[str]:
    """The document variables the DOCVARIABLE fields under *roots* show."""
    instructions = [el.text or "" for root in roots for el in root.iter(qn("w:instrText"))]
    instructions += [el.get(qn("w:instr")) or "" for root in roots for el in root.iter(qn("w:fldSimple"))]
    return sorted({m.group(1) for text in instructions for m in [re.match(r"\s*DOCVARIABLE\s+(\S+)", text)] if m})


def _part_info(kind: str, type_: str, part) -> FrontMatterPart:
    return FrontMatterPart(kind=kind, type=type_, partname=str(part.partname), text=_text([part.element])[:120],
                           images=sum(1 for r in part.rels.values() if not r.is_external and r.reltype == RT.IMAGE))


def _script_warning(language: str, text: str) -> Optional[str]:
    script = get_language(language).script
    if script in ("han", "japanese") and not re.search(r"[぀-ヿ㐀-鿿豈-﫿]", text):
        return f"the first page, headers and footers have no Chinese or Japanese characters: is this the {language} template?"
    if script == "greek" and not re.search(r"[Ͱ-Ͽἀ-῿]", text):
        return f"the first page, headers and footers have no Greek letters: is this the {language} template?"
    return None


def extract_front_matter(
    template_path: str | Path,
    package_path: str | Path,
    preview_path: Optional[str | Path] = None,
    *,
    template_id: str,
    version: int,
    language: str,
    source_filename: Optional[str] = None,
) -> FrontMatterReport:
    """Cut a language template down to its front matter, written to *package_path* (and a preview).

    A template whose first page cannot be found gets a report with that problem and no package. Raises
    ``FrontMatterError`` when the file is not a Word document.
    """
    data = Path(template_path).read_bytes()
    report = FrontMatterReport(template_id=template_id, version=version, language=language,
                               source_filename=source_filename, source_sha256=hashlib.sha256(data).hexdigest())
    try:
        doc = Document(BytesIO(data))
    except Exception as exc:
        raise FrontMatterError(f"the file cannot be opened as a Word document: {exc}") from exc
    body = doc.element.body
    try:
        cover = find_cover(body)
    except FrontMatterError as exc:
        report.problems.append(str(exc))
        return report
    report.warnings += cover.warnings
    sects = section_props(body)
    effective = effective_refs(sects)
    body_index = 1 if cover.end == "section_break" else 0
    if len(sects) > body_index + 1:
        report.warnings.append(f"the template has {len(sects)} sections: only the first page's and the first body "
                               "section's headers and footers are used")
    report.cover_end, report.sections = cover.end, len(sects)
    report.title_page_cover, report.title_page_body = title_page(sects[0]), title_page(sects[body_index])
    report.even_and_odd = even_and_odd(doc)

    # The package: the first page, then an empty body page with the body section's headers and footers.
    blocks = list(cover.blocks)
    if cover.next_block is not None:
        _end_with_page_break(blocks)
    body_sect = copy.deepcopy(sects[body_index])
    set_refs(body_sect, effective[body_index], title_page(sects[body_index]))
    keep = set(blocks)
    for el in list(body):
        if el not in keep:
            body.remove(el)
    body.append(w_el("p"))
    body.append(body_sect)

    stories = list({rId: doc.part.related_parts[rId] for refs in (effective[0], effective[body_index])
                    for rId in refs.values() if rId in doc.part.related_parts}.values())
    roots = blocks + [p.element for p in stories]
    report.cover_parts = [_part_info(k, t, doc.part.related_parts[r]) for (k, t), r in sorted(effective[0].items())]
    report.body_parts = [_part_info(k, t, doc.part.related_parts[r])
                         for (k, t), r in sorted(effective[body_index].items())]
    report.cover_blocks = len(blocks)
    report.cover_text = _text(blocks)[:_EXCERPT]
    report.styles = sorted(_style_ids(roots))

    removed = _strip_comments(roots)
    if removed:
        report.warnings.append("comments on the first page or in the headers and footers were removed")
    if any(el.tag in _TRACKED for root in roots for el in root.iter(*_TRACKED)):
        report.problems.append("the first page, headers or footers have tracked changes: accept or reject them in "
                               "Word and upload the template again")
    if any(True for b in blocks for _ in b.iter(qn("w:footnoteReference"), qn("w:endnoteReference"))):
        report.problems.append("the first page has footnotes or endnotes, which cannot be copied")
    images = set()
    for _, _, rId in _rel_attrs(blocks):
        rel = doc.part.rels.get(rId)
        if rel is None or rel.reltype in (RT.HEADER, RT.FOOTER):
            continue
        if rel.reltype not in _COPYABLE:
            report.problems.append(f"the first page holds a {_rel_name(rel.reltype)} object, which cannot be copied")
        elif not rel.is_external and rel.reltype == RT.IMAGE:
            images.add(hashlib.sha1(rel.target_part.blob).hexdigest())
    for part in stories:
        for rel in part.rels.values():
            if rel.reltype not in _COPYABLE:
                report.problems.append(f"{part.partname} holds a {_rel_name(rel.reltype)} object, which cannot be "
                                       "copied")
            elif not rel.is_external:
                images.add(hashlib.sha1(rel.target_part.blob).hexdigest())
        if _num_uses([part.element]):
            report.warnings.append(f"{part.partname} uses list numbering, which is not copied")
    report.images = len(images)
    if any(True for root in roots for _ in root.iter(qn("w:dataBinding"))):
        report.warnings.append("a content control on the first page or in a header or footer is bound to document "
                               "data; it shows its saved text")
    variables = _docvariables(roots)
    if variables:
        report.warnings.append(
            f"the first page, headers or footers have DOCVARIABLE fields ({', '.join(variables)}): Vault fills their "
            "values in. The language template's own document variables are not copied, so until then a field shows "
            "its saved text, or nothing once Word refreshes the fields")
    text = _text(roots)
    if not report.cover_text and not any(True for b in blocks for _ in b.iter(qn("w:drawing"), qn("w:pict"))):
        report.warnings.append("the first page is empty")
    script = _script_warning(language, text)
    if script:
        report.warnings.append(script)
    report.problems = list(dict.fromkeys(report.problems))
    report.warnings = list(dict.fromkeys(report.warnings))

    _drop_unused_rels(doc.part)
    package_path = Path(package_path)
    package_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(package_path))
    if preview_path is not None:
        write_preview(package_path, preview_path)
    return report


def write_preview(package_path: str | Path, preview_path: str | Path) -> None:
    """The package with two placeholder body pages, so the body's first-page and following headers show."""
    doc = Document(str(package_path))
    doc.add_paragraph(PREVIEW_TEXT[0]).add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph(PREVIEW_TEXT[1])
    Path(preview_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(preview_path))


def save_report(report: FrontMatterReport, path: str | Path) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(report.to_clean_dict(), indent=2, ensure_ascii=False), encoding="utf-8")


# ── The swap ──────────────────────────────────────────────────────────


@dataclass
class _Package:
    doc: object
    cover: Cover
    sects: list
    refs: list[Refs]

    @property
    def body_index(self) -> int:
        return 1 if self.cover.end == "section_break" else 0

    def role(self, index: int) -> int:
        """The package section a document section takes its headers and footers from."""
        return 0 if self.cover.end == "section_break" and index == 0 else self.body_index


def _open_package(package_path: str | Path) -> _Package:
    doc = Document(str(package_path))
    body = doc.element.body
    sects = section_props(body)
    return _Package(doc, find_cover(body), sects, effective_refs(sects))


def apply_front_matter(doc, package_path: str | Path, language: str) -> FrontMatterSwap:
    """Replace the document's first page, headers and footers with the package's.

    The document's body sections keep their own page size, margins and page numbering; they take the package's
    first body section's headers and footers and its different-first-page setting. The first page takes the
    package's whole section. Raises ``FrontMatterError`` when either first page cannot be found.
    """
    package = _open_package(package_path)
    swap = FrontMatterSwap()
    body = doc.element.body
    cover = find_cover(body)
    for el in cover.blocks:
        remove(el)
    if cover.next_block is not None:  # the package's first page ends with its own break
        for node in cover.next_block.findall(f"{qn('w:pPr')}/{qn('w:pageBreakBefore')}"):
            node.getparent().remove(node)
    swap.blocks_removed = len(cover.blocks)

    names = {b.get(qn("w:name")) for b in body.iter(qn("w:bookmarkStart"))}
    next_id = next_bookmark_id(body)
    copies = [copy.deepcopy(el) for el in package.cover.blocks]
    remaining = _blocks(body)
    anchor = remaining[0] if remaining else body.find(W_SECTPR)
    for el in copies:
        if anchor is not None:
            anchor.addprevious(el)
        else:
            body.append(el)
    swap.blocks_inserted = len(copies)

    images = _Images(doc.part.package)
    for el, name, rId in list(_rel_attrs(copies)):
        if el.tag not in REF_TAGS:  # header and footer references are rewired with the sections
            el.set(name, _copy_rel(package.doc.part, rId, doc.part, images))

    part_map: dict[str, str] = {}
    for refs in (package.refs[0], package.refs[package.body_index]):
        for (kind, _), rId in refs.items():
            if rId not in part_map:
                part_map[rId] = _copy_story(package.doc.part.related_parts[rId], kind, doc.part, images)
                swap.parts_added.append(str(doc.part.related_parts[part_map[rId]].partname))
    sects = section_props(body)
    for i, sect in enumerate(sects):
        role = package.role(i)
        set_refs(sect, {key: part_map[rId] for key, rId in package.refs[role].items()},
                 title_page(package.sects[role]))
    swap.sections_rewired = len(sects)
    swap.parts_dropped = _drop_unused_rels(doc.part)
    _set_even_and_odd(doc, even_and_odd(package.doc))

    stories = [doc.part.related_parts[rId].element for rId in dict.fromkeys(part_map.values())]
    copied_styles = _copy_styles(package.doc, doc, copies + stories, remaining, language, swap)
    _copy_numbering(package.doc, doc, copies + copied_styles, swap)
    if _num_uses(stories):
        swap.warnings.append("the language template's headers or footers use list numbering, which is not copied")
    _renumber_bookmarks(copies, names, next_id, swap)
    renumber_drawings(doc)
    return swap


# ── The check ─────────────────────────────────────────────────────────


def _canonical(el, part) -> bytes:
    """A block's XML with relationships named by their targets' content; ids that a copy renumbers are left out."""
    el = copy.deepcopy(el)
    for node, name, rId in list(_rel_attrs(el)):
        node.set(name, _identity(part, rId))
    for node in list(el.iter(qn("w:bookmarkStart"), qn("w:bookmarkEnd"), *REF_TAGS, qn("w:titlePg"))):
        remove(node)
    for node in el.iter(qn("wp:docPr")):
        node.set("id", "")
    for node in _num_uses([el]):
        node.set(qn("w:val"), "")
    return etree.tostring(el, method="c14n", exclusive=True)


def _story_signature(part) -> tuple:
    return part.blob, tuple(sorted((rId, _identity(part, rId)) for rId in part.rels))


def front_matter_problems(doc, package_path: str | Path) -> list[str]:
    """How the document's first page, headers and footers differ from the package's ([] when they are the same)."""
    try:
        package = _open_package(package_path)
        cover = find_cover(doc.element.body)
    except FrontMatterError as exc:
        return [str(exc)]
    problems = []
    ours = [_canonical(el, doc.part) for el in cover.blocks]
    theirs = [_canonical(el, package.doc.part) for el in package.cover.blocks]
    if ours != theirs:
        problems.append(f"the first page differs from the language template's ({len(ours)} blocks, the template "
                        f"has {len(theirs)})")
    sects = section_props(doc.element.body)
    for i, refs in enumerate(effective_refs(sects)):
        role = package.role(i)
        want = package.refs[role]
        label = "first page" if role == 0 and package.cover.end == "section_break" else "body"
        if set(refs) != set(want):
            problems.append(f"section {i + 1} ({label}) has headers and footers {sorted(refs)}, the language template "
                            f"{sorted(want)}")
            continue
        for key in sorted(want):
            if _story_signature(doc.part.related_parts[refs[key]]) != \
                    _story_signature(package.doc.part.related_parts[want[key]]):
                problems.append(f"section {i + 1} ({label}): the {key[1]} {key[0]} differs from the language "
                                "template's")
        if title_page(sects[i]) != title_page(package.sects[role]):
            problems.append(f"section {i + 1} ({label}): the different first page setting differs from the language "
                            "template's")
    if even_and_odd(doc) != even_and_odd(package.doc):
        problems.append("the even and odd pages setting differs from the language template's")
    for part in doc.part.package.iter_parts():
        if not hasattr(part, "element"):
            continue
        missing = sorted({rId for _, _, rId in _rel_attrs(part.element) if rId not in part.rels})
        if missing:
            problems.append(f"{part.partname} refers to missing relationships {missing}")
    return problems
