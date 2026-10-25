"""The translated final document (docs/translation/IMPLEMENTATION_PLAN.md, Phase T6).

The English final was rendered in PREPARING from the job's snapshot, with ``template_text.harvest_hook`` at the
renderer's body hook. The translation is rendered by the **same** call on translated inputs:
- ``claim``: a copy of the section drafts with each claim's translated text (``{{ref:...}}`` tokens kept);
- ``table_cell`` / ``table_header``: a copy of the source model with translated cells and headers (``unit.text``
  stays English, so the renderer still takes the row from its cells);
- ``xref``: a copy of the assembled document whose phrases are translated, the new number back in place of ``⟦1⟧``
  (still a REF field);
- ``template_text``: ``template_text.apply_hook`` at the harvest point, before the table of contents is rebuilt.

Then, before the document is saved (``finish``): the language template's first page, headers and footers when the
job has one (else the English ones stay), the Word language of the body, styles and defaults, and for Chinese and
Japanese the East Asian font. The post-render check compares the headers and footers with the language template's
(``front_matter_problems``), or with the English template's when there is none.

A segment without a usable translation (none, empty, or a cross-reference that lost its number) shows its English;
a warning names it, and its QA issue stays open for the reviewer.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Optional

from docx import Document
from docx.oxml.ns import qn

from app.schemas.v2 import (
    AssembledDocument,
    FrontMatterSwap,
    LanguageCode,
    NumberMap,
    RenderMode,
    RenderReport,
    SectionDraft,
    Segment,
    SegmentKind,
    SlotPlan,
    SourceDocument,
    TemplateModel,
)
from app.services.migration_v2.render import render_document
from app.services.migration_v2.render.xml import W_RPR, w_el

from .front_matter import FrontMatterError, apply_front_matter, find_cover, front_matter_problems
from .languages import get_language
from .segments import NUMBER_PLACEHOLDER
from .template_text import apply_hook, strip_tags

TRANSLATED = "translated"  # scope of the translated document's docx and render_report

DEFAULT_EAST_ASIA_FONTS = {"zh-CN": "Microsoft YaHei", "ja-JP": "Yu Gothic"}
# Render warnings that compare translated source headers with the template's English ones.
_ENGLISH_ONLY_WARNINGS = re.compile(r": source column '.*' is under template column '")
# rPr children that follow w:lang (CT_RPr order).
_AFTER_LANG = {qn(f"w:{t}") for t in ("eastAsianLayout", "specVanish", "oMath", "rPrChange")}


# ── Translated inputs ─────────────────────────────────────────────────


def usable(segment: Segment) -> Optional[str]:
    """The segment's translation, or None when there is nothing to show (no translation, or only markers)."""
    text = segment.target_text
    return text if text is not None and strip_tags(text).strip() else None


@dataclass
class Fallbacks:
    """Segments shown in English, and why (render warnings)."""

    kept: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def keep(self, segment: Segment, why: str = "no translation") -> None:
        self.kept.append(segment.segment_id)
        self.warnings.append(f"{segment.segment_id}: {why}; shown in English")


def translated_drafts(drafts: Iterable[SectionDraft], segments: Iterable[Segment], fallbacks: Fallbacks) -> list[SectionDraft]:
    """Copies of *drafts* whose claims carry their translations."""
    by_claim = {(s.locator.section_id, s.locator.slot_id, s.locator.claim_id): s
                for s in segments if s.kind == SegmentKind.CLAIM}
    out = []
    for draft in drafts:
        copy = draft.model_copy(deep=True)
        for slot in copy.slots:
            for claim in slot.claims:
                segment = by_claim.get((copy.target_section_id, slot.slot_id, claim.claim_id))
                if segment is None:
                    continue
                text = usable(segment)
                if text is None:
                    fallbacks.keep(segment)
                else:
                    claim.text = text
        out.append(copy)
    return out


def _joined_column(units: list, col: int, text_of) -> str:
    return " / ".join(dict.fromkeys(text_of(c) for u in units for c in u.table_ref.cells if c.col == col and text_of(c)))


def translated_source(source: SourceDocument, segments: Iterable[Segment], fallbacks: Fallbacks) -> SourceDocument:
    """A copy of *source* with translated table cells and headers. A header applies to every row of its table; a
    header that was only its rows' text joined (the renderer then drops it) is joined again from the translations."""
    english = {u.unit_id: u for u in source.iter_units()}
    copy = source.model_copy(deep=True)
    units = {u.unit_id: u for u in copy.iter_units()}
    tables: dict[str, list] = defaultdict(list)
    for unit in units.values():
        if unit.table_ref is not None:
            tables[unit.table_ref.table_id].append(unit)

    headers: list[Segment] = []
    for segment in segments:
        unit = units.get(segment.locator.unit_id or "")
        if segment.kind not in (SegmentKind.TABLE_CELL, SegmentKind.TABLE_HEADER) or unit is None or unit.table_ref is None:
            continue
        if segment.kind == SegmentKind.TABLE_HEADER:
            headers.append(segment)
            continue
        cell = next((c for c in unit.table_ref.cells if c.col == segment.locator.col), None)
        text = usable(segment)
        if cell is None:
            fallbacks.keep(segment, "table cell not found")
        elif text is None:
            fallbacks.keep(segment)
        elif segment.locator.paragraph is not None and segment.locator.paragraph < len(cell.paragraphs):
            cell.paragraphs[segment.locator.paragraph] = text
            cell.text = " ".join(cell.paragraphs)
        else:
            cell.text = text

    for segment in headers:
        table = units[segment.locator.unit_id].table_ref.table_id
        col = segment.locator.col
        rows = tables[table]
        english_rows = [english[u.unit_id] for u in rows]
        if english_rows[0].table_ref.header_cells[col] == _joined_column(english_rows, col, lambda c: c.text):
            text = _joined_column(rows, col, lambda c: c.text) or usable(segment)
        else:
            text = usable(segment)
        if text is None:
            fallbacks.keep(segment)
            continue
        for unit in rows:
            if col < len(unit.table_ref.header_cells):
                unit.table_ref.header_cells[col] = text
    return copy


def translated_assembly(assembled: Optional[AssembledDocument], segments: Iterable[Segment],
                        fallbacks: Fallbacks) -> Optional[AssembledDocument]:
    """A copy of *assembled* whose reference phrases are translated, the new number where ``⟦1⟧`` was."""
    if assembled is None:
        return None
    copy = assembled.model_copy(deep=True)
    for segment in segments:
        if segment.kind != SegmentKind.XREF or segment.locator.ref_index >= len(copy.refs):
            continue
        ref = copy.refs[segment.locator.ref_index]
        text = usable(segment)
        if text is None:
            fallbacks.keep(segment)
        elif NUMBER_PLACEHOLDER in segment.source_text:
            if text.count(NUMBER_PLACEHOLDER) != 1:
                fallbacks.keep(segment, f"the reference's number {NUMBER_PLACEHOLDER} is lost or doubled")
                continue
            at = text.index(NUMBER_PLACEHOLDER)
            ref.text = strip_tags(text[:at]) + ref.number + strip_tags(text[at + len(NUMBER_PLACEHOLDER):])
            ref.number_at = len(strip_tags(text[:at]))
        else:
            ref.text = strip_tags(text)
    return copy


# ── Language and fonts ────────────────────────────────────────────────


def _cover_elements(doc) -> set:
    """Every element of the first page (kept exactly as the language template or the English template has it)."""
    try:
        cover = find_cover(doc.element.body)
    except FrontMatterError:
        return set()
    return {el for block in cover.blocks for el in block.iter()}


def _default_rpr(doc):
    """``docDefaults/rPrDefault/rPr``, created when missing."""
    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = w_el("docDefaults")
        styles.insert(0, defaults)
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = w_el("rPrDefault")
        defaults.insert(0, rpr_default)
    rpr = rpr_default.find(W_RPR)
    if rpr is None:
        rpr = w_el("rPr")
        rpr_default.append(rpr)
    return rpr


def _set_lang(node, code: str, east_asian: bool) -> None:
    node.set(qn("w:val"), code)
    if east_asian:
        node.set(qn("w:eastAsia"), code)


def set_language(doc, language: LanguageCode | str) -> None:
    """The body (first page left out), the styles and the document defaults are in *language* for Word's
    proofing; for Chinese and Japanese the East Asian language too. Headers and footers are not touched."""
    lang = get_language(language)
    code, east = lang.word_lang, lang.east_asian
    cover = _cover_elements(doc)
    for node in doc.element.body.iter(qn("w:lang")):
        if node not in cover:
            _set_lang(node, code, east)
    for node in doc.styles.element.iter(qn("w:lang")):
        _set_lang(node, code, east)
    rpr = _default_rpr(doc)
    node = rpr.find(qn("w:lang"))
    if node is None:
        node = w_el("lang")
        follower = next((c for c in rpr if c.tag in _AFTER_LANG), None)
        follower.addprevious(node) if follower is not None else rpr.append(node)
    _set_lang(node, code, east)
    for node in doc.settings.element.iter(qn("w:themeFontLang")):
        _set_lang(node, code, east)


def _east_asia_of(doc) -> Optional[str]:
    """The East Asian font a document names in its defaults or its Normal style."""
    defaults = doc.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:rFonts')}")
    normal = next((s for s in doc.styles.element.findall(qn("w:style"))
                   if s.get(qn("w:type")) == "paragraph" and s.get(qn("w:default")) in ("1", "true")), None)
    fonts = [defaults, normal.find(f"{qn('w:rPr')}/{qn('w:rFonts')}") if normal is not None else None]
    return next((f.get(qn("w:eastAsia")) for f in fonts if f is not None and f.get(qn("w:eastAsia"))), None)


def east_asia_font(language: LanguageCode | str, package_path: Optional[str | Path] = None,
                   settings: Any = None) -> Optional[str]:
    """The East Asian font for Chinese and Japanese: the language template's, else the settings' fallback."""
    lang = get_language(language)
    if not lang.east_asian:
        return None
    if package_path is not None:
        found = _east_asia_of(Document(str(package_path)))
        if found:
            return found
    fonts = getattr(settings, "translation_east_asia_fonts", None) or DEFAULT_EAST_ASIA_FONTS
    return fonts.get(lang.code.value)


def _set_east_asia(fonts, font: str) -> None:
    fonts.set(qn("w:eastAsia"), font)
    fonts.attrib.pop(qn("w:eastAsiaTheme"), None)


def set_east_asia_font(doc, font: str) -> None:
    """*font* for East Asian text: in the document defaults, and wherever a style or a body run (first page left
    out) names an East Asian font of its own."""
    rpr = _default_rpr(doc)
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = w_el("rFonts")
        rpr.insert(0, fonts)
    _set_east_asia(fonts, font)
    cover = _cover_elements(doc)
    named = (qn("w:eastAsia"), qn("w:eastAsiaTheme"))
    for root, skip in ((doc.styles.element, set()), (doc.element.body, cover)):
        for node in root.iter(qn("w:rFonts")):
            if node not in skip and any(a in node.attrib for a in named):
                _set_east_asia(node, font)


# ── The render ────────────────────────────────────────────────────────


@dataclass
class TranslatedRender:
    report: RenderReport
    applied: dict[str, int]                        # segment kind → segments written in the target language
    kept: list[str]                                # segments shown in English
    swap: Optional[FrontMatterSwap] = None
    east_asia_font: Optional[str] = None


def render_translation(
    job_id: str,
    language: LanguageCode | str,
    segments: list[Segment],
    template: TemplateModel,
    drafts: list[SectionDraft],
    source: SourceDocument,
    out_path: str | Path,
    *,
    version: int = 1,
    slot_plan: Optional[SlotPlan] = None,
    accepted_gaps: Iterable[str] = (),
    assembled: Optional[AssembledDocument] = None,
    number_map: Optional[NumberMap] = None,
    front_matter: Optional[str | Path] = None,
    settings: Any = None,
    page_counter=None,
) -> TranslatedRender:
    """Render the translated final document of *segments* (with their latest translations) at *out_path*.

    *front_matter* is the job's language-template package; without one the English first page, headers and footers
    stay. Raises ``FrontMatterError`` when the package's or the document's first page cannot be found, and the
    renderer's ``RenderError`` / ``RenderBlocked``.
    """
    fallbacks = Fallbacks()
    t_drafts = translated_drafts(drafts, segments, fallbacks)
    t_source = translated_source(source, segments, fallbacks)
    t_assembled = translated_assembly(assembled, segments, fallbacks)
    template_segments = [s for s in segments if s.kind == SegmentKind.TEMPLATE_TEXT]
    hook_warnings: list[str] = []
    written: list[str] = []
    font = east_asia_font(language, front_matter, settings)
    swaps: list[FrontMatterSwap] = []

    def finish(doc) -> None:
        if front_matter is not None:
            swaps.append(apply_front_matter(doc, front_matter, get_language(language).word_lang))
        set_language(doc, language)
        if font:
            set_east_asia_font(doc, font)

    header_check = (lambda doc: front_matter_problems(doc, front_matter)) if front_matter is not None else None
    report = render_document(
        template, t_drafts, t_source, out_path, job_id=job_id, version=version, slot_plan=slot_plan,
        mode=RenderMode.FINAL, accepted_gaps=accepted_gaps, page_counter=page_counter, assembled=t_assembled,
        number_map=number_map, body_hook=apply_hook(template_segments, hook_warnings, written), finish=finish,
        header_check=header_check,
    )
    swap = swaps[0] if swaps else None
    kept = fallbacks.kept + [s.segment_id for s in template_segments if s.segment_id not in written]
    warnings = [w for w in report.warnings if not _ENGLISH_ONLY_WARNINGS.search(w)]
    warnings += fallbacks.warnings + hook_warnings + [f"first page, headers and footers: {w}" for w in
                                                      (swap.warnings if swap else [])]
    report.warnings = list(dict.fromkeys(warnings))
    applied: dict[str, int] = {}
    for segment in segments:
        if segment.segment_id not in kept:
            applied[segment.kind.value] = applied.get(segment.kind.value, 0) + 1
    return TranslatedRender(report, applied, kept, swap, font)
