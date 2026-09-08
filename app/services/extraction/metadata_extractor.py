"""SOP Metadata Extractor.

Extracts Document Title, Document Name, Document Number, Version, and Type
from the first-page preamble table (or header text) of an SOP document.
Generates a deterministic document_id; gpdat_version comes from sop_records
history (see ``SopStore.get_next_version``), not a separate counter file.
"""

import re
from pathlib import Path

import fitz  # PyMuPDF
import pdfplumber
import docx
from loguru import logger


class SOPMetadataExtractor:
    """Service to extract SOP identification metadata from first page tables."""

    _TITLE_KEYS = ("document title", "title")
    _NAME_KEYS = ("document name", "doc name")
    _NUMBER_KEYS = ("document number", "doc number", "doc no", "sop number", "document id")
    _VERSION_KEYS = ("version", "ver", "rev", "revision")
    _TYPE_KEYS = ("type/subtype", "type", "subtype", "document type", "doc type")

    # A trailing cell that itself packs a second "Label: value" pair, e.g. a
    # "Document ID: X" | "Version: Y" row split across two table columns.
    _LABEL_VALUE_RE = re.compile(r"^([A-Za-z][A-Za-z /]{1,30}):\s*(.+)$", re.DOTALL)

    @classmethod
    def sanitize_string(cls, val: str) -> str:
        """Sanitize a string to alphanumeric, hyphen, underscore, or dot."""
        if not val:
            return ""
        clean = re.sub(r"[^A-Za-z0-9_\-.]", "_", val.strip())
        clean = re.sub(r"_+", "_", clean)
        return clean.strip("_")

    @classmethod
    def generate_document_id(
        cls,
        doc_name: str,
        doc_number: str,
        doc_version: str,
        fallback_filename: str = "",
    ) -> str:
        """Generate a deterministic document ID from SOP metadata fields."""
        clean_num = cls.sanitize_string(doc_number)
        clean_name = cls.sanitize_string(doc_name)
        clean_ver = cls.sanitize_string(doc_version)

        if not clean_name and fallback_filename:
            stem = Path(fallback_filename).stem
            clean_name = cls.sanitize_string(stem)

        if not clean_num:
            clean_num = "DOC"
        if not clean_name:
            clean_name = "UNKNOWN"
        if not clean_ver:
            clean_ver = "1.0"

        return f"{clean_num}_{clean_name}_v{clean_ver}"

    @classmethod
    def _matches(cls, key: str, keywords: tuple[str, ...]) -> bool:
        """Whole-word keyword match so short abbreviations (e.g. 'rev', 'ver')
        don't false-positive inside unrelated words like 'Reviewer' or 'Server'.
        """
        return any(re.search(rf"\b{re.escape(k)}\b", key) for k in keywords)

    @classmethod
    def _clean_cell_value(cls, text: str) -> str:
        """Strip stray watermark-stamp fragments that bled into a table cell.

        A diagonal/rotated watermark stamp intersecting a narrow preamble
        cell surfaces as a lone 1-2 letter line mixed into the cell's real
        text. No genuine preamble value wraps onto a line by itself like
        that, so any such line is dropped rather than matched against a
        fixed list of watermark words.
        """
        lines = [
            line for raw_line in text.split("\n")
            if (line := raw_line.strip()) and not (len(line) <= 2 and line.isalpha())
        ]
        return "\n".join(lines)

    @classmethod
    def _row_pairs(cls, cells: list) -> list[tuple[str, str]]:
        """Turn a table row into (label, value) pairs.

        Most preamble rows are a plain 2-cell label/value pair. Some SOPs pack
        a *second* label/value pair into a trailing cell of the same row (e.g.
        "Document ID: X" next to "Version: Y"); those are split out too so
        that value isn't silently dropped.
        """
        texts = [str(c or "").strip() for c in cells]
        pairs: list[tuple[str, str]] = []
        if len(texts) >= 2 and texts[0]:
            pairs.append((texts[0], cls._clean_cell_value(texts[1])))
        for extra in texts[2:]:
            m = cls._LABEL_VALUE_RE.match(extra)
            if m:
                pairs.append((m.group(1).strip(), cls._clean_cell_value(m.group(2).strip())))
        return pairs

    @classmethod
    def _classify_row(cls, cells: list, values: dict[str, str]) -> None:
        """Classify a table row's (label, value) pairs into ``values`` in place."""
        for key_raw, val in cls._row_pairs(cells):
            key = key_raw.lower()
            if not values["title"] and cls._matches(key, cls._TITLE_KEYS):
                values["title"] = val
            elif not values["name"] and cls._matches(key, cls._NAME_KEYS):
                values["name"] = val
            elif not values["number"] and cls._matches(key, cls._NUMBER_KEYS):
                values["number"] = val
            elif not values["version"] and cls._matches(key, cls._VERSION_KEYS):
                values["version"] = val
            elif not values["type"] and cls._matches(key, cls._TYPE_KEYS):
                values["type"] = val

    @classmethod
    def extract_from_file(
        cls,
        file_path: str,
        fallback_filename: str = "",
    ) -> tuple[str, str, str, str, str]:
        """Extract (document_title, document_name, document_number, document_version,
        document_type) from file.

        ``document_title`` comes from a "Document Title"/"Title" row,
        ``document_name`` from a "Document Name"/"Doc Name" row, and
        ``document_number`` from "Document Number"/"Document ID" (some SOPs
        use "Document ID" instead of a name/number pair).
        """
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext == ".pdf":
            return cls._extract_from_pdf(path)
        elif ext == ".docx":
            return cls._extract_from_docx(path)
        else:
            return "", "", "", "", ""

    @classmethod
    def _extract_from_pdf(cls, file_path: Path) -> tuple[str, str, str, str, str]:
        values = {"title": "", "name": "", "number": "", "version": "", "type": ""}

        try:
            with pdfplumber.open(file_path) as pdf:
                if len(pdf.pages) > 0:
                    page0 = pdf.pages[0]
                    tables = page0.extract_tables()
                    for table in tables:
                        for row in table:
                            if not row:
                                continue
                            cls._classify_row(row, values)
        except Exception as e:  # noqa: BLE001
            logger.warning(
                "PDF table metadata extraction failed for '{f}': {e}", f=file_path.name, e=e
            )

        doc_title, doc_name, doc_number, doc_version, doc_type = (
            values["title"], values["name"], values["number"], values["version"], values["type"]
        )

        if not (doc_title and doc_name and doc_number and doc_version):
            try:
                doc = fitz.open(file_path)
                if len(doc) > 0:
                    text = doc[0].get_text("text")
                    if not doc_number:
                        m = re.search(
                            r"(?i)(?:Document\s*Number|Document\s*ID|Number):\s*([A-Za-z0-9\-_.]{3,})",
                            text,
                        )
                        if m:
                            doc_number = m.group(1).strip()
                    if not doc_version:
                        m = re.search(
                            r"(?i)(?:Version|Ver|Rev|Revision):\s*([0-9.]+)",
                            text,
                        )
                        if m:
                            doc_version = m.group(1).strip()
                    if not doc_title:
                        m = re.search(
                            r"(?i)(?:Document\s*Title|Title):\s*([^\r\n]+)",
                            text,
                        )
                        if m:
                            doc_title = m.group(1).strip()
                    if not doc_name:
                        m = re.search(
                            r"(?i)Document\s*Name:\s*([^\r\n]+)",
                            text,
                        )
                        if m:
                            doc_name = m.group(1).strip()
                doc.close()
            except Exception as e:  # noqa: BLE001
                logger.warning(
                    "PDF text metadata fallback failed for '{f}': {e}", f=file_path.name, e=e
                )

        return doc_title, doc_name, doc_number, doc_version, doc_type

    @classmethod
    def _extract_from_docx(cls, file_path: Path) -> tuple[str, str, str, str, str]:
        values = {"title": "", "name": "", "number": "", "version": "", "type": ""}
        doc_title = doc_name = doc_number = doc_version = doc_type = ""

        try:
            doc = docx.Document(file_path)
            for table in doc.tables[:3]:
                for row in table.rows:
                    cls._classify_row([c.text for c in row.cells], values)

            doc_title, doc_name, doc_number, doc_version, doc_type = (
                values["title"], values["name"], values["number"], values["version"], values["type"]
            )

            if not (doc_title and doc_name and doc_number and doc_version):
                full_text = "\n".join(p.text for p in doc.paragraphs[:20])
                if not doc_number:
                    m = re.search(
                        r"(?i)(?:Document\s*Number|Document\s*ID|Number):\s*([A-Za-z0-9\-_.]{3,})",
                        full_text,
                    )
                    if m:
                        doc_number = m.group(1).strip()
                if not doc_version:
                    m = re.search(
                        r"(?i)(?:Version|Ver|Rev|Revision):\s*([0-9.]+)",
                        full_text,
                    )
                    if m:
                        doc_version = m.group(1).strip()
                if not doc_title:
                    m = re.search(
                        r"(?i)(?:Document\s*Title|Title):\s*([^\r\n]+)",
                        full_text,
                    )
                    if m:
                        doc_title = m.group(1).strip()
                if not doc_name:
                    m = re.search(
                        r"(?i)Document\s*Name:\s*([^\r\n]+)",
                        full_text,
                    )
                    if m:
                        doc_name = m.group(1).strip()
        except Exception as e:  # noqa: BLE001
            logger.warning(
                "DOCX metadata extraction failed for '{f}': {e}", f=file_path.name, e=e
            )

        return doc_title, doc_name, doc_number, doc_version, doc_type
